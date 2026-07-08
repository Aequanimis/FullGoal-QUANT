from __future__ import annotations

import csv
import math
import os
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path.home() / "Desktop" / "Fullgoal" / "课题研究"
DEFAULT_INPUT_PDF = (
    PROJECT_ROOT
    / "处理后数据"
    / "境外策略ETF数据"
    / "原始数据"
    / "screener-results-2026-07-03.pdf"
)
OUTPUT_DIR = PROJECT_ROOT / "处理后数据" / "境外策略ETF数据" / "处理结果"
INPUT_PDF = Path(os.environ.get("ETF_COM_INPUT_PDF", str(DEFAULT_INPUT_PDF)))
OUTPUT_XLSX = Path(os.environ.get("ETF_COM_OUTPUT_XLSX", str(OUTPUT_DIR / "etf_com_master_extracted_v1.xlsx")))
OUTPUT_CSV = Path(os.environ.get("ETF_COM_OUTPUT_CSV", str(OUTPUT_DIR / "etf_com_master_extracted_v1.csv")))
OUTPUT_LOG = Path(os.environ.get("ETF_COM_OUTPUT_LOG", str(OUTPUT_DIR / "extraction_log_v1.txt")))
COMPARE_PREVIOUS_CSV = Path(
    os.environ.get("ETF_COM_COMPARE_PREVIOUS_CSV", str(OUTPUT_DIR / "etf_com_master_extracted_20260706.csv"))
)


MODULES = {
    "classification": {
        "pages": range(0, 5),
        "fields": [
            "ticker",
            "fund_name",
            "issuer",
            "asset_class",
            "category",
            "strategy",
            "niche",
            "segments",
            "selection_criteria",
            "weighting_scheme",
        ],
    },
    "liquidity": {
        "pages": range(5, 8),
        "fields": [
            "ticker",
            "fund_name",
            "aum_raw",
            "expense_ratio_raw",
            "avg_daily_dollar_vol_raw",
            "avg_daily_share_vol_raw",
            "spread_data_pct_raw",
            "dividend_yield_raw",
            "pe_raw",
            "pb_raw",
        ],
    },
    "flows_returns": {
        "pages": range(8, 11),
        "fields": [
            "ticker",
            "fund_name",
            "one_year_fund_flow_raw",
            "ytd_fund_flow_raw",
            "one_month_fund_flow_raw",
            "one_year_total_return_raw",
            "three_year_total_return_raw",
            "five_year_total_return_raw",
            "etn_raw",
            "leveraged_raw",
        ],
    },
    "geo": {
        "pages": range(11, 14),
        "fields": [
            "ticker",
            "fund_name",
            "leveraged_ratio",
            "inverse_raw",
            "region",
            "geography",
            "economic_development",
        ],
    },
}


MASTER_FIELDS = [
    "ticker",
    "fund_name",
    "issuer",
    "asset_class",
    "category",
    "strategy",
    "niche",
    "segments",
    "selection_criteria",
    "weighting_scheme",
    "aum_raw",
    "aum_usd",
    "expense_ratio_raw",
    "expense_ratio_pct",
    "avg_daily_dollar_vol_raw",
    "avg_daily_dollar_vol_usd",
    "avg_daily_share_vol_raw",
    "avg_daily_share_vol",
    "spread_data_pct_raw",
    "spread_data_pct",
    "dividend_yield_raw",
    "dividend_yield_pct",
    "pe_raw",
    "pe",
    "pb_raw",
    "pb",
    "one_year_fund_flow_raw",
    "one_year_fund_flow_usd",
    "ytd_fund_flow_raw",
    "ytd_fund_flow_usd",
    "one_month_fund_flow_raw",
    "one_month_fund_flow_usd",
    "one_year_total_return_raw",
    "one_year_total_return_pct",
    "three_year_total_return_raw",
    "three_year_total_return_pct",
    "five_year_total_return_raw",
    "five_year_total_return_pct",
    "etn_raw",
    "etn",
    "leveraged_raw",
    "leveraged",
    "leveraged_ratio",
    "inverse_raw",
    "inverse",
    "region",
    "geography",
    "economic_development",
    "preliminary_pool_type",
    "preliminary_strategy_level_1",
    "preliminary_strategy_level_2",
    "classification_reason",
    "is_excluded_by_features",
    "is_strategy_candidate",
    "is_extended_strategy_candidate",
    "source_pdf",
    "source_pages",
    "extract_time",
    "text_quality_warning",
    "parse_warning",
]


WATERMARK_LINE_RE = re.compile(r"^(?:[etfcom.\s]|etf\.com)+$", re.IGNORECASE)


def normalize_space(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def clean_text_cell(value: Any) -> str:
    """Remove obvious ETF.com watermark fragments while preserving readable text."""
    if value is None:
        return ""
    text = str(value).replace("\r", "\n")
    parts: List[str] = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if WATERMARK_LINE_RE.match(line) and len(line.replace(" ", "")) <= 7:
            continue
        parts.append(line)
    text = normalize_space(" ".join(parts))
    text = re.sub(r"\betf\s*\.\s*com\b", "", text, flags=re.IGNORECASE)
    text = normalize_space(text)
    return text


def keep_non_watermark_object(obj: Dict[str, Any]) -> bool:
    """Keep table text and remove ETF.com watermark glyphs before table extraction.

    The current ETF.com PDF stores watermark letters as text-layer chars, not OCR.
    They are upright, but use a different font/size/color profile:
    Helvetica-Bold, large size, pale gray. Filtering these chars before
    table extraction prevents strings such as 'fAlternative', 'Ut.S.', etc.
    """
    if obj.get("object_type") != "char":
        return True
    if not obj.get("upright", True):
        return False
    size = float(obj.get("size", 0) or 0)
    font = obj.get("fontname", "") or ""
    color = str(obj.get("non_stroking_color"))
    if "Helvetica-Bold" in font and size > 10 and color == "(0.902,)":
        return False
    return True


def clean_ticker(value: Any) -> str:
    """Ticker cells can contain lowercase watermark chars: e.g. AAAeAt -> AAAA."""
    if value is None:
        return ""
    compact = re.sub(r"[^A-Z0-9]", "", str(value))
    if 1 <= len(compact) <= 6:
        return compact
    m = re.search(r"[A-Z][A-Z0-9]{0,5}", compact)
    return m.group(0) if m else ""


def clean_numeric_raw(value: Any, kind: str) -> str:
    text = clean_text_cell(value)
    if not text:
        return ""
    text = text.replace("−", "-").replace("—", "-").replace("–", "-")
    if kind == "percent":
        stripped = re.sub(r"[A-Za-z]", "", text)
        m = re.search(r"-?\s*(?:undefined|[\d,]+(?:\.\d+)?)\s*%", stripped, re.IGNORECASE)
        if m:
            return normalize_space(m.group(0).replace(" ", ""))
        if text.strip() in {"-", "--"}:
            return "-"
    elif kind == "amount":
        # Preserve K/M/B suffix only when it appears at the end of the numeric token.
        m = re.search(r"-?\s*\$?\s*[\d,]+(?:\.\d+)?\s*[KMB]?", text, re.IGNORECASE)
        if m:
            return normalize_space(m.group(0).replace(" ", "")).upper()
        if text.strip() in {"-", "--"}:
            return "-"
    elif kind == "number":
        m = re.search(r"-?\d+(?:\.\d+)?", text.replace(",", ""))
        if m:
            return m.group(0)
        if text.strip() in {"-", "--"}:
            return "-"
    return text


def parse_amount(raw: Any) -> Optional[float]:
    text = clean_numeric_raw(raw, "amount")
    if not text or text in {"-", "--"}:
        return None
    m = re.search(r"(-?)\$?([\d,]+(?:\.\d+)?)([KMB]?)", text, re.IGNORECASE)
    if not m:
        return None
    sign = -1 if m.group(1) == "-" else 1
    value = float(m.group(2).replace(",", ""))
    suffix = m.group(3).upper()
    multiplier = {"": 1, "K": 1_000, "M": 1_000_000, "B": 1_000_000_000}.get(suffix, 1)
    return sign * value * multiplier


def parse_percent(raw: Any) -> Optional[float]:
    text = clean_numeric_raw(raw, "percent")
    if not text or text in {"-", "--"} or "undefined" in text.lower():
        return None
    m = re.search(r"(-?[\d,]+(?:\.\d+)?)\s*%", text)
    if not m:
        return None
    return float(m.group(1).replace(",", ""))


def parse_number(raw: Any) -> Optional[float]:
    text = clean_numeric_raw(raw, "number")
    if not text or text in {"-", "--"}:
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None


def parse_bool(raw: Any) -> Optional[bool]:
    text = clean_text_cell(raw).strip().lower()
    text = re.sub(r"[^a-z]", "", text)
    if text.endswith("no"):
        text = "no"
    if text.endswith("yes"):
        text = "yes"
    if text == "yes":
        return True
    if text == "no":
        return False
    return None


def clean_leveraged_ratio(raw: Any) -> str:
    text = clean_text_cell(raw)
    if not text:
        return ""
    text = text.replace("−", "-").replace("—", "-").replace("–", "-")
    if re.fullmatch(r"[etfcom.\s-]+", text, flags=re.IGNORECASE):
        return "-"
    m = re.search(r"\b\d+(?:\.\d+)?X\b", text, flags=re.IGNORECASE)
    if m:
        return m.group(0).upper()
    compact_alpha = re.sub(r"[^A-Za-z]", "", text).lower()
    if compact_alpha in {"no", "nfo", "nto", "eno", "fno", "ono"}:
        return "-"
    if text.strip() in {"-", "--"}:
        return "-"
    return text.strip()


def detect_page_fields(module_name: str, header_row: List[Any], default_fields: List[str]) -> Tuple[List[str], Optional[str]]:
    header_text = " ".join(clean_text_cell(c) for c in header_row)
    lower = header_text.lower()
    compact = re.sub(r"[^a-z0-9]", "", lower)
    if module_name == "liquidity":
        has_one_year_flow = "yrfundflow" in compact
        has_avg_share = "avgdailyshare" in compact
        if has_one_year_flow and not has_avg_share:
            return [
                "ticker",
                "fund_name",
                "aum_raw",
                "expense_ratio_raw",
                "avg_daily_dollar_vol_raw",
                "spread_data_pct_raw",
                "pe_raw",
                "pb_raw",
                "dividend_yield_raw",
                "one_year_fund_flow_raw",
            ], "liquidity page includes 1 Yr Fund Flow and omits Avg. Daily Share Vol.; mapped by actual header"
        if has_avg_share:
            return default_fields, None
    if module_name == "flows_returns":
        if "leveragedratio" in compact and "yrfundflow" not in compact:
            return [
                "ticker",
                "fund_name",
                "ytd_fund_flow_raw",
                "one_month_fund_flow_raw",
                "one_year_total_return_raw",
                "three_year_total_return_raw",
                "five_year_total_return_raw",
                "etn_raw",
                "leveraged_raw",
                "leveraged_ratio",
            ], "flows_returns page starts at YTD Fund Flow and includes Leveraged Ratio; mapped by actual header"
        return default_fields, None
    if module_name == "geo":
        if "inverse" in lower and "region" in lower and "geography" in lower and "avgdailyshare" in compact:
            return [
                "ticker",
                "fund_name",
                "inverse_raw",
                "region",
                "geography",
                "economic_development",
                "avg_daily_share_vol_raw",
            ], "geo page includes Avg. Daily Share Vol.; mapped by actual header"
        if "inverse" in lower and "region" in lower and "geography" in lower:
            return ["ticker", "fund_name", "inverse_raw", "region", "geography", "economic_development"], None
        if "Inverse" in header_text and "Leveraged Ratio" not in header_text and len(header_row) <= 3:
            return ["ticker", "fund_name", "inverse_raw"], (
                "only contains Inverse column; leveraged_ratio/region/geography/economic_development left blank"
            )
    return default_fields, None


def extract_module(pdf: pdfplumber.PDF, module_name: str) -> Tuple[Dict[str, Dict[str, Any]], List[str], List[str]]:
    config = MODULES[module_name]
    fields = config["fields"]
    rows_by_ticker: Dict[str, Dict[str, Any]] = {}
    duplicate_tickers: List[str] = []
    warnings: List[str] = []
    for page_idx in config["pages"]:
        if page_idx >= len(pdf.pages):
            warnings.append(f"{module_name}: page {page_idx + 1} out of range")
            continue
        filtered_page = pdf.pages[page_idx].filter(keep_non_watermark_object)
        found_tables = filtered_page.find_tables() or []
        tables = []
        if found_tables:
            # Crop to the largest table region after watermark filtering.
            # This keeps header/footer/disclaimer text out of extraction.
            table_obj = max(found_tables, key=lambda t: (t.bbox[2] - t.bbox[0]) * (t.bbox[3] - t.bbox[1]))
            cropped = filtered_page.crop(table_obj.bbox)
            tables = cropped.extract_tables() or [table_obj.extract()]
        else:
            tables = filtered_page.extract_tables() or []
        if not tables:
            warnings.append(f"{module_name}: no table extracted on page {page_idx + 1}")
            continue
        table = max(tables, key=len)
        page_fields, mapping_note = detect_page_fields(module_name, table[0] if table else [], fields)
        if mapping_note:
            warnings.append(f"{module_name}: page {page_idx + 1} {mapping_note}")
        for row in table[1:]:
            if not row or all(not clean_text_cell(c) for c in row):
                continue
            raw_ticker = row[0] if len(row) > 0 else ""
            ticker = clean_ticker(raw_ticker)
            if not ticker or ticker.lower().startswith("ticker") or ticker in {"PAGE"}:
                continue
            rec: Dict[str, Any] = {"ticker": ticker, "_source_page": page_idx + 1}
            for i, field in enumerate(page_fields[1:], start=1):
                rec[field] = clean_text_cell(row[i]) if i < len(row) else ""
            if ticker in rows_by_ticker:
                duplicate_tickers.append(ticker)
                warnings.append(f"{module_name}: duplicate ticker {ticker} on page {page_idx + 1}")
                # Keep the row with more populated fields.
                existing_score = sum(bool(v) for k, v in rows_by_ticker[ticker].items() if not k.startswith("_"))
                new_score = sum(bool(v) for k, v in rec.items() if not k.startswith("_"))
                if new_score > existing_score:
                    rows_by_ticker[ticker] = rec
            else:
                rows_by_ticker[ticker] = rec
    return rows_by_ticker, duplicate_tickers, warnings


def normalize_raw_fields(row: Dict[str, Any]) -> None:
    amount_fields = [
        "aum_raw",
        "avg_daily_dollar_vol_raw",
        "avg_daily_share_vol_raw",
        "one_year_fund_flow_raw",
        "ytd_fund_flow_raw",
        "one_month_fund_flow_raw",
    ]
    percent_fields = [
        "expense_ratio_raw",
        "spread_data_pct_raw",
        "dividend_yield_raw",
        "one_year_total_return_raw",
        "three_year_total_return_raw",
        "five_year_total_return_raw",
    ]
    number_fields = ["pe_raw", "pb_raw"]
    for field in amount_fields:
        if field in row:
            row[field] = clean_numeric_raw(row[field], "amount")
    for field in percent_fields:
        if field in row:
            row[field] = clean_numeric_raw(row[field], "percent")
    for field in number_fields:
        if field in row:
            row[field] = clean_numeric_raw(row[field], "number")
    if "leveraged_ratio" in row:
        row["leveraged_ratio"] = clean_leveraged_ratio(row.get("leveraged_ratio"))
    for field in ["etn_raw", "leveraged_raw", "inverse_raw"]:
        if field in row:
            val = parse_bool(row[field])
            row[field] = "Yes" if val is True else "No" if val is False else clean_text_cell(row[field])


def add_cleaned_values(row: Dict[str, Any]) -> None:
    row["aum_usd"] = parse_amount(row.get("aum_raw"))
    row["avg_daily_dollar_vol_usd"] = parse_amount(row.get("avg_daily_dollar_vol_raw"))
    row["avg_daily_share_vol"] = parse_amount(row.get("avg_daily_share_vol_raw"))
    row["one_year_fund_flow_usd"] = parse_amount(row.get("one_year_fund_flow_raw"))
    row["ytd_fund_flow_usd"] = parse_amount(row.get("ytd_fund_flow_raw"))
    row["one_month_fund_flow_usd"] = parse_amount(row.get("one_month_fund_flow_raw"))
    row["expense_ratio_pct"] = parse_percent(row.get("expense_ratio_raw"))
    row["spread_data_pct"] = parse_percent(row.get("spread_data_pct_raw"))
    row["dividend_yield_pct"] = parse_percent(row.get("dividend_yield_raw"))
    row["one_year_total_return_pct"] = parse_percent(row.get("one_year_total_return_raw"))
    row["three_year_total_return_pct"] = parse_percent(row.get("three_year_total_return_raw"))
    row["five_year_total_return_pct"] = parse_percent(row.get("five_year_total_return_raw"))
    row["pe"] = parse_number(row.get("pe_raw"))
    row["pb"] = parse_number(row.get("pb_raw"))
    row["etn"] = parse_bool(row.get("etn_raw"))
    row["leveraged"] = parse_bool(row.get("leveraged_raw"))
    row["inverse"] = parse_bool(row.get("inverse_raw"))


def detect_text_quality_warning(row: Dict[str, Any]) -> str:
    warnings: List[str] = []
    fund_name = str(row.get("fund_name", "") or "")
    geography = str(row.get("geography", "") or "")
    asset_class = str(row.get("asset_class", "") or "")
    standard_asset_classes = {
        "Equity",
        "Fixed Income",
        "Commodities",
        "Alternatives",
        "Asset Allocation",
        "Currency",
    }

    if re.search(r"^(?:[fetco]\s+|[fetco](?=[A-Z]))", fund_name):
        warnings.append("fund_name 开头疑似残留孤立水印字符")
    if re.search(r"(?:\bf\s*\.\s*|\be\s+|\bt\s+|tf\s*\.\s*|etf\s*\.\s*com)", fund_name, flags=re.IGNORECASE):
        warnings.append("fund_name 中疑似残留水印片段")
    if "ETF" not in fund_name and re.search(r"\b(E T F|ET F|E TF)\b", fund_name, flags=re.IGNORECASE):
        warnings.append("fund_name ETF 字样疑似被拆分")
    if asset_class and asset_class not in standard_asset_classes:
        warnings.append(f"asset_class 非标准值: {asset_class}")
    if re.search(r"\b(?:Ut\.S\.|eU\.S\.|eUt\.S\.|U\.[cftoe]S\.|U[cftoe]\.S\.)\b", geography):
        warnings.append(f"geography 疑似污染: {geography}")

    vocab_fields = ["category", "strategy", "selection_criteria", "weighting_scheme"]
    for field in vocab_fields:
        value = str(row.get(field, "") or "")
        if re.search(r"\b(?:Equcity|Fixmed|Incoome|Proprie\s+\w+|Singcle|Vaniclla)\b", value):
            warnings.append(f"{field} 出现明显拼写污染: {value}")
        if re.search(r"\b[fetco]\s*\.\s*", value):
            warnings.append(f"{field} 疑似残留水印片段")
    return "；".join(dict.fromkeys(warnings))


def text_pool(row: Dict[str, Any], fields: Iterable[str]) -> str:
    return " | ".join(str(row.get(f, "") or "") for f in fields)


def find_keywords(row: Dict[str, Any], fields: Iterable[str], keywords: List[str]) -> List[Tuple[str, str]]:
    hits: List[Tuple[str, str]] = []
    for field in fields:
        value = str(row.get(field, "") or "")
        low = value.lower()
        for kw in keywords:
            if kw.lower() in low:
                hits.append((field, kw))
    return hits


def classify(row: Dict[str, Any]) -> None:
    row["preliminary_strategy_level_1"] = ""
    row["preliminary_strategy_level_2"] = ""
    row["classification_reason"] = ""
    row["is_excluded_by_features"] = False
    row["is_strategy_candidate"] = False
    row["is_extended_strategy_candidate"] = False

    segments_text = str(row.get("segments", "") or "")
    fund_name = str(row.get("fund_name", "") or "")
    leveraged_ratio = str(row.get("leveraged_ratio", "") or "").strip()
    exclusion_hits: List[str] = []
    if row.get("leveraged") is True:
        exclusion_hits.append("leveraged=True")
    if row.get("inverse") is True:
        exclusion_hits.append("inverse=True")
    if row.get("etn") is True:
        exclusion_hits.append("etn=True")
    if leveraged_ratio and leveraged_ratio != "-":
        exclusion_hits.append(f"leveraged_ratio={leveraged_ratio}")
    if re.search(r"\b(Leveraged|Inverse)\b", segments_text, re.IGNORECASE):
        exclusion_hits.append("segments contains Leveraged/Inverse")
    if re.search(r"\b(2X|3X|Ultra|Bear|Bull 2X|Daily Target|Daily ETF)\b", fund_name, re.IGNORECASE):
        exclusion_hits.append("fund_name contains leveraged/inverse keyword")
    if exclusion_hits:
        row["preliminary_pool_type"] = "排除_杠杆反向ETN"
        row["is_excluded_by_features"] = True
        row["classification_reason"] = "; ".join(exclusion_hits)
        return

    extended_keywords = [
        "Structured Outcome",
        "Buffer",
        "Defined Protection",
        "Defined Outcome",
        "Autocallable",
        "Barrier",
        "Income Strategy",
        "Yield Premium",
        "Premium Income",
        "Covered Call",
        "BuyWrite",
        "Collared",
        "Laddered Buffer",
        "Downside Hedge",
        "Target Outcome",
    ]
    extended_fields = ["category", "strategy", "niche", "segments", "fund_name"]
    extended_hits = find_keywords(row, extended_fields, extended_keywords)
    if extended_hits:
        row["preliminary_pool_type"] = "扩展策略ETF"
        row["is_strategy_candidate"] = True
        row["is_extended_strategy_candidate"] = True
        hit_text = "; ".join(f"{field}:{kw}" for field, kw in extended_hits[:6])
        combined = " ".join(kw for _, kw in extended_hits).lower()
        if any(k in combined for k in ["buffer", "defined outcome", "defined protection", "structured outcome", "target outcome", "laddered buffer"]):
            lvl1 = "Buffer/Defined Outcome"
        elif any(k in combined for k in ["autocallable", "barrier"]):
            lvl1 = "Autocallable/Barrier Income"
        elif any(k in combined for k in ["yield premium", "premium income", "covered call", "buywrite", "income strategy"]):
            lvl1 = "Option Income/Premium Income"
        elif any(k in combined for k in ["downside hedge", "protection"]):
            lvl1 = "Downside Hedge/Protection"
        else:
            lvl1 = "Extended Strategy"
        row["preliminary_strategy_level_1"] = lvl1
        row["preliminary_strategy_level_2"] = lvl1
        row["classification_reason"] = f"命中扩展策略关键词：{hit_text}"
        return

    core_priority = [
        ("Dividend", ["High Dividend Yield", "Dividend"]),
        ("Free Cash Flow", ["Free Cash Flow", "FCF", "Cash Cow"]),
        ("Low Volatility", ["Low Vol", "Min Vol", "Minimum Volatility"]),
        ("Quality", ["Quality"]),
        ("Value", ["Value"]),
        ("Growth", ["Growth"]),
        ("Momentum", ["Momentum"]),
        ("Factor / Multi-Factor", ["Multi-Factor", "Active Factor", "Alpha Capture", "Factor"]),
        ("Fundamental Weighting", ["Fundamental"]),
        ("Equal Weight", ["Equal"]),
        ("Shareholder Yield", ["Shareholder Yield", "Buyback"]),
        ("Moat", ["Moat"]),
        ("Yield Enhanced", ["Yield Enhanced"]),
        ("Sector Rotation", ["Sector Rotation"]),
    ]
    core_fields = [
        "fund_name",
        "category",
        "strategy",
        "niche",
        "segments",
        "selection_criteria",
        "weighting_scheme",
    ]
    core_hits_all: List[Tuple[str, str]] = []
    selected_lvl1 = ""
    for lvl1, keywords in core_priority:
        hits = find_keywords(row, core_fields, keywords)
        if hits and not selected_lvl1:
            selected_lvl1 = lvl1
            core_hits_all.extend(hits)
            break
    # Keep additional evidence for the same row.
    for _, keywords in core_priority:
        for hit in find_keywords(row, core_fields, keywords):
            if hit not in core_hits_all:
                core_hits_all.append(hit)
    if selected_lvl1:
        row["preliminary_pool_type"] = "核心策略ETF_候选"
        row["is_strategy_candidate"] = True
        row["is_extended_strategy_candidate"] = False
        row["preliminary_strategy_level_1"] = selected_lvl1
        row["preliminary_strategy_level_2"] = selected_lvl1
        row["classification_reason"] = "命中核心策略关键词：" + "; ".join(
            f"{field}:{kw}" for field, kw in core_hits_all[:8]
        )
        return

    strategy = str(row.get("strategy", "") or "").strip().lower()
    selection = str(row.get("selection_criteria", "") or "").strip().lower()
    weighting = str(row.get("weighting_scheme", "") or "").strip().lower()
    if strategy == "active" or selection == "proprietary" or weighting == "proprietary":
        row["preliminary_pool_type"] = "观察池_主动或专有策略待核验"
        row["classification_reason"] = "Active 或 Proprietary 口径但未命中核心/扩展策略关键词"
        return

    asset_class = str(row.get("asset_class", "") or "").strip().lower()
    category = str(row.get("category", "") or "").strip().lower()
    if (
        (category == "sector" and strategy == "vanilla")
        or asset_class == "commodities"
        or (asset_class == "fixed income" and any(k in category for k in ["broad market", "corporate", "bank loans"]))
        or (selection == "market cap" and weighting == "market cap" and strategy == "vanilla")
    ):
        row["preliminary_pool_type"] = "排除_普通宽基主题商品债券"
        row["classification_reason"] = "未命中策略关键词，符合普通宽基/主题/商品/债券排除规则"
        return

    row["preliminary_pool_type"] = "待人工复核"
    row["classification_reason"] = "未命中排除、核心策略、扩展策略或主动/专有观察规则"


def merge_modules(module_data: Dict[str, Dict[str, Dict[str, Any]]], extract_time: str) -> List[Dict[str, Any]]:
    all_tickers = sorted(set().union(*(set(data.keys()) for data in module_data.values())))
    master: List[Dict[str, Any]] = []
    for ticker in all_tickers:
        row: Dict[str, Any] = {field: "" for field in MASTER_FIELDS}
        row["ticker"] = ticker
        source_pages: List[str] = []
        parse_warnings: List[str] = []

        class_rec = module_data.get("classification", {}).get(ticker, {})
        if class_rec:
            for field in MODULES["classification"]["fields"][1:]:
                row[field] = class_rec.get(field, "")
            source_pages.append(f"classification:p{class_rec.get('_source_page')}")
        else:
            parse_warnings.append("missing classification module")

        # Product classification module fund_name is authoritative.
        authoritative_fund_name = row.get("fund_name", "")

        for module_name in ["liquidity", "flows_returns", "geo"]:
            rec = module_data.get(module_name, {}).get(ticker)
            if not rec:
                parse_warnings.append(f"missing {module_name} module")
                continue
            source_pages.append(f"{module_name}:p{rec.get('_source_page')}")
            for field, value in rec.items():
                if field == "fund_name" or field.startswith("_") or field == "ticker":
                    continue
                row[field] = value
            if not authoritative_fund_name and rec.get("fund_name"):
                row["fund_name"] = rec.get("fund_name", "")

        normalize_raw_fields(row)
        add_cleaned_values(row)
        classify(row)
        row["text_quality_warning"] = detect_text_quality_warning(row)
        row["source_pdf"] = str(INPUT_PDF)
        row["source_pages"] = "; ".join(source_pages)
        row["extract_time"] = extract_time
        if row["text_quality_warning"]:
            parse_warnings.append(row["text_quality_warning"])
        row["parse_warning"] = "; ".join(parse_warnings)
        master.append(row)
    return master


def write_csv(rows: List[Dict[str, Any]]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    with OUTPUT_CSV.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=MASTER_FIELDS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k) for k in MASTER_FIELDS})


def excel_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    return value


def write_xlsx(rows: List[Dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "master"
    ws.append(MASTER_FIELDS)
    for row in rows:
        ws.append([excel_value(row.get(field)) for field in MASTER_FIELDS])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    text_fields = {"ticker", "source_pdf", "source_pages", "extract_time", "parse_warning"}
    numeric_fields = {
        "aum_usd",
        "avg_daily_dollar_vol_usd",
        "avg_daily_share_vol",
        "one_year_fund_flow_usd",
        "ytd_fund_flow_usd",
        "one_month_fund_flow_usd",
        "expense_ratio_pct",
        "spread_data_pct",
        "dividend_yield_pct",
        "one_year_total_return_pct",
        "three_year_total_return_pct",
        "five_year_total_return_pct",
        "pe",
        "pb",
    }
    for col_idx, field in enumerate(MASTER_FIELDS, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(field)
        for cell in ws[col_letter][1:]:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 80))
            if field in numeric_fields:
                cell.number_format = "#,##0.00"
            elif field in text_fields:
                cell.number_format = "@"
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 45)
    wb.save(OUTPUT_XLSX)


def build_log(
    pdf_pages: int,
    module_data: Dict[str, Dict[str, Dict[str, Any]]],
    duplicate_tickers: List[str],
    module_warnings: List[str],
    rows: List[Dict[str, Any]],
) -> str:
    lines: List[str] = []
    lines.append("ETF.com PDF extraction log v1")
    lines.append(f"extract_time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    lines.append(f"source_pdf: {INPUT_PDF}")
    lines.append(f"output_xlsx: {OUTPUT_XLSX}")
    lines.append(f"output_csv: {OUTPUT_CSV}")
    lines.append("")
    lines.append(f"1. PDF 总页数: {pdf_pages}")
    module_labels = {
        "classification": "2. 产品分类模块提取 ticker 数",
        "liquidity": "3. 规模费率流动性模块提取 ticker 数",
        "flows_returns": "4. 资金流收益特征模块提取 ticker 数",
        "geo": "5. 杠杆反向地域模块提取 ticker 数",
    }
    for key, label in module_labels.items():
        lines.append(f"{label}: {len(module_data.get(key, {}))}")
    lines.append(f"6. 合并后唯一 ticker 数: {len(rows)}")
    lines.append(f"7. 重复 ticker 数: {len(set(duplicate_tickers))}")
    if duplicate_tickers:
        lines.append(f"   重复 ticker 明细: {', '.join(sorted(set(duplicate_tickers)))}")
    union_tickers = set(row["ticker"] for row in rows)
    lines.append("8. 各模块未匹配 ticker 数:")
    for key in MODULES:
        missing = sorted(union_tickers - set(module_data.get(key, {}).keys()))
        lines.append(f"   - {key}: {len(missing)}" + (f" ({', '.join(missing[:20])})" if missing else ""))
    lines.append("9. preliminary_pool_type 的数量分布:")
    dist = Counter(row.get("preliminary_pool_type") for row in rows)
    for k, v in dist.most_common():
        lines.append(f"   - {k}: {v}")
    lines.append("10. 前 10 行 master 数据预览:")
    preview_fields = ["ticker", "fund_name", "issuer", "preliminary_pool_type", "preliminary_strategy_level_1"]
    for row in rows[:10]:
        lines.append("   " + " | ".join(str(row.get(f, "")) for f in preview_fields))
    warning_count = sum(bool(row.get("parse_warning")) for row in rows)
    lines.append(f"11. parse_warning 非空的记录数量: {warning_count}")
    text_warning_count = sum(bool(row.get("text_quality_warning")) for row in rows)
    lines.append(f"12. text_quality_warning 非空的记录数量: {text_warning_count}")
    if module_warnings:
        lines.append("")
        lines.append("module_warnings:")
        for w in module_warnings:
            lines.append(f"   - {w}")
    lines.append("")
    lines.append("重点样本验收:")
    sample_tickers = ["ACWV", "ABFL", "AAPR", "ABUF", "AALG", "AAPU", "AGQ", "AGIQ", "AGIX"]
    by_ticker = {row["ticker"]: row for row in rows}
    for ticker in sample_tickers:
        row = by_ticker.get(ticker)
        if not row:
            lines.append(f"   - {ticker}: 未提取")
            continue
        lines.append(
            f"   - {ticker}: {row.get('preliminary_pool_type')} | "
            f"{row.get('preliminary_strategy_level_1')} | "
            f"{row.get('classification_reason')}"
        )
    lines.append("")
    lines.append("重点样本修复前后对比:")
    previous_rows = read_previous_rows(COMPARE_PREVIOUS_CSV)
    for ticker in ["AAA", "AAAC", "AAAD", "ACWV", "ABUF", "AGIX"]:
        old = previous_rows.get(ticker, {})
        new = by_ticker.get(ticker, {})
        if not new:
            lines.append(f"   - {ticker}: v2 未提取")
            continue
        lines.append(
            f"   - {ticker}: "
            f"fund_name [{old.get('fund_name', '')} -> {new.get('fund_name', '')}]；"
            f"issuer [{old.get('issuer', '')} -> {new.get('issuer', '')}]；"
            f"asset_class [{old.get('asset_class', '')} -> {new.get('asset_class', '')}]；"
            f"geography [{old.get('geography', '')} -> {new.get('geography', '')}]；"
            f"text_quality_warning={new.get('text_quality_warning', '')}"
        )
    return "\n".join(lines)


def read_previous_rows(path: Path) -> Dict[str, Dict[str, str]]:
    if not path.exists():
        return {}
    with path.open("r", newline="", encoding="utf-8-sig") as f:
        return {row.get("ticker", ""): row for row in csv.DictReader(f) if row.get("ticker")}


def main() -> None:
    if not INPUT_PDF.exists():
        raise FileNotFoundError(f"Input PDF not found: {INPUT_PDF}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    extract_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    module_data: Dict[str, Dict[str, Dict[str, Any]]] = {}
    duplicate_tickers: List[str] = []
    module_warnings: List[str] = []
    with pdfplumber.open(str(INPUT_PDF)) as pdf:
        pdf_pages = len(pdf.pages)
        for module_name in MODULES:
            data, duplicates, warnings = extract_module(pdf, module_name)
            module_data[module_name] = data
            duplicate_tickers.extend(duplicates)
            module_warnings.extend(warnings)

    rows = merge_modules(module_data, extract_time)
    write_csv(rows)
    write_xlsx(rows)
    log_text = build_log(pdf_pages, module_data, duplicate_tickers, module_warnings, rows)
    OUTPUT_LOG.write_text(log_text, encoding="utf-8-sig")
    print(log_text)


if __name__ == "__main__":
    main()
