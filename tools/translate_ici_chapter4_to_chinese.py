from __future__ import annotations

import re
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


INPUT_DIR = Path(r"C:\Users\lvdon\Desktop\Fullgoal\ICI课题\26-fb-data-ch4")
OUTPUT_DIR = INPUT_DIR / "中文数据集"
FILES = [f"4.{i}.xlsx" for i in range(1, 8)]
TEXT_FILE = "data_methods_and_assumptions.txt"

ALLOWED_WORDS = {
    "ETF", "ETFs", "ICI", "IRA", "IRAs", "DC", "Cboe", "Bloomberg", "Refinitiv", "USD", "US",
}

TRANSLATIONS: dict[str, str] = {
    "Figure 4.1": "图4.1",
    "Figure 4.2": "图4.2",
    "Figure 4.3": "图4.3",
    "Figure 4.4": "图4.4",
    "Figure 4.5": "图4.5",
    "Figure 4.6": "图4.6",
    "Figure 4.7": "图4.7",
    "Year": "年份",
    "Annual": "年度",
    "Percent": "百分比",
    "Billions of dollars, year-end": "十亿美元，年末",
    "Billions of dollars, annual": "十亿美元，年度",
    "Trillions of dollars": "万亿美元",
    "Sources": "资料来源",
    "Source": "资料来源",
    "Note": "注",
    "All ETFs": "全部ETF",
    "Number of ETFs": "ETF数量",
    "Total Net Assets of ETFs Surpassed $13 Trillion in 2025": "ETF总净资产在2025年突破13万亿美元",
    "Large-cap domestic equity": "美国大盘股票",
    "Other domestic equity": "其他美国股票",
    "Global/International equity": "全球/国际股票",
    "Bond": "债券",
    "Hybrid and commodity1": "混合及商品¹",
    "Commodity ETFs include funds—both registered and not registered under the Investment Company Act of 1940—that invest primarily in commodities, currencies, and futures.": "商品ETF包括根据《1940年投资公司法》注册及未注册的基金，这些基金主要投资于商品、货币和期货。",
    "Note: The first bond, hybrid, and commodity ETFs were opened in 2002, 2007, and 2004, respectively.": "注：首只债券ETF、混合ETF和商品ETF分别设立于2002年、2007年和2004年。",
    "ETF Secondary Market Trading Averaged 28% of Daily US Stock Trading in 2025": "2025年ETF二级市场交易平均占美国每日股票交易量的28%",
    "Percentage of total US stock market trading volume, annual": "占美国股票市场总交易量的比例，年度",
    "Average daily ETF share of total US stock market trading": "ETF平均每日交易量占美国股票市场总交易量的比例",
    "Minimum daily ETF share for the year": "年内ETF日交易占比最低值",
    "Maximum daily ETF share for the year": "年内ETF日交易占比最高值",
    "Date of maximum": "最高值日期",
    "Date of minimum": "最低值日期",
    "Domestic Equity ETFs Have Had Minimal Impact on Underlying US Stocks": "美国国内股票ETF对底层美国股票影响有限",
    "Domestic equity ETF primary market activity": "美国国内股票ETF一级市场活动",
    "Value of company stock traded": "公司股票成交额",
    "Domestic equity ETF primary market activity as a share of company stock traded": "美国国内股票ETF一级市场活动占公司股票成交额比例",
    "Primary market activity is measured as the total of gross issuance and gross redemptions.": "一级市场活动按总申购和总赎回之和衡量。",
    "Net Share Issuance of ETFs Surged to a Record $1.5 Trillion in 2025": "ETF净份额发行额在2025年跃升至创纪录的1.5万亿美元",
    "Data for net share issuance include reinvested dividends.": "净份额发行额数据包含再投资股息。",
    "Number of ETFs Entering and Exiting the Industry": "ETF进入和退出行业的数量",
    "Opened": "新设立",
    "Liquidated/Merged": "清盘/合并",
    "Data include ETFs that invest primarily in other ETFs.": "数据包括主要投资于其他ETF的ETF。",
    "ETF–Owning Households Are From All Demographic Groups": "ETF持有家庭覆盖各类人口群体",
    "Percentage of ETF–owning households, 2025": "ETF持有家庭占比，2025年",
    "Who are they?": "他们是谁？",
    "What do they own?": "他们持有什么？",
    "Where do they hold ETFs?": "他们在哪里持有ETF？",
    "Why do they invest in ETFs?": "他们为什么投资ETF？",
    "Aged 35 to 64": "35至64岁",
    "College graduates": "大学毕业生",
    "Employed (full- or part-time)": "就业中，包括全职或兼职",
    "Household income under $150,000": "家庭收入低于15万美元",
    "Retired": "已退休",
    "Baby Boomers": "婴儿潮一代",
    "Generation X": "X世代",
    "Millennials": "千禧一代",
    "Own individual retirement accounts (IRAs)": "持有IRA个人退休账户",
    "Own defined contribution (DC) retirement plan accounts": "持有DC缴费确定型退休计划账户",
    "Own equity ETFs": "持有股票ETF",
    "Own mutual funds": "持有共同基金",
    "Own ETFs inside IRAs": "在IRA个人退休账户中持有ETF",
    "Own ETFs inside taxable brokerage accounts": "在应税证券经纪账户中持有ETF",
    "Save for retirement": "为退休储蓄",
    "Cost effectiveness of ETFs1": "ETF具有成本效益¹",
    "Diversification of investments through ETFs1": "通过ETF实现投资分散化¹",
    "Ability to sell ETF shares any time during the day1": "能够在日内任意时间卖出ETF份额¹",
    "Characteristics of ETF-Owning Households": "ETF持有家庭特征",
    "All US households": "全美家庭",
    "Households owning ETFs": "ETF持有家庭",
    "Households owning mutual funds": "共同基金持有家庭",
    "Households owning individual stocks": "个股持有家庭",
    "Median": "中位数",
    "Age of household survey respondent": "家庭调查受访者年龄",
    "Household income1": "家庭收入¹",
    "Household financial assets2": "家庭金融资产²",
    "Percentage of households": "家庭占比",
    "Married or living with a partner": "已婚或与伴侣同居",
    "College or postgraduate degree": "大学或研究生学历",
    "Retired from lifetime occupation": "从长期职业中退休",
    "IRA(s)": "IRA个人退休账户",
    "DC retirement plan account(s)": "DC缴费确定型退休计划账户",
    "Household’s willingness to take financial risk": "家庭承担金融风险的意愿",
    "Substantial risk for substantial gain": "为获得可观收益愿意承担重大风险",
    "Above-average risk for above-average gain": "为获得高于平均的收益愿意承担高于平均的风险",
    "Average risk for average gain": "为获得平均收益愿意承担平均风险",
    "Below-average risk for below-average gain": "为获得低于平均的收益愿意承担低于平均的风险",
    "Unwilling to take any risk": "不愿承担任何风险",
    "Domestic equity ETF primary market activity1\nTrillions of dollars": "美国国内股票ETF一级市场活动¹\n万亿美元",
    "Value of company stock traded\nTrillions of dollars": "公司股票成交额\n万亿美元",
    "Domestic equity ETF primary market activity as a share of company stock traded\nPercent": "美国国内股票ETF一级市场活动占公司股票成交额比例\n百分比",
    "Note: Data for net share issuance include reinvested dividends.": "注：净份额发行额数据包含再投资股息。",
    "Note: Data include ETFs that invest primarily in other ETFs.": "注：数据包括主要投资于其他ETF的ETF。",
    "Sources: Investment Company Institute, Bloomberg, Refinitiv, and Cboe Exchange, Inc.": "资料来源：美国投资公司协会、Bloomberg、Refinitiv 和 Cboe Exchange, Inc.",
    "Source: ICI Research Data Release, “Profile of ETF-Owning Households, 2025”": "资料来源：ICI Research Data Release，《Profile of ETF-Owning Households, 2025》",
    "Percentage of ETF-owning households indicating the feature was “very important”, “somewhat important”, or “not very important” in their ETF purchase decision.": "ETF持有家庭中表示该特征在其ETF购买决策中“非常重要”“有些重要”或“不太重要”的比例。",
    "Total reported is household income before taxes in 2024.": "报告总额为2024年税前家庭收入。",
    "Household financial assets include assets in employer-sponsored retirement plans but exclude the household’s primary residence.": "家庭金融资产包括雇主发起退休计划中的资产，但不包括家庭主要居所。",
    "Household owns": "家庭持有",
    "Household survey respondent": "家庭调查受访者",
    "Households": "家庭",
    "households": "家庭",
    "owning": "持有",
    "ETFs": "ETF",
    "funds": "基金",
    "mutual": "共同",
    "individual": "个股",
    "stocks": "股票",
    "All US": "全美",
}

MONTHS = {
    "Jan": "1月", "Feb": "2月", "Mar": "3月", "Apr": "4月", "May": "5月", "Jun": "6月",
    "Jul": "7月", "Aug": "8月", "Sep": "9月", "Oct": "10月", "Nov": "11月", "Dec": "12月",
}

METHODS_ZH = """方法与假设

除非另有说明，本书所有数据均适用以下方法：

» 美国注册投资公司数据仅包括向美国投资公司协会（ICI）报送统计信息的公司。这些公司的资产至少覆盖行业资产的98%。

» 为避免重复计算，基金中基金被排除在数据之外。

» 由于四舍五入，美元金额和百分比加总后可能与所列总数不完全一致。

» 美国注册投资公司数据包括未根据《1940年投资公司法》注册的交易型开放式指数基金（ETF）。

» 长期基金包括股票基金、混合基金和债券基金。

数据可能会被修订。虽然由独立来源提供的信息和数据被认为可靠，但美国投资公司协会不对其准确性、完整性或及时性负责。独立来源表达的观点不一定代表协会观点。如对相关材料有疑问或意见，请直接联系数据来源方。
"""


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def translate_months(s: str) -> str:
    def repl(match: re.Match[str]) -> str:
        return f"{MONTHS[match.group(1)]}{int(match.group(2))}日"

    return re.sub(r"\b(" + "|".join(MONTHS) + r")\s+(\d{1,2})\b", repl, s)


def translate_string(s: str) -> tuple[str, str]:
    original = s
    stripped = s.strip()
    if stripped in TRANSLATIONS:
        return s.replace(stripped, TRANSLATIONS[stripped]), "词典"

    month_translated = translate_months(s)
    if month_translated != s:
        return month_translated, "月份日期"

    # Footnotes often appear as 1Text or 2Text in the raw workbook.
    m = re.match(r"^([12])(.+)$", stripped, flags=re.S)
    if m and m.group(2).strip() in TRANSLATIONS:
        marker = "¹" if m.group(1) == "1" else "²"
        return marker + TRANSLATIONS[m.group(2).strip()], "脚注词典"

    if "\n" in s:
        translated_lines = []
        changed = False
        for line in s.split("\n"):
            line_translated, mode = translate_string(line)
            translated_lines.append(line_translated)
            changed = changed or mode != "未覆盖"
        if changed:
            return "\n".join(translated_lines), "多行词典"

    return original, "未覆盖"


def has_untranslated_english(s: str) -> bool:
    cleaned = s
    for word in ALLOWED_WORDS:
        cleaned = re.sub(rf"\b{re.escape(word)}\b", "", cleaned)
    return bool(re.search(r"[A-Za-z]{4,}", cleaned))


def apply_yahei(workbook) -> None:
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    old = copy(cell.font)
                    cell.font = Font(
                        name="微软雅黑",
                        size=old.sz,
                        bold=old.bold,
                        italic=old.italic,
                        vertAlign=old.vertAlign,
                        underline=old.underline,
                        strike=old.strike,
                        color=old.color,
                    )
        for row_num in range(1, min(ws.max_row, 5) + 1):
            for cell in ws[row_num]:
                old = copy(cell.font)
                cell.font = Font(
                    name="微软雅黑",
                    size=old.sz,
                    bold=True,
                    italic=old.italic,
                    vertAlign=old.vertAlign,
                    underline=old.underline,
                    strike=old.strike,
                    color=old.color,
                )
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            current = ws.column_dimensions[letter].width
            if current is None:
                max_len = 0
                for cell in ws[letter][:50]:
                    max_len = max(max_len, len(text(cell.value)))
                ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 35)


def add_translation_note(workbook) -> None:
    if "翻译说明" in workbook.sheetnames:
        del workbook["翻译说明"]
    ws = workbook.create_sheet("翻译说明")
    rows = [
        ("项目", "内容"),
        ("数据来源", "ICI 2026 Investment Company Fact Book"),
        ("处理说明", "本文件由英文原始数据翻译为中文，数值未作修改"),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(name="微软雅黑", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 70
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="微软雅黑")


def add_figure_47_tidy(workbook) -> None:
    if "整理版_中文" in workbook.sheetnames:
        del workbook["整理版_中文"]
    source = workbook["图4.7"]
    ws = workbook.create_sheet("整理版_中文")
    ws.append(["指标", "全美家庭", "ETF持有家庭", "共同基金持有家庭", "个股持有家庭"])
    for r in range(7, 28):
        metric = source.cell(r, 1).value
        if metric is None or str(metric).strip() == "":
            continue
        values = [source.cell(r, c).value for c in [3, 5, 7, 9]]
        ws.append([str(metric).strip()] + values)
    for cell in ws[1]:
        cell.font = Font(name="微软雅黑", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="微软雅黑")
    for col, width in {"A": 40, "B": 16, "C": 18, "D": 20, "E": 18}.items():
        ws.column_dimensions[col].width = width


def create_translation_table(path: Path, translation_usage: dict[str, str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "词条翻译表"
    ws.append(["英文词条", "中文翻译"])
    for key in sorted(set(TRANSLATIONS) | set(translation_usage)):
        ws.append([key, TRANSLATIONS.get(key, translation_usage.get(key, ""))])
    for cell in ws[1]:
        cell.font = Font(name="微软雅黑", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 90
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="微软雅黑")
    wb.save(path)


def create_untranslated_table(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "未翻译词条"
    ws.append(["文件名", "sheet名", "单元格", "原始文本"])
    for item in rows:
        ws.append([item["文件名"], item["sheet名"], item["单元格"], item["原始文本"]])
    for cell in ws[1]:
        cell.font = Font(name="微软雅黑", bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 100
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="微软雅黑")
    wb.save(path)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    untranslated: list[dict] = []
    translation_usage: dict[str, str] = {}
    output_paths: list[Path] = []

    for filename in FILES:
        input_path = INPUT_DIR / filename
        output_path = OUTPUT_DIR / f"{input_path.stem}_中文.xlsx"
        wb = load_workbook(input_path)
        original_sheet_names = list(wb.sheetnames)
        for ws in wb.worksheets:
            old_title = ws.title
            if re.fullmatch(r"4\.[1-7]", old_title):
                ws.title = f"图{old_title}"
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        continue
                    if isinstance(cell.value, str) and cell.value.strip():
                        old_value = cell.value
                        new_value, mode = translate_string(old_value)
                        if new_value != old_value:
                            cell.value = new_value
                            translation_usage[old_value.strip()] = new_value.strip()
                        elif mode == "未覆盖" and has_untranslated_english(old_value):
                            untranslated.append({
                                "文件名": filename,
                                "sheet名": old_title,
                                "单元格": cell.coordinate,
                                "原始文本": old_value,
                            })
        apply_yahei(wb)
        if filename == "4.7.xlsx":
            add_figure_47_tidy(wb)
        add_translation_note(wb)
        wb.save(output_path)
        output_paths.append(output_path)

    methods_output = OUTPUT_DIR / "data_methods_and_assumptions_中文.txt"
    methods_output.write_text(METHODS_ZH, encoding="utf-8")
    output_paths.append(methods_output)

    translation_table = OUTPUT_DIR / "ICI_Chapter4_词条翻译表.xlsx"
    create_translation_table(translation_table, translation_usage)
    output_paths.append(translation_table)

    untranslated_table = OUTPUT_DIR / "ICI_Chapter4_未翻译词条_待人工核对.xlsx"
    create_untranslated_table(untranslated_table, untranslated)
    output_paths.append(untranslated_table)

    readme = OUTPUT_DIR / "README_中文数据集说明.txt"
    readme_text = "\n".join([
        "ICI Chapter 4 中文数据集说明",
        "",
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"输入目录：{INPUT_DIR}",
        f"输出目录：{OUTPUT_DIR}",
        "输入文件：4.1.xlsx、4.2.xlsx、4.3.xlsx、4.4.xlsx、4.5.xlsx、4.6.xlsx、4.7.xlsx、data_methods_and_assumptions.txt",
        "翻译范围：sheet名、标题、字段名、分类名、脚注、资料来源、月份日期、文本单元格。",
        "数值是否被修改：未修改。",
        f"未翻译词条数量：{len(untranslated)}",
        "说明：ETF、ICI、IRA、DC、Cboe、Bloomberg、Refinitiv、USD 等专有缩写按要求保留。",
    ])
    readme.write_text(readme_text, encoding="utf-8")
    output_paths.append(readme)

    # Quality check: reopen outputs and scan remaining long English.
    remaining_long_english = []
    for path in output_paths:
        if path.suffix.lower() != ".xlsx" or "未翻译词条" in path.name:
            continue
        wb = load_workbook(path, data_only=False)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and has_untranslated_english(cell.value):
                        remaining_long_english.append((path.name, ws.title, cell.coordinate, cell.value))

    print("输出文件路径：")
    for path in output_paths:
        print(f"  {path}")
    print(f"未翻译词条数量：{len(untranslated)}")
    print(f"中文Excel中仍存在未翻译英文长句数量：{len(remaining_long_english)}")
    if remaining_long_english:
        print("未翻译英文长句样例：")
        for item in remaining_long_english[:20]:
            print(f"  {item[0]} | {item[1]} | {item[2]} | {item[3]}")
    print("处理完成：ICI Chapter 4 中文数据集已生成")


if __name__ == "__main__":
    main()
