from pathlib import Path

from openpyxl import load_workbook


paths = [
    Path.home() / "Desktop" / "Fullgoal" / "课题研究" / "处理后数据" / "境外策略ETF数据" / "处理结果" / "境外ETF全量数据_清洗标准化_v1.xlsx",
    Path.home() / "Desktop" / "Fullgoal" / "课题研究" / "原始数据" / "全市场ETF基础信息.csv.xlsx",
    Path.home() / "Desktop" / "Fullgoal" / "ICI课题" / "课题更正" / "data_extraction_top10_amc_etf_full_data.xlsx",
]

for path in paths:
    print("\nPATH", path)
    print("exists", path.exists())
    if not path.exists():
        continue
    wb = load_workbook(path, read_only=True, data_only=True)
    print("sheets", wb.sheetnames)
    for ws in wb.worksheets[:3]:
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        print("sheet", ws.title, "rows", ws.max_row, "cols", ws.max_column)
        print("headers", headers[:80])
        for vals in ws.iter_rows(min_row=2, max_row=4, values_only=True):
            print("sample", vals[: min(20, len(vals))])
