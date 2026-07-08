from __future__ import annotations

import importlib.util
import math
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE_DIR = Path.home() / "Desktop" / "Fullgoal"
WORK_DIR = BASE_DIR / "课题研究" / "FullGoal-QUANT"
ICI_DIR = BASE_DIR / "ICI课题" / "课题更正"
US_FULL = BASE_DIR / "课题研究" / "处理后数据" / "境外策略ETF数据" / "处理结果" / "境外ETF全量数据_清洗标准化_v1.xlsx"
US_TOP10_FALLBACK = ICI_DIR / "美国ETF_TOP10生态评分数据包_20260706.xlsx"
CN_FULL = BASE_DIR / "课题研究" / "原始数据" / "全市场ETF基础信息.csv.xlsx"
CN_SUPPLEMENT = ICI_DIR / "data_extraction_top10_amc_etf_full_data.xlsx"
CN_REFERENCE = ICI_DIR / "中国ETF_TOP10生态评分数据包_补充交易跟踪版.xlsx"
OUTPUT = ICI_DIR / "中美ETF_TOP10五维生态评分_规模筛选版.xlsx"

US_BUILDER = WORK_DIR / "03_scripts" / "build_us_top10_etf_ecosystem_score.py"
CN_BUILDER = WORK_DIR / "03_scripts" / "build_cn_top10_etf_ecosystem_score.py"

SCORE_HEADERS = [
    "规模排名",
    "生态排名",
    "管理人",
    "ETF管理规模",
    "市场份额",
    "ETF产品数量",
    "产品广度",
    "费率优势",
    "交易活跃",
    "多元能力",
    "旗舰产品",
    "生态评分",
    "评分口径说明",
]

DIMENSIONS = ["产品广度", "费率优势", "交易活跃", "多元能力", "旗舰产品"]
US_NOTE = "美国样本先按ETF管理规模筛选TOP10，再在TOP10内部进行五维等权生态评分；分数用于美国市场内部比较及中美路径映射，不代表中美绝对实力直接比较。"
CN_NOTE = "中国样本先按2026年6月底ETF管理规模筛选TOP10，再在TOP10内部进行五维等权生态评分；交易活跃使用2026-07-07补充交易数据修正；分数用于中国市场内部比较及中美路径映射，不代表中美绝对实力直接比较。"

RULE_ROWS = [
    ["产品广度", "20%", "产品线是否铺得全", "ETF产品数量、指数/主题覆盖数量", "越多越好", "看产品货架宽度，不评价规模"],
    ["费率优势", "20%", "产品成本是否有竞争力", "AUM加权综合费率", "费率越低越好", "ETF竞争的基础能力，简单直接"],
    ["交易活跃", "20%", "产品是否被市场真实交易使用", "成交效率、活跃产品占比", "成交效率越高、活跃产品占比越高越好", "规模已用于筛选，因此用成交效率而非绝对成交额"],
    ["多元能力", "20%", "是否能把不同投资能力转化为ETF产品", "多元产品规模占比、能力模块覆盖数", "占比越高、覆盖越多越好", "衡量策略、跨境、债券、商品、主题等能力ETF化"],
    ["旗舰产品", "20%", "是否有客户容易识别的代表产品", "最大ETF规模占比、Top3产品规模占比", "占比越高越好", "看是否形成代表性旗舰产品，而非单纯绝对规模"],
]

RULE_NOTE = "本评分为自建研究框架，用于ETF管理人生态位比较，不代表官方排名。样本先按ETF管理规模筛选各自市场TOP10管理人，规模只作为入选标准，不进入生态评分。五个维度采用等权，代表指标在各自市场TOP10内部标准化至60—100分。美国和中国使用同一套五维框架，但分数在各自市场内部标准化，因此中美之间主要比较结构、生态位和学习路径，不直接比较绝对实力。"


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def norm_code(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip().upper())


def code_core(value: Any) -> str:
    return norm_code(value).split(".")[0]


def to_float(value: Any) -> float | None:
    if value in (None, "", "-"):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("%", "")
    try:
        return float(text)
    except ValueError:
        return None


def sum_field(rows: list[dict[str, Any]], field: str) -> float:
    return sum((r.get(field) or 0) for r in rows if r.get(field) is not None)


def weighted_avg(rows: list[dict[str, Any]], value_field: str, weight_field: str) -> float | None:
    valid = [(r.get(value_field), r.get(weight_field)) for r in rows if r.get(value_field) is not None and (r.get(weight_field) or 0) > 0]
    total_weight = sum(w for _, w in valid)
    if total_weight <= 0:
        return None
    return sum(v * w for v, w in valid) / total_weight


def standardize(values: dict[str, float | None], positive: bool) -> tuple[dict[str, float | None], list[str]]:
    warnings: list[str] = []
    valid = [v for v in values.values() if v is not None]
    if not valid:
        return {k: None for k in values}, ["全样本缺失"]
    min_v, max_v = min(valid), max(valid)
    scores: dict[str, float | None] = {}
    for key, value in values.items():
        if value is None:
            scores[key] = None
            warnings.append(f"{key}缺失")
            continue
        if max_v == min_v:
            score = 80.0
        elif positive:
            score = 60 + 40 * (value - min_v) / (max_v - min_v)
        else:
            score = 60 + 40 * (max_v - value) / (max_v - min_v)
        scores[key] = max(60.0, min(100.0, score))
    return scores, warnings


def average_dimension(manager: str, scores: list[float | None], dim: str, warn: list[str]) -> float:
    valid = [s for s in scores if s is not None]
    if not valid:
        warn.append(f"{manager}-{dim}所有代表指标缺失，采用80中性分")
        return 80.0
    if len(valid) < len(scores):
        warn.append(f"{manager}-{dim}部分代表指标缺失，按可用指标计算")
    return sum(valid) / len(valid)


def is_diversified_us(row: dict[str, Any]) -> bool:
    asset = str(row.get("asset_class") or "").lower()
    category = str(row.get("category") or "").lower()
    strategy = str(row.get("strategy") or "").lower()
    niche = str(row.get("niche") or "").lower()
    segments = str(row.get("segments") or "").lower()
    text = " ".join([asset, category, strategy, niche, segments, str(row.get("fund_name") or "").lower()])
    if row.get("active_etf_flag") or row.get("smart_beta_factor_flag") or row.get("non_plain_beta_flag"):
        return True
    return any(token in text for token in ["fixed income", "bond", "commodity", "gold", "currency", "sector", "theme", "international", "emerging", "covered call", "dividend", "factor", "value", "growth", "momentum", "low volatility"])


def module_count_us(rows: list[dict[str, Any]]) -> int:
    flags = {
        "宽基": False,
        "行业主题": False,
        "策略/Smart Beta": False,
        "跨境": False,
        "债券": False,
        "商品": False,
        "货币": False,
        "主动/增强": False,
    }
    for r in rows:
        asset = str(r.get("asset_class") or "").lower()
        category = str(r.get("category") or "").lower()
        strategy = str(r.get("strategy") or "").lower()
        text = " ".join([asset, category, strategy, str(r.get("niche") or "").lower(), str(r.get("segments") or "").lower(), str(r.get("fund_name") or "").lower()])
        if "equity" in asset and any(t in text for t in ["total market", "large cap", "size and style", "broad-based"]):
            flags["宽基"] = True
        if any(t in text for t in ["sector", "theme", "thematic", "technology", "healthcare", "financial", "energy", "semiconductor"]):
            flags["行业主题"] = True
        if r.get("smart_beta_factor_flag") or any(t in text for t in ["factor", "dividend", "quality", "value", "growth", "momentum", "low volatility", "equal weight"]):
            flags["策略/Smart Beta"] = True
        if any(t in text for t in ["international", "global", "emerging", "developed markets", "china", "japan", "europe"]):
            flags["跨境"] = True
        if r.get("fixed_income_flag") or any(t in text for t in ["fixed income", "bond", "treasury", "municipal"]):
            flags["债券"] = True
        if any(t in text for t in ["commodity", "gold", "silver", "oil", "natural gas"]):
            flags["商品"] = True
        if any(t in text for t in ["money market", "cash", "ultra short"]):
            flags["货币"] = True
        if r.get("active_etf_flag") or "active" in strategy or "enhanced" in text:
            flags["主动/增强"] = True
    return sum(flags.values())


def coverage_count_us(rows: list[dict[str, Any]]) -> int:
    values = set()
    for r in rows:
        seg = str(r.get("segments") or "").strip()
        if seg:
            values.add(seg)
            continue
        niche = str(r.get("niche") or "").strip()
        cat = str(r.get("category") or "").strip()
        if niche:
            values.add(niche)
        elif cat:
            values.add(cat)
        else:
            values.add(str(r.get("fund_name") or "").strip())
    return len({v for v in values if v})


def read_us_products() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    info = {
        "path": US_FULL if US_FULL.exists() else US_TOP10_FALLBACK,
        "from_full_pool": US_FULL.exists(),
        "raw_count": None,
        "excluded_count": None,
        "included_count": None,
        "manager_count": None,
        "warn": [],
    }
    if US_FULL.exists():
        us = load_module(US_BUILDER, "us_builder")
        wb = load_workbook(US_FULL, read_only=True, data_only=True)
        ws, headers, fmap = us.detect_main_sheet(wb)
        all_rows = us.build_clean_rows(ws, headers, fmap)
        rows = [r for r in all_rows if not r.get("etn_flag")]
        info.update({
            "raw_count": len(all_rows),
            "excluded_count": len(all_rows) - len(rows),
            "included_count": len(rows),
            "manager_count": len({r["issuer"] for r in rows}),
            "sheet": ws.title,
        })
        return rows, info
    info["warn"].append("美国样本使用既有TOP10数据包，未能从全量产品池重新筛选，需人工确认TOP10名单。")
    wb = load_workbook(US_TOP10_FALLBACK, read_only=True, data_only=True)
    rows = []
    for r in read_rows(wb, "09_Top_Products_By_Issuer"):
        rows.append({
            "issuer": r.get("issuer"),
            "fund_name": r.get("fund_name"),
            "ticker": r.get("ticker"),
            "asset_class": r.get("asset_class"),
            "category": r.get("category"),
            "strategy": r.get("strategy"),
            "niche": r.get("niche"),
            "segments": None,
            "aum_usd": to_float(r.get("aum_usd")),
            "expense_ratio_pct": to_float(r.get("expense_ratio_pct")),
            "avg_daily_dollar_vol_usd": to_float(r.get("avg_daily_dollar_vol_usd")),
            "active_etf_flag": str(r.get("strategy") or "").lower() == "active",
            "smart_beta_factor_flag": False,
            "non_plain_beta_flag": str(r.get("strategy") or "").lower() != "vanilla",
            "fixed_income_flag": str(r.get("asset_class") or "").lower() == "fixed income",
        })
    info.update({"raw_count": len(rows), "excluded_count": 0, "included_count": len(rows), "manager_count": len({r["issuer"] for r in rows}), "sheet": "09_Top_Products_By_Issuer"})
    return rows, info


def read_rows(wb, sheet: str) -> list[dict[str, Any]]:
    ws = wb[sheet]
    headers = [str(c.value or "").strip() for c in ws[1]]
    out = []
    for vals in ws.iter_rows(min_row=2, values_only=True):
        row = {h: vals[i] if i < len(vals) else None for i, h in enumerate(headers) if h}
        if any(v not in (None, "") for v in row.values()):
            out.append(row)
    return out


def read_cn_products() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    cn = load_module(CN_BUILDER, "cn_builder")
    wb = load_workbook(CN_FULL, read_only=True, data_only=True)
    ws, headers, fmap = cn.detect_main_sheet(wb)
    raw_nonblank = sum(1 for vals in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in vals))
    rows, excluded_links, data_as_of = cn.build_rows(ws, headers, fmap)
    for r in rows:
        r["fund_code_norm"] = norm_code(r.get("fund_code"))
        r["fund_code_core"] = code_core(r.get("fund_code"))
    info = {
        "path": CN_FULL,
        "sheet": ws.title,
        "data_as_of": "2026-06-30 / 原始主数据截至2026年6月底",
        "raw_count": raw_nonblank,
        "excluded_count": raw_nonblank - len(rows),
        "excluded_links": excluded_links,
        "included_count": len(rows),
        "manager_count": len({r["manager"] for r in rows}),
        "from_full_pool": True,
        "warn": [],
        "supplement_path": CN_SUPPLEMENT,
    }
    return rows, info


def read_cn_supplement() -> list[dict[str, Any]]:
    if not CN_SUPPLEMENT.exists():
        return []
    wb = load_workbook(CN_SUPPLEMENT, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    out = []
    for vals in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, vals))
        code = norm_code(row.get("基金代码"))
        adv_m = to_float(row.get("近60日日均成交额(百万元)"))
        out.append({
            "fund_code": code,
            "fund_code_core": code_core(code),
            "adv_60d_cny": adv_m * 1_000_000 if adv_m is not None else None,
        })
    return out


def attach_cn_supplement(rows: list[dict[str, Any]], info: dict[str, Any]) -> None:
    supp = read_cn_supplement()
    by_code = {r["fund_code"]: r for r in supp}
    by_core = defaultdict(list)
    for r in supp:
        by_core[r["fund_code_core"]].append(r)
    matched = 0
    matched_supp = set()
    for row in rows:
        s = by_code.get(row["fund_code_norm"])
        if not s:
            candidates = by_core.get(row["fund_code_core"], [])
            s = candidates[0] if candidates else None
        if s:
            matched += 1
            matched_supp.add(s["fund_code"])
            row["adv_60d_cny"] = s.get("adv_60d_cny")
        else:
            row["adv_60d_cny"] = None
    info["supplement_rows"] = len(supp)
    info["supplement_matched_rows"] = matched
    info["supplement_unmatched_rows"] = len(supp) - len(matched_supp)
    info["supplement_match_rate"] = matched / len(rows) if rows else 0


def compute_market_metrics(rows: list[dict[str, Any]], market: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    manager_key = "issuer" if market == "US" else "manager"
    aum_key = "aum_usd" if market == "US" else "aum_cny"
    fee_key = "expense_ratio_pct" if market == "US" else "fee_pct"
    adv_key = "avg_daily_dollar_vol_usd" if market == "US" else "adv_60d_cny"
    active_threshold = 10_000_000 if market == "US" else 100_000_000
    total_market_aum = sum_field(rows, aum_key)
    by_manager: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        manager = str(row.get(manager_key) or "").strip()
        if manager:
            by_manager[manager].append(row)
    ranked = sorted(by_manager.items(), key=lambda kv: sum_field(kv[1], aum_key), reverse=True)[:10]
    metrics = []
    for rank, (manager, rs) in enumerate(ranked, 1):
        total_aum = sum_field(rs, aum_key)
        sorted_aum = sorted(rs, key=lambda r: r.get(aum_key) or 0, reverse=True)
        etf_count = len(rs)
        if market == "US":
            coverage_count = coverage_count_us(rs)
            diversified_rows = [r for r in rs if is_diversified_us(r)]
            modules = module_count_us(rs)
        else:
            coverage_count = len({r.get("index_name") or r.get("category_clean") or r.get("fund_name") for r in rs if r.get("index_name") or r.get("category_clean") or r.get("fund_name")})
            diversified_rows = [r for r in rs if r.get("non_plain_beta_flag")]
            flags = {
                "宽基": any(r.get("broad_based_flag") for r in rs),
                "行业主题": any(r.get("industry_theme_flag") for r in rs),
                "策略/Smart Beta": any(r.get("smart_beta_strategy_flag") for r in rs),
                "跨境": any(r.get("cross_border_flag") for r in rs),
                "债券": any(r.get("fixed_income_flag") for r in rs),
                "商品": any(r.get("commodity_flag") for r in rs),
                "货币": any(r.get("money_flag") for r in rs),
                "主动/增强": any("增强" in str(r.get("fund_name") or "") or "主动" in str(r.get("fund_name") or "") for r in rs),
            }
            modules = sum(flags.values())
        total_adv = sum_field(rs, adv_key)
        adv_known_count = sum(1 for r in rs if r.get(adv_key) is not None)
        metric = {
            "规模排名": rank,
            "manager": manager,
            "total_aum": total_aum,
            "market_share": total_aum / total_market_aum if total_market_aum else None,
            "etf_count": etf_count,
            "coverage_count": coverage_count,
            "weighted_fee": weighted_avg(rs, fee_key, aum_key),
            "trading_efficiency": total_adv / total_aum if total_aum and adv_known_count else None,
            "active_product_share": sum(1 for r in rs if (r.get(adv_key) or 0) >= active_threshold) / etf_count if adv_known_count else None,
            "diversified_aum_share": sum_field(diversified_rows, aum_key) / total_aum if total_aum else None,
            "module_count": modules,
            "top1_aum_share": (sorted_aum[0].get(aum_key) or 0) / total_aum if total_aum and sorted_aum else None,
            "top3_aum_share": sum_field(sorted_aum[:3], aum_key) / total_aum if total_aum else None,
            "adv_known_count": adv_known_count,
        }
        metrics.append(metric)
    return metrics, {"top10_managers": [m["manager"] for m in metrics], "total_market_aum": total_market_aum, "manager_count": len(by_manager)}


def score_market(metrics: list[dict[str, Any]], note: str) -> tuple[list[dict[str, Any]], list[str]]:
    managers = [m["manager"] for m in metrics]
    warn: list[str] = []
    specs = {
        "etf_count": True,
        "coverage_count": True,
        "weighted_fee": False,
        "trading_efficiency": True,
        "active_product_share": True,
        "diversified_aum_share": True,
        "module_count": True,
        "top1_aum_share": True,
        "top3_aum_share": True,
    }
    raw = {k: {m["manager"]: m.get(k) for m in metrics} for k in specs}
    scores = {}
    for field, positive in specs.items():
        scores[field], field_warn = standardize(raw[field], positive)
        warn.extend([f"{field}: {w}" for w in field_warn])
    by_manager = {m["manager"]: m for m in metrics}
    out = []
    for manager in managers:
        dims = {
            "产品广度": average_dimension(manager, [scores["etf_count"][manager], scores["coverage_count"][manager]], "产品广度", warn),
            "费率优势": average_dimension(manager, [scores["weighted_fee"][manager]], "费率优势", warn),
            "交易活跃": average_dimension(manager, [scores["trading_efficiency"][manager], scores["active_product_share"][manager]], "交易活跃", warn),
            "多元能力": average_dimension(manager, [scores["diversified_aum_share"][manager], scores["module_count"][manager]], "多元能力", warn),
            "旗舰产品": average_dimension(manager, [scores["top1_aum_share"][manager], scores["top3_aum_share"][manager]], "旗舰产品", warn),
        }
        total = sum(dims.values()) / len(dims)
        base = by_manager[manager]
        out.append({
            "规模排名": base["规模排名"],
            "管理人": manager,
            "ETF管理规模": base["total_aum"],
            "市场份额": base["market_share"],
            "ETF产品数量": base["etf_count"],
            **{d: round(v, 1) for d, v in dims.items()},
            "生态评分": round(total, 1),
            "评分口径说明": note,
        })
    out.sort(key=lambda r: r["生态评分"], reverse=True)
    for rank, row in enumerate(out, 1):
        row["生态排名"] = rank
    return out, warn


def write_score_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(name)
    ws.append(SCORE_HEADERS)
    for row in rows:
        ws.append([row.get(h) for h in SCORE_HEADERS])
    style_sheet(ws, score_sheet=True)


def write_rules_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("评分细则")
    ws.append(["维度", "权重", "这项看什么", "实际计算指标", "评分方向", "设计说明"])
    for row in RULE_ROWS:
        ws.append(row)
    start = ws.max_row + 2
    ws.cell(start, 1).value = "说明"
    ws.cell(start, 1).font = Font(bold=True, color="1F4E78")
    ws.cell(start + 1, 1).value = RULE_NOTE
    ws.merge_cells(start_row=start + 1, start_column=1, end_row=start + 1, end_column=6)
    style_sheet(ws, score_sheet=False, filter_ref="A1:F6")
    ws.row_dimensions[start + 1].height = 90


def write_info_sheet(wb: Workbook, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("数据信息")
    headers = ["项目", "美国", "中国", "说明"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    style_sheet(ws, score_sheet=False)


def style_sheet(ws, score_sheet: bool, filter_ref: str | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = filter_ref or ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if score_sheet and cell.column in {1, 2, 6}:
                cell.number_format = "0"
            if score_sheet and cell.column == 4:
                cell.number_format = '#,##0'
            if score_sheet and cell.column == 5:
                cell.number_format = "0.0%"
            if score_sheet and 7 <= cell.column <= 12:
                cell.number_format = "0.0"
            if score_sheet and cell.column == 12:
                cell.font = Font(bold=True)
    for col_idx in range(1, ws.max_column + 1):
        max_len = len(str(ws.cell(1, col_idx).value or ""))
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                max_len = max(max_len, min(len(str(value)), 50))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 50)
    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 24 if row_idx == 1 else 42


def info_rows(us_info: dict[str, Any], cn_info: dict[str, Any], us_market: dict[str, Any], cn_market: dict[str, Any], us_warn: list[str], cn_warn: list[str]) -> list[dict[str, Any]]:
    return [
        {"项目": "主数据文件路径", "美国": str(us_info["path"]), "中国": str(cn_info["path"]), "说明": "美国优先使用全量清洗标准化产品池；中国使用6月底全市场ETF基础信息。"},
        {"项目": "交易补充数据路径", "美国": "来自美国主数据Avg. Daily $ Vol.", "中国": str(CN_SUPPLEMENT), "说明": "中国补充数据仅用于交易活跃维度，不改变6月底样本池。"},
        {"项目": "数据截至日期", "美国": "2026-07-06", "中国": "主数据2026年6月底；交易补充2026-07-07", "说明": "严格区分样本/AUM时点与交易补充时点。"},
        {"项目": "原始ETF产品数量", "美国": us_info.get("raw_count"), "中国": cn_info.get("raw_count"), "说明": "读取主数据后的非空产品记录数量。"},
        {"项目": "剔除非ETF/联接基金数量", "美国": us_info.get("excluded_count"), "中国": cn_info.get("excluded_count"), "说明": "美国主要剔除ETN；中国剔除联接基金、LOF、非ETF等。"},
        {"项目": "纳入评分ETF产品数量", "美国": us_info.get("included_count"), "中国": cn_info.get("included_count"), "说明": "用于按管理人汇总AUM和筛选TOP10的产品池。"},
        {"项目": "全市场管理人数量", "美国": us_market.get("manager_count"), "中国": cn_market.get("manager_count"), "说明": "纳入产品池后的管理人数量。"},
        {"项目": "TOP10筛选标准", "美国": "按ETF AUM/USD汇总降序取前10", "中国": "按2026年6月底ETF AUM/CNY汇总降序取前10", "说明": "规模只用于TOP10筛选，不进入生态评分。"},
        {"项目": "TOP10管理人名单", "美国": "；".join(us_market.get("top10_managers", [])), "中国": "；".join(cn_market.get("top10_managers", [])), "说明": "规模排名口径名单。"},
        {"项目": "是否从全量产品池重新筛选TOP10", "美国": "是" if us_info.get("from_full_pool") else "否-WARN", "中国": "是", "说明": "若美国为否，需人工确认TOP10名单。"},
        {"项目": "AUM字段来源", "美国": "aum_usd", "中国": "最新基金规模(亿) -> aum_cny", "说明": "金额字段用于筛选TOP10、市场份额、旗舰占比。"},
        {"项目": "费率字段来源", "美国": "expense_ratio_pct，按AUM加权", "中国": "管理费率_支持历史(%) + 托管费率_支持历史(%)，按AUM加权", "说明": "费率优势为反向指标。"},
        {"项目": "成交字段来源", "美国": "avg_daily_dollar_vol_usd", "中国": "近60日日均成交额(百万元) × 1,000,000", "说明": "交易活跃使用成交效率和活跃产品占比。"},
        {"项目": "多元产品分类字段来源", "美国": "asset_class/category/strategy/niche/segments及主动、Smart Beta标记", "中国": "基金名称、跟踪指数、投资类型二级分类衍生模块标记", "说明": "统计多元AUM占比和能力模块覆盖数。"},
        {"项目": "旗舰产品字段来源", "美国": "产品级AUM排序计算Top1/Top3占比", "中国": "产品级AUM排序计算Top1/Top3占比", "说明": "不使用最大ETF绝对规模评分。"},
        {"项目": "中国交易补充数据匹配率", "美国": "不适用", "中国": f"{cn_info.get('supplement_matched_rows', 0)}/{cn_info.get('included_count', 0)} = {cn_info.get('supplement_match_rate', 0):.1%}; 补充未匹配{cn_info.get('supplement_unmatched_rows', 0)}条", "说明": "未匹配补充记录不新增入样本池。"},
        {"项目": "缺失字段WARN", "美国": "；".join(us_info.get("warn", []) + us_warn) or "无", "中国": "；".join(cn_info.get("warn", []) + cn_warn) or "无", "说明": "任一维度单指标或中性分在此列示。"},
        {"项目": "口径限制说明", "美国": "美国与中国分别在各自TOP10内部标准化", "中国": "中国主数据以2026年6月底为准；7.07补充交易数据只用于交易活跃", "说明": "不得将中美分数解释为绝对实力直接比较。"},
        {"项目": "QA结论", "美国": "PASS" if us_info.get("included_count") and not us_info.get("warn") else "WARN", "中国": "PASS" if cn_info.get("included_count") else "WARN", "说明": "输出4个sheet；未修改输入文件；未使用旧版六维得分或旧生态总分。"},
    ]


def main() -> None:
    us_rows, us_info = read_us_products()
    cn_rows, cn_info = read_cn_products()
    attach_cn_supplement(cn_rows, cn_info)
    us_metrics, us_market = compute_market_metrics(us_rows, "US")
    cn_metrics, cn_market = compute_market_metrics(cn_rows, "CN")
    us_scores, us_warn = score_market(us_metrics, US_NOTE)
    cn_scores, cn_warn = score_market(cn_metrics, CN_NOTE)

    wb = Workbook()
    wb.remove(wb.active)
    write_score_sheet(wb, "美国TOP10五维评分", us_scores)
    write_score_sheet(wb, "中国TOP10五维评分", cn_scores)
    write_rules_sheet(wb)
    write_info_sheet(wb, info_rows(us_info, cn_info, us_market, cn_market, us_warn, cn_warn))
    wb.save(OUTPUT)

    verify = load_workbook(OUTPUT, read_only=True, data_only=True)
    print(f"美国输入文件路径: {us_info['path']}")
    print(f"中国输入文件路径: {cn_info['path']}")
    print(f"中国交易补充文件路径: {CN_SUPPLEMENT}")
    print(f"输出文件路径: {OUTPUT}")
    print(f"美国是否从全量产品池重新筛选TOP10: {'是' if us_info.get('from_full_pool') else '否'}")
    print("中国是否从全量产品池重新筛选TOP10: 是")
    print("美国TOP10规模排名:")
    for row in sorted(us_metrics, key=lambda r: r["规模排名"]):
        print(f"- {row['规模排名']}. {row['manager']}: AUM={row['total_aum']:.0f}, ETF数量={row['etf_count']}")
    print("中国TOP10规模排名:")
    for row in sorted(cn_metrics, key=lambda r: r["规模排名"]):
        print(f"- {row['规模排名']}. {row['manager']}: AUM={row['total_aum']:.0f}, ETF数量={row['etf_count']}")
    print("美国TOP10生态评分排名:")
    for row in us_scores:
        print(f"- {row['生态排名']}. {row['管理人']}: {row['生态评分']:.1f}")
    print("中国TOP10生态评分排名:")
    for row in cn_scores:
        print(f"- {row['生态排名']}. {row['管理人']}: {row['生态评分']:.1f}")
    print("各维度字段来源:")
    print("- 产品广度: ETF产品数量 + 指数/主题覆盖数量")
    print("- 费率优势: AUM加权综合费率/expense ratio")
    print("- 交易活跃: 成交效率 + 活跃产品占比")
    print("- 多元能力: 多元产品规模占比 + 能力模块覆盖数")
    print("- 旗舰产品: 最大ETF规模占比 + Top3产品规模占比")
    all_warn = us_info.get("warn", []) + cn_info.get("warn", []) + us_warn + cn_warn
    print("缺失字段WARN:")
    print("无" if not all_warn else "\n".join(f"- {w}" for w in all_warn))
    print(f"输出sheet列表: {verify.sheetnames}")
    qa = "PASS" if verify.sheetnames == ["美国TOP10五维评分", "中国TOP10五维评分", "评分细则", "数据信息"] and len(us_scores) == 10 and len(cn_scores) == 10 else "WARN"
    print(f"QA结论: {qa}")


if __name__ == "__main__":
    main()
