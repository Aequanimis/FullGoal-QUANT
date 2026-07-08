from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE_DIR = Path.home() / "Desktop" / "Fullgoal" / "ICI课题" / "课题更正"
INPUT = BASE_DIR / "中美ETF_TOP10六维评分_战略侧重版.xlsx"
OUTPUT = BASE_DIR / "中美ETF_TOP10六维评分_展示优化版.xlsx"

SCORE_SHEETS = ["美国TOP10六维评分", "中国TOP10六维评分"]
OUTPUT_HEADERS = [
    "排名",
    "管理人",
    "产品广度",
    "规模沉淀",
    "费率优势",
    "交易活跃",
    "多元能力",
    "旗舰产品",
    "生态评分",
    "评分口径说明",
]

RENAME_MAP = {
    "产品覆盖度": "产品广度",
    "产品矩阵广度": "产品广度",
    "产品布局": "产品广度",
    "产品矩阵广度_score": "产品广度",
    "规模实力": "规模沉淀",
    "规模沉淀能力": "规模沉淀",
    "规模沉淀能力_score": "规模沉淀",
    "费率竞争": "费率优势",
    "成本竞争力": "费率优势",
    "成本竞争力_score": "费率优势",
    "交易活跃度": "交易活跃",
    "交易活跃": "交易活跃",
    "交易流动性": "交易活跃",
    "交易流动性_score": "交易活跃",
    "多元产品能力": "多元能力",
    "能力迁移": "多元能力",
    "能力迁移程度": "多元能力",
    "能力迁移程度_score": "多元能力",
    "旗舰产品力": "旗舰产品",
    "入口心智": "旗舰产品",
    "入口产品心智": "旗舰产品",
    "入口产品心智_score": "旗舰产品",
    "生态总分_战略侧重版": "生态评分",
    "total_ecosystem_score_us_comparable": "生态评分",
    "total_ecosystem_score": "生态评分",
}

DROP_FIELDS = {
    "中国增强评分",
    "产品管理质量_aux_score",
    "total_ecosystem_score_cn_enhanced",
    "total_ecosystem_score_cn_enhanced_score",
    "生态总分_中国增强版",
}

WEIGHTS = {
    "产品广度": 0.15,
    "规模沉淀": 0.20,
    "费率优势": 0.15,
    "交易活跃": 0.15,
    "多元能力": 0.20,
    "旗舰产品": 0.15,
}

US_NOTE = (
    "美国TOP10内部标准化评分，采用战略侧重权重；用于美国市场内部排序及与中国管理人画像映射，"
    "不代表中美绝对实力直接比较。"
)
CN_NOTE = (
    "中国TOP10内部标准化评分，采用战略侧重权重；交易活跃使用补充交易数据修正，用于中国市场内部排序"
    "及与美国管理人画像映射，不代表中美绝对实力直接比较。"
)

RULE_ROWS = [
    [
        "产品广度",
        "15%",
        "ETF产品线是否完整，覆盖资产类别和细分赛道是否够广",
        "ETF产品数量、资产类别覆盖数、产品类别覆盖数、指数/主题覆盖数",
        "越多越好、越均衡越好",
        "产品覆盖是基础，但不鼓励单纯铺产品，因此权重为15%",
    ],
    [
        "规模沉淀",
        "20%",
        "ETF产品是否真正做大，是否沉淀出大规模产品",
        "ETF总管理规模、10亿以上产品数量、50亿以上产品数量、Top5产品规模",
        "越大越好",
        "规模是ETF生态有效性的核心验证，因此权重为20%",
    ],
    [
        "费率优势",
        "15%",
        "ETF产品成本是否有竞争力",
        "AUM加权综合费率、Top10产品加权费率、低费率产品规模占比、平均综合费率",
        "费率越低越好",
        "低费率是ETF竞争基础，但不是唯一决定因素，因此权重为15%",
    ],
    [
        "交易活跃",
        "15%",
        "ETF产品是否被市场真实交易使用",
        "近60日日均成交额、高流动性产品数量、高流动性产品规模占比、IOPV溢折率、成交额集中度",
        "成交越活跃越好，溢折率和集中度越低越好",
        "交易活跃体现ETF工具属性，但与规模和旗舰产品有一定重叠，因此权重为15%",
    ],
    [
        "多元能力",
        "20%",
        "是否能把不同资产、策略、主题、跨境、固收、商品等能力转化为ETF产品",
        "非纯宽基产品规模占比、策略/Smart Beta产品规模占比、跨境ETF规模占比、债券/商品ETF规模占比、行业主题ETF规模占比、能力模块覆盖数",
        "覆盖越广、规模占比越高越好",
        "体现管理人从产品供给走向能力封装，是本课题关注重点，因此权重为20%",
    ],
    [
        "旗舰产品",
        "15%",
        "是否拥有客户容易识别、规模和交易表现突出的代表性ETF",
        "最大ETF规模、最大ETF成交额、最大ETF规模占比、Top3产品规模占比、旗舰产品识别度",
        "代表产品越突出越好",
        "旗舰产品有助于形成客户认知，但与规模沉淀有一定重叠，因此权重为15%",
    ],
]

EXPLANATIONS = [
    "本评分为自建研究框架，用于ETF管理人生态位比较，不代表官方排名。",
    "各底层指标先在各自市场TOP10管理人样本内标准化至60—100分，再汇总为六维得分。",
    "金额类指标使用log处理；费率、IOPV溢折率、成交集中度等为反向指标。",
    "美国和中国使用同一套六维框架和同一套战略侧重权重，但分数是在各自市场TOP10样本内标准化得到。",
    "因此，美国sheet主要用于美国管理人内部比较，中国sheet主要用于中国管理人内部比较。",
    "中美之间主要比较六维结构、生态位和学习路径，不应直接将中国管理人的分数和美国管理人的分数理解为绝对实力比较。",
]

WEIGHT_NOTE = (
    "本评分采用“战略侧重版”权重设计。产品广度是ETF业务基础，但不鼓励单纯铺产品，因此权重为15%；"
    "规模沉淀体现ETF产品是否真正做大，是生态有效性的核心验证，因此权重为20%；"
    "费率优势是ETF竞争的重要基础，但不是唯一决定因素，因此权重为15%；"
    "交易活跃体现ETF作为交易工具的使用情况，但与规模和旗舰产品存在一定重叠，因此权重为15%；"
    "多元能力体现管理人能否将策略、跨境、固收、商品、行业主题等能力转化为ETF产品，是本课题关注重点，因此权重为20%；"
    "旗舰产品体现管理人是否拥有客户容易识别的代表性产品，但与规模沉淀有一定重叠，因此权重为15%。"
)


def to_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").replace("%", "").strip())
    except ValueError:
        return None


def read_rows(ws) -> tuple[list[str], list[dict[str, Any]]]:
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {headers[i]: values[i] if i < len(values) else None for i in range(len(headers)) if headers[i]}
        if any(v not in (None, "") for v in row.values()):
            rows.append(row)
    return headers, rows


def build_field_mapping(headers: list[str]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for header in headers:
        if header in DROP_FIELDS:
            continue
        new_name = RENAME_MAP.get(header, header)
        if new_name in OUTPUT_HEADERS:
            mapping[header] = new_name
    return mapping


def normalize_score_rows(headers: list[str], rows: list[dict[str, Any]], note: str) -> tuple[list[dict[str, Any]], dict[str, str], list[str]]:
    warnings: list[str] = []
    mapping = build_field_mapping(headers)
    missing = [h for h in OUTPUT_HEADERS if h not in ("评分口径说明", "生态评分") and h not in mapping.values()]
    if missing:
        warnings.append(f"WARN: 缺少必要字段映射：{missing}")

    out_rows = []
    for source_row in rows:
        out = {header: None for header in OUTPUT_HEADERS}
        for old_name, new_name in mapping.items():
            out[new_name] = source_row.get(old_name)
        out["评分口径说明"] = note

        if out.get("生态评分") in (None, ""):
            if all(to_float(out.get(dim)) is not None for dim in WEIGHTS):
                out["生态评分"] = sum(to_float(out[dim]) * weight for dim, weight in WEIGHTS.items())
                warnings.append("WARN: 原表未识别生态评分字段，已按战略侧重权重重新计算生态评分")
            else:
                warnings.append(f"WARN: {out.get('管理人')} 缺少生态评分且六维不完整，无法计算")

        for dim in list(WEIGHTS) + ["生态评分"]:
            value = to_float(out.get(dim))
            out[dim] = round(value, 1) if value is not None else None
        rank_value = to_float(out.get("排名"))
        out["排名"] = int(rank_value) if rank_value is not None else None
        out_rows.append(out)

    return out_rows, mapping, warnings


def style_score_sheet(ws) -> None:
    apply_base_style(ws, max_width=50)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if 3 <= cell.column <= 9:
                cell.number_format = "0.0"
            if cell.column == 1:
                cell.number_format = "0"
            if cell.column == 9:
                cell.font = Font(bold=True)
            if cell.column == 10:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 42


def apply_base_style(ws, max_width: int = 50, filter_ref: str | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = filter_ref or ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col_idx in range(1, ws.max_column + 1):
        max_len = len(str(ws.cell(1, col_idx).value or ""))
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                max_len = max(max_len, min(len(str(value)), max_width))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), max_width)
    ws.row_dimensions[1].height = 24


def write_score_sheet(wb: Workbook, name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(name)
    ws.append(OUTPUT_HEADERS)
    for row in rows:
        ws.append([row.get(header) for header in OUTPUT_HEADERS])
    style_score_sheet(ws)


def write_rules_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("评分细则")
    headers = ["维度", "权重", "这项看什么", "主要参考指标", "评分方向", "权重设置原因"]
    ws.append(headers)
    for row in RULE_ROWS:
        ws.append(row)

    start = ws.max_row + 2
    ws.cell(start, 1).value = "评分说明"
    ws.cell(start, 1).font = Font(bold=True, color="1F4E78")
    for idx, text in enumerate(EXPLANATIONS, start + 1):
        ws.cell(idx, 1).value = idx - start
        ws.cell(idx, 2).value = text
        ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=6)

    weight_start = ws.max_row + 2
    ws.cell(weight_start, 1).value = "权重设计说明"
    ws.cell(weight_start, 1).font = Font(bold=True, color="1F4E78")
    ws.cell(weight_start + 1, 1).value = WEIGHT_NOTE
    ws.merge_cells(start_row=weight_start + 1, start_column=1, end_row=weight_start + 1, end_column=6)

    apply_base_style(ws, max_width=50, filter_ref="A1:F7")
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 48
    ws.row_dimensions[weight_start + 1].height = 90


def main() -> None:
    if not INPUT.exists():
        raise FileNotFoundError(f"输入文件不存在：{INPUT}")

    source_wb = load_workbook(INPUT, read_only=True, data_only=True)
    output_wb = Workbook()
    output_wb.remove(output_wb.active)

    all_warnings: list[str] = []
    field_maps: dict[str, dict[str, str]] = {}
    for sheet_name in SCORE_SHEETS:
        if sheet_name not in source_wb.sheetnames:
            raise ValueError(f"输入文件缺少sheet：{sheet_name}")
        headers, source_rows = read_rows(source_wb[sheet_name])
        note = US_NOTE if sheet_name.startswith("美国") else CN_NOTE
        rows, mapping, warnings = normalize_score_rows(headers, source_rows, note)
        field_maps[sheet_name] = mapping
        all_warnings.extend([f"{sheet_name}: {w}" for w in warnings])
        write_score_sheet(output_wb, sheet_name, rows)

    write_rules_sheet(output_wb)
    output_wb.save(OUTPUT)

    verify_wb = load_workbook(OUTPUT, read_only=False, data_only=True)
    print(f"输入文件路径: {INPUT}")
    print(f"输出文件路径: {OUTPUT}")
    for sheet_name in SCORE_SHEETS:
        print(f"{sheet_name}字段改名前后对照:")
        for old_name, new_name in field_maps[sheet_name].items():
            if old_name != new_name:
                print(f"- {old_name} -> {new_name}")
    print(f"是否成功新增“评分细则”sheet: {'是' if '评分细则' in verify_wb.sheetnames else '否'}")
    print(f"输出sheet列表: {verify_wb.sheetnames}")
    if all_warnings:
        for warning in all_warnings:
            print(warning)
    else:
        print("WARN: 无")


if __name__ == "__main__":
    main()
