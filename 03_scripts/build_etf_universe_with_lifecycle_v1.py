from __future__ import annotations

import math
import re
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(r"C:\Users\lvdon\Desktop\Fullgoal\课题研究\处理后数据\境外策略ETF数据")
SCREENER_XLSX = BASE_DIR / "原始数据" / "境外ETF全量数据 0706.xlsx"
LAUNCHES_XLSX = BASE_DIR / "处理结果" / "ETF_Launches_汇总.xlsx"
CLOSURES_XLSX = BASE_DIR / "处理结果" / "ETF_Closures_汇总.xlsx"
OUTPUT_DIR = BASE_DIR / "处理结果"

OUTPUT_XLSX = OUTPUT_DIR / "境外ETF全量数据_清洗整合_含LaunchesClosures_v1.xlsx"
OUTPUT_CSV = OUTPUT_DIR / "ETF_Universe_All_清洗整合_v1.csv"
OUTPUT_LOG = OUTPUT_DIR / "境外ETF全量数据_清洗整合_含LaunchesClosures_v1_日志.txt"

SCREENER_SOURCE_FILE = SCREENER_XLSX.name

SCREENER_FIELDS = [
    "Ticker",
    "Fund Name",
    "Issuer",
    "Asset Class",
    "Category",
    "Strategy",
    "Niche",
    "Segments",
    "Selection Criteria",
    "Weighting Scheme",
    "AUM",
    "Expense Ratio",
    "Avg. Daily $ Vol.",
    "Spread Data %",
    "P/E",
    "P/B",
    "Dividend Yield",
    "1 Yr Fund Flow",
    "YTD Fund Flow",
    "1 MO Fund Flow",
    "1 YR TR",
    "3 Yr. Total Return",
    "5 Yr. Total Return",
    "ETN",
    "Leveraged",
    "Leveraged Ratio",
    "Inverse",
    "Region",
    "Geography",
    "Economic Development",
    "Avg. Daily Share Vol.",
]

FINAL_COLUMNS = [
    "Ticker",
    "Fund Name",
    "Issuer",
    "Brand",
    "Asset Class",
    "Category",
    "Strategy",
    "Niche",
    "Segments",
    "Selection Criteria",
    "Weighting Scheme",
    "Status",
    "Data_Type",
    "Data_Source",
    "Is_Current_Screener",
    "Inception Date",
    "Inception Year",
    "Closure Date",
    "Closure Year",
    "Lifecycle_Match_Flag",
    "AUM",
    "AUM_USD",
    "AUM_USD_100M",
    "Expense Ratio",
    "Expense_Ratio_Pct",
    "Spread Data %",
    "Spread_Data_Pct",
    "Avg. Daily $ Vol.",
    "Avg_Daily_Dollar_Vol_USD",
    "Avg_Daily_Dollar_Vol_USD_100M",
    "Avg. Daily Share Vol.",
    "Avg_Daily_Share_Vol",
    "P/E",
    "PE",
    "P/B",
    "PB",
    "Dividend Yield",
    "Dividend_Yield_Pct",
    "1 YR TR",
    "Total_Return_1Y_Pct",
    "3 Yr. Total Return",
    "Total_Return_3Y_Pct",
    "5 Yr. Total Return",
    "Total_Return_5Y_Pct",
    "1 Yr Fund Flow",
    "Fund_Flow_1Y_USD",
    "YTD Fund Flow",
    "Fund_Flow_YTD_USD",
    "1 MO Fund Flow",
    "Fund_Flow_1MO_USD",
    "ETN",
    "Is_ETN",
    "Leveraged",
    "Is_Leveraged",
    "Leveraged Ratio",
    "Leveraged_Ratio_Num",
    "Inverse",
    "Is_Inverse",
    "Region",
    "Geography",
    "Economic Development",
    "Screener_Source_File",
    "Launch_Source_File",
    "Closure_Source_File",
    "Raw_Row_ID",
]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    text = str(value).replace("\ufeff", "").replace("\r", " ").replace("\n", " ")
    text = re.sub(r"[\u00a0\u1680\u180e\u2000-\u200a\u2028\u2029\u202f\u205f\u3000]", " ", text)
    text = re.sub(r"[\u200b-\u200d\u2060]", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_name_for_match(value: Any) -> str:
    text = clean_text(value).lower()
    text = text.replace("&", " and ")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\b(etf|fund|trust|shares|portfolio|index)\b", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def is_missing_like(value: Any) -> bool:
    text = clean_text(value)
    return text == "" or text in {"-", "--", "N/A", "n/a", "nan", "NaN"} or text.lower() == "undefined%"


def parse_amount(value: Any) -> tuple[float | None, str | None]:
    if is_missing_like(value):
        return None, None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value), None
    text = clean_text(value).replace("−", "-").replace("—", "-").replace("–", "-")
    text = text.replace(",", "")
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative = True
        text = text[1:-1].strip()
    match = re.fullmatch(r"(-)?\$?\s*([0-9]+(?:\.[0-9]+)?)\s*([KMBkmb]?)", text)
    if not match:
        return None, f"amount_parse_failed:{clean_text(value)}"
    sign = -1 if match.group(1) or negative else 1
    multiplier = {"": 1, "K": 1_000, "M": 1_000_000, "B": 1_000_000_000}[match.group(3).upper()]
    return sign * float(match.group(2)) * multiplier, None


def parse_percent(value: Any) -> tuple[float | None, str | None]:
    if is_missing_like(value):
        return None, None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value) * 100, None
    text = clean_text(value).replace("−", "-").replace("—", "-").replace("–", "-").replace(",", "")
    try:
        if text.endswith("%"):
            return float(text[:-1].strip()), None
        return float(text) * 100, None
    except ValueError:
        return None, f"percent_parse_failed:{clean_text(value)}"


def parse_number(value: Any) -> tuple[float | None, str | None]:
    if is_missing_like(value):
        return None, None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value), None
    text = clean_text(value).replace(",", "")
    try:
        return float(text), None
    except ValueError:
        return None, f"number_parse_failed:{clean_text(value)}"


def parse_yes_no(value: Any) -> tuple[str, str | None]:
    if is_missing_like(value):
        return "", None
    text = clean_text(value).lower()
    if text in {"yes", "y", "true", "1"}:
        return "Yes", None
    if text in {"no", "n", "false", "0"}:
        return "No", None
    return "", f"yes_no_parse_failed:{clean_text(value)}"


def parse_leveraged_ratio(value: Any) -> tuple[float | None, str | None]:
    if is_missing_like(value):
        return None, None
    text = clean_text(value).upper().replace(" ", "")
    match = re.fullmatch(r"([+-]?[0-9]+(?:\.[0-9]+)?)X?", text)
    if not match:
        return None, f"leveraged_ratio_parse_failed:{clean_text(value)}"
    return float(match.group(1)), None


def parse_dates(values: pd.Series) -> pd.Series:
    raw = values.fillna("").astype(str).map(clean_text)
    parsed = pd.Series(pd.NaT, index=raw.index, dtype="datetime64[ns]")
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
        remaining = parsed.isna() & raw.ne("")
        if not remaining.any():
            break
        parsed.loc[remaining] = pd.to_datetime(raw.loc[remaining], format=fmt, errors="coerce")
    remaining = parsed.isna() & raw.ne("")
    if remaining.any():
        parsed.loc[remaining] = pd.to_datetime(raw.loc[remaining], errors="coerce")
    return parsed


def date_to_text(series: pd.Series) -> pd.Series:
    return series.dt.strftime("%Y-%m-%d").fillna("")


def year_to_int_series(series: pd.Series) -> pd.Series:
    return series.dt.year.astype("Int64")


def name_similarity(left: Any, right: Any) -> float:
    left_norm = normalize_name_for_match(left)
    right_norm = normalize_name_for_match(right)
    if not left_norm or not right_norm:
        return 0.0
    return SequenceMatcher(None, left_norm, right_norm).ratio()


def read_excel_first_sheet(path: Path) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=0, dtype=object, keep_default_na=False)


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.columns = [clean_text(col) for col in out.columns]
    return out


def clean_launches(path: Path) -> pd.DataFrame:
    launches = normalize_columns(read_excel_first_sheet(path))
    for col in ["Inception Date", "Inception Year", "Ticker", "Fund Name", "Source_File"]:
        if col not in launches.columns:
            launches[col] = ""
    launches["Ticker"] = launches["Ticker"].map(clean_text).str.upper()
    launches["Fund Name"] = launches["Fund Name"].map(clean_text)
    launches["_fund_name_norm"] = launches["Fund Name"].map(normalize_name_for_match)
    parsed = parse_dates(launches["Inception Date"])
    launches["_parsed_inception_date"] = parsed
    launches["Inception Date"] = date_to_text(parsed)
    launches["Inception Year"] = year_to_int_series(parsed)
    launches["Source_File"] = launches["Source_File"].map(clean_text)
    return launches


def clean_closures(path: Path) -> tuple[pd.DataFrame, int]:
    closures = normalize_columns(read_excel_first_sheet(path))
    for col in ["Closure Date", "Closure Year", "Ticker", "Fund Name", "Status", "Source_File"]:
        if col not in closures.columns:
            closures[col] = ""
    closures["Ticker"] = closures["Ticker"].map(clean_text).str.upper()
    closures["Fund Name"] = closures["Fund Name"].map(clean_text)
    closures["_fund_name_norm"] = closures["Fund Name"].map(normalize_name_for_match)
    parsed = parse_dates(closures["Closure Date"])
    closures["_parsed_closure_date"] = parsed
    closures["Closure Date"] = date_to_text(parsed)
    closures["Closure Year"] = year_to_int_series(parsed)
    closures["Status"] = "Closed"
    closures["Source_File"] = closures["Source_File"].map(clean_text)
    before = len(closures)
    closures = closures.drop_duplicates(subset=["Ticker", "Fund Name", "Closure Date"], keep="first").copy()
    return closures, before - len(closures)


def first_valid_by_earliest(group: pd.DataFrame, date_col: str) -> pd.Series:
    sorted_group = group.sort_values([date_col, "Fund Name", "Source_File"], na_position="last", kind="stable")
    return sorted_group.iloc[0]


def prepare_launch_indexes(launches: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    sorted_launches = launches.sort_values(
        ["Ticker", "_parsed_inception_date", "Fund Name", "Source_File"],
        na_position="last",
        kind="stable",
    )
    by_ticker = sorted_launches.loc[sorted_launches["Ticker"].ne("")].drop_duplicates(
        subset=["Ticker"],
        keep="first",
    )
    by_name = sorted_launches.loc[
        sorted_launches["Ticker"].ne("") & sorted_launches["_fund_name_norm"].ne("")
    ].drop_duplicates(
        subset=["Ticker", "_fund_name_norm"],
        keep="first",
    )
    return by_ticker, by_name


def build_active_current(
    screener_raw: pd.DataFrame,
    launches: pd.DataFrame,
    launch_by_ticker: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, int]]:
    active = normalize_columns(screener_raw)
    missing_fields = [col for col in SCREENER_FIELDS if col not in active.columns]
    for col in SCREENER_FIELDS:
        if col not in active.columns:
            active[col] = ""

    for col in SCREENER_FIELDS:
        active[col] = active[col].map(clean_text)
    active["Ticker"] = active["Ticker"].str.upper()
    active["Brand"] = active["Issuer"]
    active["Raw_Row_ID"] = [f"Screener_{i}" for i in range(1, len(active) + 1)]
    active["Data_Source"] = "ETF.com Screener"
    active["Data_Type"] = "Active_Current"
    active["Status"] = "Active"
    active["Is_Current_Screener"] = 1
    active["Screener_Source_File"] = SCREENER_SOURCE_FILE
    active["Closure Date"] = ""
    active["Closure Year"] = pd.Series([pd.NA] * len(active), dtype="Int64")
    active["Closure_Source_File"] = ""

    amount_fields = {
        "AUM": "AUM_USD",
        "Avg. Daily $ Vol.": "Avg_Daily_Dollar_Vol_USD",
        "Avg. Daily Share Vol.": "Avg_Daily_Share_Vol",
        "1 Yr Fund Flow": "Fund_Flow_1Y_USD",
        "YTD Fund Flow": "Fund_Flow_YTD_USD",
        "1 MO Fund Flow": "Fund_Flow_1MO_USD",
    }
    percent_fields = {
        "Expense Ratio": "Expense_Ratio_Pct",
        "Spread Data %": "Spread_Data_Pct",
        "Dividend Yield": "Dividend_Yield_Pct",
        "1 YR TR": "Total_Return_1Y_Pct",
        "3 Yr. Total Return": "Total_Return_3Y_Pct",
        "5 Yr. Total Return": "Total_Return_5Y_Pct",
    }
    number_fields = {"P/E": "PE", "P/B": "PB"}
    bool_fields = {"ETN": "Is_ETN", "Leveraged": "Is_Leveraged", "Inverse": "Is_Inverse"}

    parse_failures: dict[str, int] = {
        "amount_parse_failed": 0,
        "percent_parse_failed": 0,
        "number_parse_failed": 0,
        "yes_no_parse_failed": 0,
        "leveraged_ratio_parse_failed": 0,
        "missing_screener_fields": len(missing_fields),
    }

    for source_col, target_col in amount_fields.items():
        parsed = active[source_col].map(parse_amount)
        active[target_col] = parsed.map(lambda item: item[0])
        parse_failures["amount_parse_failed"] += int(parsed.map(lambda item: item[1] is not None).sum())

    active["AUM_USD_100M"] = active["AUM_USD"] / 100_000_000
    active["Avg_Daily_Dollar_Vol_USD_100M"] = active["Avg_Daily_Dollar_Vol_USD"] / 100_000_000

    for source_col, target_col in percent_fields.items():
        parsed = active[source_col].map(parse_percent)
        active[target_col] = parsed.map(lambda item: item[0])
        parse_failures["percent_parse_failed"] += int(parsed.map(lambda item: item[1] is not None).sum())

    for source_col, target_col in number_fields.items():
        parsed = active[source_col].map(parse_number)
        active[target_col] = parsed.map(lambda item: item[0])
        parse_failures["number_parse_failed"] += int(parsed.map(lambda item: item[1] is not None).sum())

    for source_col, target_col in bool_fields.items():
        parsed = active[source_col].map(parse_yes_no)
        active[target_col] = parsed.map(lambda item: item[0])
        parse_failures["yes_no_parse_failed"] += int(parsed.map(lambda item: item[1] is not None).sum())

    ratio_parsed = active["Leveraged Ratio"].map(parse_leveraged_ratio)
    active["Leveraged_Ratio_Num"] = ratio_parsed.map(lambda item: item[0])
    parse_failures["leveraged_ratio_parse_failed"] = int(ratio_parsed.map(lambda item: item[1] is not None).sum())

    ticker_counts = launches.loc[launches["Ticker"].ne("")].groupby("Ticker").size()
    launch_pick = launch_by_ticker.set_index("Ticker", drop=False)

    launch_dates: list[str] = []
    launch_years: list[Any] = []
    launch_sources: list[str] = []
    flags: list[str] = []
    matched_names: list[str] = []
    similarities: list[float] = []

    for _, row in active.iterrows():
        ticker = row["Ticker"]
        count = int(ticker_counts.get(ticker, 0))
        if count == 0:
            launch_dates.append("")
            launch_years.append(pd.NA)
            launch_sources.append("")
            flags.append("No_Launch_Match")
            matched_names.append("")
            similarities.append(0.0)
            continue
        match = launch_pick.loc[ticker]
        similarity = name_similarity(row["Fund Name"], match["Fund Name"])
        flag = "Multiple_Launch_Matches" if count > 1 else "Matched_By_Ticker"
        if count == 1 and similarity < 0.82:
            flag = "Fund_Name_Mismatch"
        launch_dates.append(match["Inception Date"])
        launch_years.append(match["Inception Year"])
        launch_sources.append(match["Source_File"])
        flags.append(flag)
        matched_names.append(match["Fund Name"])
        similarities.append(round(similarity, 4))

    active["Inception Date"] = launch_dates
    active["Inception Year"] = pd.Series(launch_years, dtype="Int64")
    active["Launch_Source_File"] = launch_sources
    active["Lifecycle_Match_Flag"] = flags
    active["_launch_matched_fund_name"] = matched_names
    active["_launch_name_similarity"] = similarities

    audit = active[
        [
            "Ticker",
            "Fund Name",
            "_launch_matched_fund_name",
            "_launch_name_similarity",
            "Inception Date",
            "Inception Year",
            "Launch_Source_File",
            "Lifecycle_Match_Flag",
            "Raw_Row_ID",
        ]
    ].rename(
        columns={
            "_launch_matched_fund_name": "Matched_Launch_Fund_Name",
            "_launch_name_similarity": "Fund_Name_Similarity",
            "Lifecycle_Match_Flag": "Launch_Match_Flag",
        }
    )
    return active[FINAL_COLUMNS].copy(), audit, parse_failures


def build_closed_history(
    closures: pd.DataFrame,
    launches: pd.DataFrame,
    launch_by_ticker: pd.DataFrame,
    launch_by_name: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    ticker_counts = launches.loc[launches["Ticker"].ne("")].groupby("Ticker").size()
    ticker_pick = launch_by_ticker.set_index("Ticker", drop=False)
    name_pick = launch_by_name.set_index(["Ticker", "_fund_name_norm"], drop=False)

    closed = pd.DataFrame(index=closures.index)
    for col in FINAL_COLUMNS:
        closed[col] = ""
    closed["Ticker"] = closures["Ticker"].values
    closed["Fund Name"] = closures["Fund Name"].values
    closed["Status"] = "Closed"
    closed["Data_Type"] = "Closed_History"
    closed["Data_Source"] = "ETF.com Closures"
    closed["Is_Current_Screener"] = 0
    closed["Closure Date"] = closures["Closure Date"].values
    closed["Closure Year"] = closures["Closure Year"].values
    closed["Closure_Source_File"] = closures["Source_File"].values
    closed["Raw_Row_ID"] = [f"Closure_{i}" for i in range(1, len(closed) + 1)]

    inception_dates: list[str] = []
    inception_years: list[Any] = []
    launch_sources: list[str] = []
    flags: list[str] = []
    matched_names: list[str] = []

    for _, row in closures.iterrows():
        ticker = row["Ticker"]
        name_norm = row["_fund_name_norm"]
        exact_key = (ticker, name_norm)
        if ticker and name_norm and exact_key in name_pick.index:
            match = name_pick.loc[exact_key]
            if isinstance(match, pd.DataFrame):
                match = first_valid_by_earliest(match, "_parsed_inception_date")
            inception_dates.append(match["Inception Date"])
            inception_years.append(match["Inception Year"])
            launch_sources.append(match["Source_File"])
            flags.append("Matched_By_Ticker_And_Name")
            matched_names.append(match["Fund Name"])
            continue

        count = int(ticker_counts.get(ticker, 0))
        if count == 0:
            inception_dates.append("")
            inception_years.append(pd.NA)
            launch_sources.append("")
            flags.append("No_Launch_Match")
            matched_names.append("")
        elif count == 1:
            match = ticker_pick.loc[ticker]
            inception_dates.append(match["Inception Date"])
            inception_years.append(match["Inception Year"])
            launch_sources.append(match["Source_File"])
            flags.append("Matched_By_Unique_Ticker")
            matched_names.append(match["Fund Name"])
        else:
            inception_dates.append("")
            inception_years.append(pd.NA)
            launch_sources.append("")
            flags.append("Ambiguous_Ticker_Reused")
            matched_names.append("")

    closed["Inception Date"] = inception_dates
    closed["Inception Year"] = pd.Series(inception_years, dtype="Int64")
    closed["Launch_Source_File"] = launch_sources
    closed["Lifecycle_Match_Flag"] = flags

    audit = closed[
        [
            "Ticker",
            "Fund Name",
            "Closure Date",
            "Closure Year",
            "Inception Date",
            "Inception Year",
            "Launch_Source_File",
            "Lifecycle_Match_Flag",
            "Closure_Source_File",
            "Raw_Row_ID",
        ]
    ].copy()
    audit.insert(2, "Matched_Launch_Fund_Name", matched_names)
    audit = audit.rename(columns={"Lifecycle_Match_Flag": "Closure_Match_Flag"})
    return closed[FINAL_COLUMNS].copy(), audit


def build_ticker_reuse_anomalies(
    launches: pd.DataFrame,
    closures: pd.DataFrame,
    active: pd.DataFrame,
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []

    closure_multi_date = closures.loc[closures["Ticker"].ne("") & closures["Closure Date"].ne("")].groupby("Ticker")[
        "Closure Date"
    ].nunique()
    for ticker in sorted(closure_multi_date[closure_multi_date > 1].index):
        group = closures.loc[closures["Ticker"].eq(ticker)]
        rows.append(
            {
                "Anomaly_Type": "Closures 中同一 Ticker 对应多个 Closure Date",
                "Ticker": ticker,
                "Record_Count": len(group),
                "Values": ", ".join(sorted(group["Closure Date"].dropna().astype(str).unique())),
                "Source_Files": ", ".join(sorted(group["Source_File"].dropna().astype(str).unique())),
            }
        )

    closure_multi_name = closures.loc[closures["Ticker"].ne("") & closures["Fund Name"].ne("")].groupby("Ticker")[
        "Fund Name"
    ].nunique()
    for ticker in sorted(closure_multi_name[closure_multi_name > 1].index):
        group = closures.loc[closures["Ticker"].eq(ticker)]
        rows.append(
            {
                "Anomaly_Type": "Closures 中同一 Ticker 对应多个 Fund Name",
                "Ticker": ticker,
                "Record_Count": len(group),
                "Values": " | ".join(sorted(group["Fund Name"].dropna().astype(str).unique())),
                "Source_Files": ", ".join(sorted(group["Source_File"].dropna().astype(str).unique())),
            }
        )

    launch_multi_date = launches.loc[launches["Ticker"].ne("") & launches["Inception Date"].ne("")].groupby("Ticker")[
        "Inception Date"
    ].nunique()
    for ticker in sorted(launch_multi_date[launch_multi_date > 1].index):
        group = launches.loc[launches["Ticker"].eq(ticker)]
        rows.append(
            {
                "Anomaly_Type": "Launches 中同一 Ticker 对应多个 Inception Date",
                "Ticker": ticker,
                "Record_Count": len(group),
                "Values": ", ".join(sorted(group["Inception Date"].dropna().astype(str).unique())),
                "Source_Files": ", ".join(sorted(group["Source_File"].dropna().astype(str).unique())),
            }
        )

    active_tickers = set(active.loc[active["Ticker"].ne(""), "Ticker"])
    closure_tickers = set(closures.loc[closures["Ticker"].ne(""), "Ticker"])
    for ticker in sorted(active_tickers & closure_tickers):
        group = closures.loc[closures["Ticker"].eq(ticker)]
        active_names = active.loc[active["Ticker"].eq(ticker), "Fund Name"].dropna().astype(str).unique()
        rows.append(
            {
                "Anomaly_Type": "当前 Active 主表中 Ticker 与 Closures 表重复出现",
                "Ticker": ticker,
                "Record_Count": len(group),
                "Values": "Active: " + " | ".join(sorted(active_names)) + " ; Closed Dates: "
                + ", ".join(sorted(group["Closure Date"].dropna().astype(str).unique())),
                "Source_Files": ", ".join(sorted(group["Source_File"].dropna().astype(str).unique())),
            }
        )

    return pd.DataFrame(rows) if rows else pd.DataFrame(columns=["Anomaly_Type", "Ticker", "Record_Count", "Values", "Source_Files"])


def build_field_coverage(universe: pd.DataFrame) -> pd.DataFrame:
    active_count = int(universe["Status"].eq("Active").sum())
    closed_count = int(universe["Status"].eq("Closed").sum())
    total_count = len(universe)

    def coverage(field: str, mask: pd.Series | None = None) -> tuple[int, float]:
        target = universe if mask is None else universe.loc[mask]
        if len(target) == 0:
            return 0, 0.0
        nonblank = target[field].notna() & target[field].astype(str).str.strip().ne("")
        count = int(nonblank.sum())
        return count, round(count / len(target), 4)

    metrics = [
        ("Active产品数量", active_count, ""),
        ("Closed产品数量", closed_count, ""),
        ("总产品记录数", total_count, ""),
    ]
    for label, field in [
        ("Inception Date 覆盖率", "Inception Date"),
        ("Closure Date 覆盖率", "Closure Date"),
        ("AUM 覆盖率", "AUM_USD"),
        ("Expense Ratio 覆盖率", "Expense_Ratio_Pct"),
        ("1Y Return 覆盖率", "Total_Return_1Y_Pct"),
        ("3Y Return 覆盖率", "Total_Return_3Y_Pct"),
        ("5Y Return 覆盖率", "Total_Return_5Y_Pct"),
    ]:
        count, ratio = coverage(field)
        metrics.append((label, ratio, f"nonblank={count}/{total_count}"))
    return pd.DataFrame(metrics, columns=["Metric", "Value", "Note"])


def build_yearly_summary(launches: pd.DataFrame, closures: pd.DataFrame, active: pd.DataFrame) -> pd.DataFrame:
    launch_counts = launches.loc[launches["Inception Year"].notna(), "Inception Year"].astype(int).value_counts()
    closure_counts = closures.loc[closures["Closure Year"].notna(), "Closure Year"].astype(int).value_counts()
    active_counts = active.loc[active["Inception Year"].notna(), "Inception Year"].astype(int).value_counts()
    years = sorted(set(launch_counts.index) | set(closure_counts.index) | set(active_counts.index))
    rows = []
    for year in years:
        launch_count = int(launch_counts.get(year, 0))
        closure_count = int(closure_counts.get(year, 0))
        rows.append(
            {
                "Year": year,
                "Launch Count": launch_count,
                "Closure Count": closure_count,
                "Active_Current_Count": int(active_counts.get(year, 0)),
                "Net Launch Count": launch_count - closure_count,
                "Note": "Closure 数据从 2016 年开始，2010-2015 年清盘数据可能缺失" if 2010 <= year <= 2015 else "",
            }
        )
    return pd.DataFrame(rows)


def build_processing_log_sheet(metrics: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame([{"Metric": key, "Value": value} for key, value in metrics.items()])


def autosize_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for worksheet in workbook.worksheets:
        if worksheet.max_row >= 1:
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_len = 0
            for cell in column_cells[:300]:
                max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
            worksheet.column_dimensions[column_letter].width = min(max(max_len + 2, 11), 52)
    workbook.save(path)


def safe_excel_value_frame(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in out.columns:
        if str(out[col].dtype) == "Int64":
            out[col] = out[col].astype(object).where(out[col].notna(), "")
    return out


def main() -> None:
    run_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    screener_raw = read_excel_first_sheet(SCREENER_XLSX)
    launches = clean_launches(LAUNCHES_XLSX)
    closures, closure_duplicate_rows_deleted = clean_closures(CLOSURES_XLSX)
    launch_by_ticker, launch_by_name = prepare_launch_indexes(launches)

    active, launch_audit, parse_failures = build_active_current(screener_raw, launches, launch_by_ticker)
    closed, closure_audit = build_closed_history(closures, launches, launch_by_ticker, launch_by_name)
    universe = pd.concat([active, closed], ignore_index=True)

    ticker_anomalies = build_ticker_reuse_anomalies(launches, closures, active)
    field_coverage = build_field_coverage(universe)
    yearly_summary = build_yearly_summary(launches, closures, active)

    launch_match_counts = launch_audit["Launch_Match_Flag"].value_counts().to_dict()
    closure_match_counts = closure_audit["Closure_Match_Flag"].value_counts().to_dict()
    closure_reuse_tickers = closures.loc[closures["Ticker"].ne("") & closures["Closure Date"].ne("")].groupby("Ticker")[
        "Closure Date"
    ].nunique()
    closure_reuse_count = int((closure_reuse_tickers > 1).sum())
    missing_screener_fields = [field for field in SCREENER_FIELDS if field not in normalize_columns(screener_raw).columns]

    metrics = {
        "输入文件_Screener": str(SCREENER_XLSX),
        "输入文件_Launches": str(LAUNCHES_XLSX),
        "输入文件_Closures": str(CLOSURES_XLSX),
        "输出文件_Excel": str(OUTPUT_XLSX),
        "输出文件_CSV": str(OUTPUT_CSV),
        "输出文件_日志": str(OUTPUT_LOG),
        "运行时间": run_time,
        "当前 Screener 原始行数": len(screener_raw),
        "Launches 原始行数": len(launches),
        "Closures 原始行数": len(closures) + closure_duplicate_rows_deleted,
        "Active 清洗后行数": len(active),
        "Closed 清洗后行数": len(closed),
        "最终总表行数": len(universe),
        "Closures 删除重复行数量": closure_duplicate_rows_deleted,
        "关键异常数量": len(ticker_anomalies),
        "Launch 日期解析失败数量": int(launches["_parsed_inception_date"].isna().sum()),
        "Closure 日期解析失败数量": int(closures["_parsed_closure_date"].isna().sum()),
        "金额解析失败数量": parse_failures["amount_parse_failed"],
        "百分比解析失败数量": parse_failures["percent_parse_failed"],
        "估值解析失败数量": parse_failures["number_parse_failed"],
        "Yes/No 解析失败数量": parse_failures["yes_no_parse_failed"],
        "Leveraged Ratio 解析失败数量": parse_failures["leveraged_ratio_parse_failed"],
        "缺失 Screener 字段": ", ".join(missing_screener_fields) if missing_screener_fields else "无",
    }

    processing_log_sheet = build_processing_log_sheet(metrics)

    universe.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        safe_excel_value_frame(universe).to_excel(writer, sheet_name="ETF_Universe_All", index=False)
        safe_excel_value_frame(active).to_excel(writer, sheet_name="Active_Current_Cleaned", index=False)
        safe_excel_value_frame(closed).to_excel(writer, sheet_name="Closed_History_Cleaned", index=False)
        field_coverage.to_excel(writer, sheet_name="Field_Coverage_Summary", index=False)
        yearly_summary.to_excel(writer, sheet_name="Yearly_Lifecycle_Summary", index=False)
        ticker_anomalies.to_excel(writer, sheet_name="Ticker_Reuse_Anomalies", index=False)
        launch_audit.to_excel(writer, sheet_name="Launch_Match_Audit", index=False)
        closure_audit.to_excel(writer, sheet_name="Closure_Match_Audit", index=False)
        processing_log_sheet.to_excel(writer, sheet_name="Processing_Log", index=False)
    autosize_workbook(OUTPUT_XLSX)

    log_lines = [
        "境外 ETF Universe 清洗整合日志",
        "=" * 50,
        *[f"{key}: {value}" for key, value in metrics.items()],
        "",
        "Launch 匹配分布:",
        *[f"- {key}: {value}" for key, value in sorted(launch_match_counts.items())],
        "",
        "Closure 匹配分布:",
        *[f"- {key}: {value}" for key, value in sorted(closure_match_counts.items())],
        "",
        "年度生命周期统计:",
        yearly_summary.to_string(index=False),
    ]
    OUTPUT_LOG.write_text("\n".join(log_lines) + "\n", encoding="utf-8-sig")

    active_match_success = int(active["Inception Date"].astype(str).str.strip().ne("").sum())
    active_no_match = len(active) - active_match_success
    closed_match_success = int(closed["Inception Date"].astype(str).str.strip().ne("").sum())
    closed_no_match = len(closed) - closed_match_success

    print("ETF Universe 清洗整合完成")
    print(f"当前 Screener 读取行数: {len(screener_raw)}")
    print(f"Launches 读取行数: {len(launches)}")
    print(f"Closures 读取行数: {len(closures) + closure_duplicate_rows_deleted}")
    print(f"Active 当前产品最终行数: {len(active)}")
    print(f"Closed 历史产品最终行数: {len(closed)}")
    print(f"ETF_Universe_All 总表最终行数: {len(universe)}")
    print(f"Active 产品 Inception Date 匹配成功/未匹配: {active_match_success} / {active_no_match}")
    print(f"Closed 产品 Inception Date 匹配成功/未匹配: {closed_match_success} / {closed_no_match}")
    print(f"Closures Ticker 复用/多清盘日期异常数: {closure_reuse_count}")
    print(f"输出文件保存目录: {OUTPUT_DIR}")
    print(f"Excel: {OUTPUT_XLSX}")
    print(f"CSV: {OUTPUT_CSV}")
    print(f"日志: {OUTPUT_LOG}")


if __name__ == "__main__":
    main()
