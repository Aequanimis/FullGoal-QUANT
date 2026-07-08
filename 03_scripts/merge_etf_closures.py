from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


INPUT_DIR = Path(r"C:\Users\lvdon\Desktop\Fullgoal\课题研究\处理后数据\境外策略ETF数据\原始数据\全量closures")
OUTPUT_DIR = Path(r"C:\Users\lvdon\Desktop\Fullgoal\课题研究\处理后数据\境外策略ETF数据\处理结果")

OUTPUT_XLSX = OUTPUT_DIR / "ETF_Closures_汇总.xlsx"
OUTPUT_CSV = OUTPUT_DIR / "ETF_Closures_汇总.csv"
OUTPUT_LOG = OUTPUT_DIR / "ETF_Closures_汇总日志.txt"

REQUIRED_COLUMNS = ["Closure Date", "Ticker", "Fund Name"]
OUTPUT_COLUMNS = ["Closure Date", "Closure Year", "Ticker", "Fund Name", "Status", "Source_File"]


@dataclass
class FileReadStat:
    file_name: str
    original_rows: int
    valid_rows: int
    date_parse_failed_rows: int
    encoding: str
    read_error: str = ""


def normalize_col_name(value: object) -> str:
    return str(value).replace("\ufeff", "").strip().lower()


def clean_text(value: object) -> str:
    text = "" if pd.isna(value) else str(value)
    text = text.replace("\ufeff", "")
    text = re.sub(r"[\u00a0\u1680\u180e\u2000-\u200a\u2028\u2029\u202f\u205f\u3000]", " ", text)
    text = re.sub(r"[\u200b-\u200d\u2060]", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def is_blank_series(series: pd.Series) -> pd.Series:
    return series.fillna("").astype(str).map(clean_text).eq("")


def read_csv_with_fallback(path: Path) -> tuple[pd.DataFrame, str]:
    encodings = ["utf-8-sig", "utf-8", "gbk", "latin1"]
    last_error: Exception | None = None
    for encoding in encodings:
        try:
            df = pd.read_csv(
                path,
                dtype=str,
                keep_default_na=False,
                encoding=encoding,
                skip_blank_lines=False,
            )
            return df, encoding
        except Exception as exc:  # noqa: BLE001 - vendor CSVs may vary by encoding and shape.
            last_error = exc
    raise RuntimeError(f"无法读取 CSV：{path}；最后错误：{last_error}") from last_error


def drop_empty_columns(df: pd.DataFrame) -> pd.DataFrame:
    kept_columns: list[str] = []
    for column in df.columns:
        column_name = str(column).strip()
        if column_name.startswith("Unnamed:") and is_blank_series(df[column]).all():
            continue
        if is_blank_series(df[column]).all():
            continue
        kept_columns.append(column)
    return df.loc[:, kept_columns].copy()


def select_required_columns(df: pd.DataFrame, file_name: str) -> pd.DataFrame:
    column_lookup = {normalize_col_name(col): col for col in df.columns}
    selected = pd.DataFrame()
    missing: list[str] = []

    for required in REQUIRED_COLUMNS:
        source_col = column_lookup.get(normalize_col_name(required))
        if source_col is None:
            missing.append(required)
        else:
            selected[required] = df[source_col]

    if missing:
        raise ValueError(f"{file_name} 缺少必要字段：{', '.join(missing)}")

    selected["Source_File"] = file_name
    return selected


def parse_closure_dates(values: Iterable[object]) -> pd.Series:
    raw = pd.Series(values, dtype="string").fillna("").str.strip()
    parsed = pd.Series(pd.NaT, index=raw.index, dtype="datetime64[ns]")

    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%Y/%m/%d"):
        remaining = parsed.isna() & raw.ne("")
        if not remaining.any():
            break
        parsed.loc[remaining] = pd.to_datetime(raw.loc[remaining], format=fmt, errors="coerce")

    remaining = parsed.isna() & raw.ne("")
    if remaining.any():
        parsed.loc[remaining] = pd.to_datetime(raw.loc[remaining], errors="coerce")

    return parsed


def build_anomaly_sheet(deduped: pd.DataFrame, ticker_multi_date: list[str], ticker_multi_name: list[str]) -> pd.DataFrame:
    anomaly_frames: list[pd.DataFrame] = []

    def append_records(mask: pd.Series, anomaly_type: str, detail: str) -> None:
        records = deduped.loc[mask, OUTPUT_COLUMNS].copy()
        if records.empty:
            return
        records.insert(0, "Anomaly_Type", anomaly_type)
        records.insert(1, "Detail", detail)
        anomaly_frames.append(records)

    append_records(deduped["_parsed_closure_date"].isna(), "Closure Date 无法解析", "Closure Date is blank or invalid")
    append_records(deduped["Ticker"].eq(""), "Ticker 为空", "Ticker is blank")
    append_records(deduped["Fund Name"].eq(""), "Fund Name 为空", "Fund Name is blank")
    append_records(deduped["Ticker"].isin(ticker_multi_date), "同一 Ticker 对应多个 Closure Date", "Ticker has multiple closure dates")
    append_records(deduped["Ticker"].isin(ticker_multi_name), "同一 Ticker 对应多个 Fund Name", "Ticker has multiple fund names")

    if not anomaly_frames:
        return pd.DataFrame(
            [["无异常", "", "", "", "", "", "", ""]],
            columns=["Anomaly_Type", "Detail", *OUTPUT_COLUMNS],
        )

    return pd.concat(anomaly_frames, ignore_index=True).drop_duplicates(
        subset=["Anomaly_Type", "Closure Date", "Ticker", "Fund Name", "Source_File"],
        keep="first",
    )


def format_multi_value_lines(df: pd.DataFrame, value_col: str, tickers: list[str]) -> list[str]:
    if not tickers:
        return ["无"]

    lines: list[str] = []
    for ticker in sorted(tickers):
        group = df.loc[df["Ticker"].eq(ticker)].copy()
        values = ", ".join(sorted(value for value in group[value_col].dropna().astype(str).unique() if value))
        source_files = ", ".join(sorted(group["Source_File"].dropna().astype(str).unique()))
        lines.append(f"- {ticker}: {value_col}=[{values}], rows={len(group)}, source_files=[{source_files}]")
    return lines


def autosize_workbook(xlsx_path: Path) -> None:
    workbook = load_workbook(xlsx_path)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    for worksheet in workbook.worksheets:
        if worksheet.max_row >= 1:
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            column_letter = get_column_letter(column_cells[0].column)
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            worksheet.column_dimensions[column_letter].width = min(max(max_len + 2, 12), 65)

    workbook.save(xlsx_path)


def main() -> None:
    if not INPUT_DIR.exists():
        raise FileNotFoundError(f"输入目录不存在：{INPUT_DIR}")

    csv_files = sorted(path for path in INPUT_DIR.iterdir() if path.is_file() and path.suffix.lower() == ".csv")
    if not csv_files:
        raise FileNotFoundError(f"输入目录下未找到 CSV 文件：{INPUT_DIR}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    file_stats: list[FileReadStat] = []
    frames: list[pd.DataFrame] = []

    for path in csv_files:
        try:
            raw_df, encoding = read_csv_with_fallback(path)
            original_rows = len(raw_df)
            cleaned_df = drop_empty_columns(raw_df)
            selected_df = select_required_columns(cleaned_df, path.name)

            for column in REQUIRED_COLUMNS:
                selected_df[column] = selected_df[column].map(clean_text)

            empty_mask = selected_df[REQUIRED_COLUMNS].apply(lambda col: col.eq("")).all(axis=1)
            header_mask = (
                selected_df["Closure Date"].str.lower().eq("closure date")
                & selected_df["Ticker"].str.lower().eq("ticker")
                & selected_df["Fund Name"].str.lower().eq("fund name")
            )
            valid_df = selected_df.loc[~empty_mask & ~header_mask].copy()
            parsed_dates = parse_closure_dates(valid_df["Closure Date"])

            file_stats.append(
                FileReadStat(
                    file_name=path.name,
                    original_rows=original_rows,
                    valid_rows=len(valid_df),
                    date_parse_failed_rows=int(parsed_dates.isna().sum()),
                    encoding=encoding,
                )
            )
            frames.append(selected_df)
        except Exception as exc:  # noqa: BLE001 - log bad files and continue.
            file_stats.append(
                FileReadStat(
                    file_name=path.name,
                    original_rows=0,
                    valid_rows=0,
                    date_parse_failed_rows=0,
                    encoding="",
                    read_error=str(exc),
                )
            )

    if not frames:
        raise RuntimeError("所有 CSV 均读取失败，未生成可合并数据。")

    combined = pd.concat(frames, ignore_index=True)
    merged_before_rows = len(combined)

    for column in REQUIRED_COLUMNS + ["Source_File"]:
        combined[column] = combined[column].map(clean_text)

    empty_mask = combined[REQUIRED_COLUMNS].apply(lambda col: col.eq("")).all(axis=1)
    empty_rows_deleted = int(empty_mask.sum())
    combined = combined.loc[~empty_mask].copy()

    repeated_header_mask = (
        combined["Closure Date"].str.lower().eq("closure date")
        & combined["Ticker"].str.lower().eq("ticker")
        & combined["Fund Name"].str.lower().eq("fund name")
    )
    repeated_header_rows_deleted = int(repeated_header_mask.sum())
    combined = combined.loc[~repeated_header_mask].copy()

    combined["Ticker"] = combined["Ticker"].str.upper()
    combined["_raw_closure_date"] = combined["Closure Date"]
    combined["_parsed_closure_date"] = parse_closure_dates(combined["_raw_closure_date"])
    combined["Closure Date"] = combined["_parsed_closure_date"].dt.strftime("%Y-%m-%d").fillna("")
    combined["Closure Year"] = combined["_parsed_closure_date"].dt.year.astype("Int64")
    combined["Status"] = "Closed"

    date_unparseable_rows = int(combined["_parsed_closure_date"].isna().sum())
    ticker_blank_rows = int(combined["Ticker"].eq("").sum())
    fund_name_blank_rows = int(combined["Fund Name"].eq("").sum())

    before_dedup_rows = len(combined)
    deduped = combined.drop_duplicates(subset=REQUIRED_COLUMNS, keep="first").copy()
    duplicate_rows_deleted = before_dedup_rows - len(deduped)

    ticker_multi_date_counts = (
        deduped.loc[deduped["Ticker"].ne("") & deduped["Closure Date"].ne("")]
        .groupby("Ticker")["Closure Date"]
        .nunique()
    )
    ticker_multi_date = ticker_multi_date_counts[ticker_multi_date_counts > 1].index.tolist()

    ticker_multi_name_counts = (
        deduped.loc[deduped["Ticker"].ne("") & deduped["Fund Name"].ne("")]
        .groupby("Ticker")["Fund Name"]
        .nunique()
    )
    ticker_multi_name = ticker_multi_name_counts[ticker_multi_name_counts > 1].index.tolist()

    final_df = deduped[OUTPUT_COLUMNS].sort_values(
        ["Closure Date", "Ticker", "Fund Name", "Source_File"],
        kind="stable",
    )
    file_summary_df = pd.DataFrame([stat.__dict__ for stat in file_stats])[
        ["file_name", "original_rows", "valid_rows", "date_parse_failed_rows", "encoding", "read_error"]
    ]
    anomaly_df = build_anomaly_sheet(deduped, ticker_multi_date, ticker_multi_name)

    final_df.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        final_df.to_excel(writer, sheet_name="ETF_Closures_All", index=False)
        file_summary_df.to_excel(writer, sheet_name="File_Summary", index=False)
        anomaly_df.to_excel(writer, sheet_name="Anomaly_Check", index=False)
    autosize_workbook(OUTPUT_XLSX)

    valid_dates = deduped["_parsed_closure_date"].dropna()
    earliest_date = valid_dates.min().strftime("%Y-%m-%d") if not valid_dates.empty else "无可解析日期"
    latest_date = valid_dates.max().strftime("%Y-%m-%d") if not valid_dates.empty else "无可解析日期"

    year_counts = (
        final_df.loc[final_df["Closure Year"].notna(), "Closure Year"]
        .astype(int)
        .value_counts()
        .sort_index()
    )
    year_count_lines = [f"- {year}: {count}" for year, count in year_counts.items()] or ["无"]

    read_error_lines = [
        f"- {stat.file_name}: {stat.read_error}" for stat in file_stats if stat.read_error
    ] or ["无"]

    log_lines: list[str] = [
        "ETF Closures 汇总日志",
        "=" * 40,
        f"输入目录: {INPUT_DIR}",
        f"输出目录: {OUTPUT_DIR}",
        f"读取 CSV 文件数: {len(csv_files)}",
        f"读取失败文件数: {sum(1 for stat in file_stats if stat.read_error)}",
        "读取失败明细:",
        *read_error_lines,
        "",
        "每个 CSV 文件原始行数:",
    ]
    for stat in file_stats:
        log_lines.append(
            f"- {stat.file_name}: 原始行数={stat.original_rows}, 有效行数={stat.valid_rows}, "
            f"日期解析失败行数={stat.date_parse_failed_rows}, 编码={stat.encoding}"
            + (f", 读取错误={stat.read_error}" if stat.read_error else "")
        )

    log_lines.extend(
        [
            "",
            f"合并前总行数: {merged_before_rows}",
            f"删除空行数量: {empty_rows_deleted}",
            f"删除重复表头行数量: {repeated_header_rows_deleted}",
            f"删除重复行数量: {duplicate_rows_deleted}",
            f"最终保留行数: {len(final_df)}",
            f"日期无法解析的行数: {date_unparseable_rows}",
            f"Ticker 为空的行数: {ticker_blank_rows}",
            f"Fund Name 为空的行数: {fund_name_blank_rows}",
            f"是否存在同一 Ticker 对应多个 Closure Date 的异常: {'是' if ticker_multi_date else '否'}",
            f"是否存在同一 Ticker 对应多个 Fund Name 的异常: {'是' if ticker_multi_name else '否'}",
            f"最早 Closure Date: {earliest_date}",
            f"最晚 Closure Date: {latest_date}",
            "",
            "每个 Closure Year 对应的清盘ETF数量:",
            *year_count_lines,
            "",
            "同一 Ticker 对应多个 Closure Date 明细:",
            *format_multi_value_lines(deduped, "Closure Date", ticker_multi_date),
            "",
            "同一 Ticker 对应多个 Fund Name 明细:",
            *format_multi_value_lines(deduped, "Fund Name", ticker_multi_name),
            "",
            f"输出 Excel: {OUTPUT_XLSX}",
            f"输出 CSV: {OUTPUT_CSV}",
            f"输出日志: {OUTPUT_LOG}",
        ]
    )
    OUTPUT_LOG.write_text("\n".join(log_lines) + "\n", encoding="utf-8-sig")

    print("ETF Closures 处理完成")
    print(f"成功读取 CSV 文件数: {sum(1 for stat in file_stats if not stat.read_error)} / {len(csv_files)}")
    print(f"最终 ETF Closure 记录数: {len(final_df)}")
    print(f"清盘日期范围: {earliest_date} 至 {latest_date}")
    print("清盘年份分布:")
    for line in year_count_lines:
        print(line)
    print(f"是否存在异常 Ticker: {'是' if ticker_multi_date or ticker_multi_name else '否'}")
    print(f"输出文件保存目录: {OUTPUT_DIR}")
    print(f"Excel: {OUTPUT_XLSX}")
    print(f"CSV: {OUTPUT_CSV}")
    print(f"日志: {OUTPUT_LOG}")


if __name__ == "__main__":
    main()
