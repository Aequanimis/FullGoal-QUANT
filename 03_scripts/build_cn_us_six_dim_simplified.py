from __future__ import annotations

import math
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE_DIR = Path.home() / "Desktop" / "Fullgoal" / "ICI课题" / "课题更正"
US_PRIMARY = BASE_DIR / "美国ETF_TOP10生态评分数据包_20260706.xlsx"
US_FALLBACK = BASE_DIR / "美国ETF_TOP10维度评分提取版_20260706.xlsx"
CN_PRIMARY = BASE_DIR / "中国ETF_TOP10生态评分数据包_补充交易跟踪版.xlsx"
CN_FALLBACK_1 = BASE_DIR / "中国ETF_TOP10维度评分提取版_补充交易跟踪版.xlsx"
CN_FALLBACK_2 = BASE_DIR / "中国ETF_TOP10生态评分数据包.xlsx"
OUTPUT = BASE_DIR / "中美ETF_TOP10六维评分_简化可解释版.xlsx"

OUTPUT_HEADERS = [
    "排名",
    "管理人",
    "产品广度",
    "规模沉淀",
    "费率优势",
    "交易活跃",
    "多元能力",
    "旗舰产品",
    "生态评分",
    "评分口径说明",
]

DIM_WEIGHTS = {
    "产品广度": 0.15,
    "规模沉淀": 0.20,
    "费率优势": 0.15,
    "交易活跃": 0.15,
    "多元能力": 0.20,
    "旗舰产品": 0.15,
}

US_NOTE = (
    "美国TOP10内部标准化评分，采用简化六维指标和战略侧重权重；用于美国市场内部排序及与中国管理人画像映射，"
    "不代表中美绝对实力直接比较。"
)
CN_NOTE = (
    "中国TOP10内部标准化评分，采用简化六维指标和战略侧重权重；交易活跃使用补充交易数据修正，"
    "用于中国市场内部排序及与美国管理人画像映射，不代表中美绝对实力直接比较。"
)

RULE_ROWS = [
    ["产品广度", "15%", "产品线是否铺得全", "ETF产品数量、资产/产品类别覆盖数量", "越多越好、越均衡越好", "产品覆盖是基础，但不鼓励单纯铺产品"],
    ["规模沉淀", "20%", "产品是否真正做大", "ETF总管理规模、大规模产品数量", "越大越好", "规模是ETF生态有效性的核心验证"],
    ["费率优势", "15%", "产品成本是否有竞争力", "AUM加权综合费率", "费率越低越好", "低费率是ETF竞争基础，但不是唯一决定因素"],
    ["交易活跃", "15%", "产品是否被市场真实交易使用", "日均成交额合计、活跃产品数量", "成交越活跃越好", "交易活跃体现ETF工具属性，但与规模和旗舰产品有一定重叠"],
    ["多元能力", "20%", "是否能把不同投资能力转化为ETF产品", "多元产品规模占比、多元产品类型覆盖数量", "占比越高、覆盖越多越好", "体现管理人从产品供给走向能力封装，是本课题关注重点"],
    ["旗舰产品", "15%", "是否有客户容易识别的代表产品", "最大ETF规模、Top3产品规模占比", "代表产品越突出越好", "旗舰产品有助于形成客户认知，但与规模沉淀有一定重叠"],
]

RULE_NOTE = (
    "本评分为自建研究框架，用于ETF管理人生态位比较，不代表官方排名。各代表指标先在各自市场TOP10管理人样本内"
    "标准化至60—100分，再汇总为六维得分。美国和中国使用同一套六维框架和战略侧重权重，但分数是在各自市场"
    "内部标准化得到，因此中美之间主要比较结构、生态位和学习路径，不直接比较绝对实力。"
)


def to_float(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("%", "")
    try:
        return float(text)
    except ValueError:
        return None


def read_rows(wb, sheet_name: str) -> list[dict[str, Any]]:
    ws = wb[sheet_name]
    headers = [str(c.value or "").strip() for c in ws[1]]
    out = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {h: values[i] if i < len(values) else None for i, h in enumerate(headers) if h}
        if any(v not in (None, "") for v in row.values()):
            out.append(row)
    return out


def choose_path(primary: Path, *fallbacks: Path) -> Path:
    for path in (primary, *fallbacks):
        if path.exists():
            return path
    raise FileNotFoundError(f"输入文件不存在：{primary} / {fallbacks}")


def choose_sheet(wb, candidates: list[str]) -> str:
    for sheet in candidates:
        if sheet in wb.sheetnames:
            return sheet
    raise ValueError(f"无法找到候选sheet：{candidates}; 当前sheet={wb.sheetnames}")


def first_existing(row: dict[str, Any], fields: list[str]) -> tuple[str | None, float | None]:
    for field in fields:
        if field in row:
            value = to_float(row.get(field))
            if value is not None:
                return field, value
    return None, None


def safe_sum(row: dict[str, Any], fields: list[str]) -> tuple[str, float | None]:
    values = []
    used = []
    for field in fields:
        if field in row:
            value = to_float(row.get(field))
            if value is not None:
                values.append(value)
                used.append(field)
    if not values:
        return "", None
    return " + ".join(used), sum(values)


def score_values(values: dict[str, float | None], positive: bool, log_transform: bool = False) -> tuple[dict[str, float], list[str]]:
    warnings = []
    transformed: dict[str, float | None] = {}
    for manager, value in values.items():
        if value is None:
            transformed[manager] = None
            warnings.append(f"WARN: {manager} 指标缺失")
        else:
            transformed[manager] = math.log1p(max(value, 0)) if log_transform else value
    valid = [v for v in transformed.values() if v is not None]
    if not valid:
        return {manager: 80.0 for manager in values}, warnings + ["WARN: 该指标全样本缺失，统一记80分"]
    min_v, max_v = min(valid), max(valid)
    out = {}
    for manager, value in transformed.items():
        if value is None:
            out[manager] = 80.0
            continue
        if max_v == min_v:
            score = 80.0
        elif positive:
            score = 60 + 40 * (value - min_v) / (max_v - min_v)
        else:
            score = 60 + 40 * (max_v - value) / (max_v - min_v)
        out[manager] = max(60.0, min(100.0, score))
    return out, warnings


def avg_available(scores: list[float | None]) -> tuple[float, bool]:
    valid = [s for s in scores if s is not None]
    if not valid:
        return 80.0, True
    return sum(valid) / len(valid), False


def build_us_metric_rows(path: Path) -> tuple[list[dict[str, Any]], str, list[str], list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    raw_sheet = choose_sheet(wb, ["04_Issuer_Raw_Metrics", "05_Raw_Metrics_TOP10"])
    rows = read_rows(wb, raw_sheet)
    metric_rows = []
    sources = [
        f"Raw Metrics sheet: {raw_sheet}",
        "产品广度: etf_count + (asset_class_count + category_count)",
        "规模沉淀: total_aum_usd(log1p) + product_count_over_10b",
        "费率优势: weighted_avg_expense_ratio_pct(反向)",
        "交易活跃: total_adv_usd(log1p) + high_liquidity_product_count",
        "多元能力: non_plain_beta_aum_share + (asset_class_count + strategy_count)",
        "旗舰产品: largest_etf_aum_usd(log1p) + top3_aum_share",
    ]
    warnings: list[str] = []
    for row in rows:
        manager = row.get("issuer") or row.get("manager")
        if manager in (None, ""):
            continue
        coverage_src, coverage = safe_sum(row, ["asset_class_count", "category_count"])
        diversified_coverage_src, diversified_coverage = safe_sum(row, ["asset_class_count", "strategy_count"])
        item = {
            "manager": str(manager).strip(),
            "etf_count": first_existing(row, ["etf_count"])[1],
            "coverage_count": coverage,
            "total_aum": first_existing(row, ["total_aum_usd", "total_aum_billion_usd"])[1],
            "large_product_count": first_existing(row, ["product_count_over_10b", "large_product_count"])[1],
            "weighted_fee": first_existing(row, ["weighted_avg_expense_ratio_pct", "weighted_expense_ratio", "weighted_avg_expense_ratio", "expense_ratio"])[1],
            "total_adv": first_existing(row, ["total_adv_usd", "total_avg_daily_dollar_volume", "total_trading_value"])[1],
            "active_product_count": first_existing(row, ["high_liquidity_product_count"])[1],
            "diversified_aum_share": first_existing(row, ["non_plain_beta_aum_share", "non_core_beta_aum_share", "active_aum_share", "smart_beta_factor_aum_share"])[1],
            "diversified_coverage_count": diversified_coverage,
            "largest_etf_aum": first_existing(row, ["largest_etf_aum_usd", "largest_etf_aum_billion_usd"])[1],
            "top3_aum_share": first_existing(row, ["top3_aum_share"])[1],
        }
        if not coverage_src:
            warnings.append(f"WARN: {manager} 产品类别覆盖数量字段缺失")
        if not diversified_coverage_src:
            warnings.append(f"WARN: {manager} 多元覆盖数量字段缺失")
        metric_rows.append(item)
    return metric_rows, raw_sheet, sources, warnings


def build_cn_metric_rows(path: Path) -> tuple[list[dict[str, Any]], str, list[str], list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    raw_sheet = choose_sheet(wb, ["05_Manager_Raw_Metrics_Updated", "05_Raw_Metrics_TOP10_Updated", "04_Manager_Raw_Metrics", "05_Raw_Metrics_TOP10"])
    rows = read_rows(wb, raw_sheet)
    metric_rows = []
    sources = [
        f"Raw Metrics sheet: {raw_sheet}",
        "产品广度: etf_count + (asset_class_count + category_count)",
        "规模沉淀: total_aum_cny(log1p) + product_count_over_50bn",
        "费率优势: weighted_avg_fee_pct(反向)",
        "交易活跃: total_adv_60d_cny(log1p) + high_liquidity_product_count",
        "多元能力: non_plain_beta_aum_share + capability_coverage_count",
        "旗舰产品: largest_etf_aum_cny(log1p) + top3_aum_share",
    ]
    warnings: list[str] = []
    for row in rows:
        manager = row.get("manager") or row.get("issuer")
        if manager in (None, ""):
            continue
        coverage_src, coverage = safe_sum(row, ["asset_class_count", "category_count"])
        item = {
            "manager": str(manager).strip(),
            "etf_count": first_existing(row, ["etf_count"])[1],
            "coverage_count": coverage,
            "total_aum": first_existing(row, ["total_aum_cny", "total_aum_100m_cny"])[1],
            "large_product_count": first_existing(row, ["product_count_over_50bn", "product_count_over_10bn", "large_product_count"])[1],
            "weighted_fee": first_existing(row, ["weighted_avg_fee_pct", "weighted_expense_ratio", "weighted_avg_expense_ratio", "expense_ratio"])[1],
            "total_adv": first_existing(row, ["total_adv_60d_cny", "total_adv_60d_100m_cny", "total_trading_value_cny"])[1],
            "active_product_count": first_existing(row, ["high_liquidity_product_count"])[1],
            "diversified_aum_share": first_existing(row, ["non_plain_beta_aum_share", "non_core_beta_aum_share", "smart_beta_strategy_aum_share", "cross_border_aum_share", "fixed_income_commodity_aum_share", "industry_theme_aum_share"])[1],
            "diversified_coverage_count": first_existing(row, ["capability_coverage_count", "asset_class_count", "category_count"])[1],
            "largest_etf_aum": first_existing(row, ["largest_etf_aum_cny", "largest_etf_aum"])[1],
            "top3_aum_share": first_existing(row, ["top3_aum_share"])[1],
        }
        if not coverage_src:
            warnings.append(f"WARN: {manager} 产品类别覆盖数量字段缺失")
        metric_rows.append(item)
    return metric_rows, raw_sheet, sources, warnings


def compute_scores(metric_rows: list[dict[str, Any]], note: str) -> tuple[list[dict[str, Any]], list[str]]:
    managers = [r["manager"] for r in metric_rows]
    by_field = {field: {r["manager"]: r.get(field) for r in metric_rows} for field in [
        "etf_count",
        "coverage_count",
        "total_aum",
        "large_product_count",
        "weighted_fee",
        "total_adv",
        "active_product_count",
        "diversified_aum_share",
        "diversified_coverage_count",
        "largest_etf_aum",
        "top3_aum_share",
    ]}
    indicator_specs = {
        "etf_count": (True, False),
        "coverage_count": (True, False),
        "total_aum": (True, True),
        "large_product_count": (True, False),
        "weighted_fee": (False, False),
        "total_adv": (True, True),
        "active_product_count": (True, False),
        "diversified_aum_share": (True, False),
        "diversified_coverage_count": (True, False),
        "largest_etf_aum": (True, True),
        "top3_aum_share": (True, False),
    }
    warnings: list[str] = []
    indicator_scores = {}
    for field, (positive, log_transform) in indicator_specs.items():
        scores, field_warnings = score_values(by_field[field], positive=positive, log_transform=log_transform)
        indicator_scores[field] = scores
        warnings.extend([f"WARN: {field}: {w}" for w in field_warnings])

    output = []
    neutral_dims: list[str] = []
    for manager in managers:
        dimensions = {}
        dimensions["产品广度"], neutral = avg_available([indicator_scores["etf_count"][manager], indicator_scores["coverage_count"][manager]])
        if neutral:
            neutral_dims.append(f"{manager}-产品广度")
        dimensions["规模沉淀"], neutral = avg_available([indicator_scores["total_aum"][manager], indicator_scores["large_product_count"][manager]])
        if neutral:
            neutral_dims.append(f"{manager}-规模沉淀")
        dimensions["费率优势"], neutral = avg_available([indicator_scores["weighted_fee"][manager]])
        if neutral:
            neutral_dims.append(f"{manager}-费率优势")
        dimensions["交易活跃"], neutral = avg_available([indicator_scores["total_adv"][manager], indicator_scores["active_product_count"][manager]])
        if neutral:
            neutral_dims.append(f"{manager}-交易活跃")
        dimensions["多元能力"], neutral = avg_available([indicator_scores["diversified_aum_share"][manager], indicator_scores["diversified_coverage_count"][manager]])
        if neutral:
            neutral_dims.append(f"{manager}-多元能力")
        dimensions["旗舰产品"], neutral = avg_available([indicator_scores["largest_etf_aum"][manager], indicator_scores["top3_aum_share"][manager]])
        if neutral:
            neutral_dims.append(f"{manager}-旗舰产品")
        total = sum(dimensions[dim] * weight for dim, weight in DIM_WEIGHTS.items())
        row = {"管理人": manager, **{k: round(v, 1) for k, v in dimensions.items()}, "生态评分": round(total, 1), "评分口径说明": note}
        output.append(row)

    if neutral_dims:
        warnings.append(f"WARN: 以下维度因所有代表指标缺失采用80中性分：{'; '.join(neutral_dims)}")
    output.sort(key=lambda r: r["生态评分"], reverse=True)
    for idx, row in enumerate(output, 1):
        row["排名"] = idx
    return output, warnings


def write_score_sheet(wb: Workbook, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(sheet_name)
    ws.append(OUTPUT_HEADERS)
    for row in rows:
        ws.append([row.get(h) for h in OUTPUT_HEADERS])
    style_sheet(ws, score_sheet=True)


def write_rules_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("评分细则")
    ws.append(["维度", "权重", "这项看什么", "实际计算指标", "评分方向", "权重设置原因"])
    for row in RULE_ROWS:
        ws.append(row)
    start = ws.max_row + 2
    ws.cell(start, 1).value = "说明"
    ws.cell(start, 1).font = Font(bold=True, color="1F4E78")
    ws.cell(start + 1, 1).value = RULE_NOTE
    ws.merge_cells(start_row=start + 1, start_column=1, end_row=start + 1, end_column=6)
    style_sheet(ws, score_sheet=False, filter_ref="A1:F7")
    ws.row_dimensions[start + 1].height = 90


def style_sheet(ws, score_sheet: bool, filter_ref: str | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = filter_ref or ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if score_sheet and 3 <= cell.column <= 9:
                cell.number_format = "0.0"
            if score_sheet and cell.column == 1:
                cell.number_format = "0"
            if score_sheet and cell.column == 9:
                cell.font = Font(bold=True)
    for col_idx in range(1, ws.max_column + 1):
        max_len = len(str(ws.cell(1, col_idx).value or ""))
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row_idx, col_idx).value
            if value is not None:
                max_len = max(max_len, min(len(str(value)), 50))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 50)
    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 24 if row_idx == 1 else 42


def main() -> None:
    us_path = choose_path(US_PRIMARY, US_FALLBACK)
    cn_path = choose_path(CN_PRIMARY, CN_FALLBACK_1, CN_FALLBACK_2)

    us_metrics, us_sheet, us_sources, us_warn = build_us_metric_rows(us_path)
    cn_metrics, cn_sheet, cn_sources, cn_warn = build_cn_metric_rows(cn_path)
    us_rows, us_score_warn = compute_scores(us_metrics, US_NOTE)
    cn_rows, cn_score_warn = compute_scores(cn_metrics, CN_NOTE)
    warnings = us_warn + us_score_warn + cn_warn + cn_score_warn

    wb = Workbook()
    wb.remove(wb.active)
    write_score_sheet(wb, "美国TOP10六维评分", us_rows)
    write_score_sheet(wb, "中国TOP10六维评分", cn_rows)
    write_rules_sheet(wb)
    wb.save(OUTPUT)

    verify = load_workbook(OUTPUT, read_only=True, data_only=True)
    print(f"美国输入文件路径: {us_path}")
    print(f"中国输入文件路径: {cn_path}")
    print(f"输出文件路径: {OUTPUT}")
    print("美国用于计算的关键字段来源:")
    for source in us_sources:
        print(f"- {source}")
    print("中国用于计算的关键字段来源:")
    for source in cn_sources:
        print(f"- {source}")
    print("美国TOP10生态评分排名:")
    for row in us_rows:
        print(f"- {row['排名']}. {row['管理人']}: {row['生态评分']:.1f}")
    print("中国TOP10生态评分排名:")
    for row in cn_rows:
        print(f"- {row['排名']}. {row['管理人']}: {row['生态评分']:.1f}")
    print(f"是否出现缺失字段WARN: {'是' if warnings else '否'}")
    for warning in warnings:
        print(warning)
    print(f"输出sheet列表: {verify.sheetnames}")


if __name__ == "__main__":
    main()
