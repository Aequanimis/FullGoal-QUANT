from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


WORK_DIR = Path(r"C:\Users\lvdon\Desktop\Fullgoal\日常工作\7.8\Morning Star")
INPUT_XLSX = WORK_DIR / "美国ETF2010-2026_合并版.xlsx"
OUTPUT_XLSX = WORK_DIR / "美国ETF2010-2026_合并版_中文表头.xlsx"
OUTPUT_LOG = WORK_DIR / "美国ETF2010-2026_中文表头日志.txt"

SOURCE_SHEET = "美国ETF2010-2026_合并"
OUTPUT_SHEET = "美国ETF2010-2026_中文表头"
SUMMARY_SHEET = "合并说明"
EXPECTED_COLS = 396


HEADER_MAP: dict[str, str] = {
    "Name": "基金名称",
    "Investment Type": "投资类型",
    "Base Currency": "基础货币",
    "Virtual Class": "虚拟份额类别",
    "Ticker": "交易代码",
    "ISIN": "ISIN代码",
    "Morningstar Category Broad Group": "晨星大类分组",
    "Morningstar Category Group": "晨星类别分组",
    "Morningstar Category": "晨星分类",
    "Global Category": "全球分类",
    "Investment Area": "投资区域",
    "Morningstar Rating Overall": "晨星综合评级",
    "US Category Group": "美国类别分组",
    "Broad Theme": "一级主题",
    "Theme": "主题",
    "Sub-Theme": "子主题",
    "Domicile": "注册地",
    "Region of Sale": "销售区域",
    "Fund Legal Structure": "基金法律结构",
    "Equity Style Box (Long)": "股票风格箱（多头）",
    "Fixed Income Surveyed Style Box": "固定收益调查风格箱",
    "Morningstar Category Group Id": "晨星类别分组ID",
    "Active/Passive Strategy Type": "主动/被动策略类型",
    "Morningstar Medalist Rating Disclosures Type": "晨星奖牌评级披露类型",
    "Morningstar Medalist Rating": "晨星奖牌评级",
    "Morningstar Medalist Rating Parent Pillar": "晨星奖牌评级-母公司支柱",
    "Morningstar Medalist Rating People Pillar": "晨星奖牌评级-团队支柱",
    "Morningstar Medalist Rating Process Pillar": "晨星奖牌评级-流程支柱",
    "Morningstar Medalist Rating Price Pillar": "晨星奖牌评级-价格支柱",
    "Morningstar Medalist Rating Performance Pillar": "晨星奖牌评级-业绩支柱",
    "Morningstar Index": "晨星指数",
    "Index Family Id": "指数系列ID",
    "Index Family Name": "指数系列名称",
    "Primary Prospectus Benchmark": "招募说明书主要基准",
    "Primary Prospectus Benchmark Id": "招募说明书主要基准ID",
    "Primary Prospectus Benchmark Inception Date": "招募说明书主要基准成立日期",
    "S&P Dow Jones Benchmark": "标普道琼斯基准",
    "FTSE/Russell Benchmark": "富时/罗素基准",
    "Firm Name": "基金公司名称",
    "Branding Name": "品牌名称",
    "Inception Date": "成立日期",
    "Manager Name": "基金经理姓名",
    "Manager History": "基金经理历史",
    "Manager Ownership Level": "基金经理持有水平",
    "Manager Tenure (Longest)": "基金经理任期（最长）",
    "Manager Tenure (Average)": "基金经理任期（平均）",
    "Strategic Beta": "策略贝塔",
    "Strategic Beta Group": "策略贝塔分组",
    "Index Selection": "指数选样方式",
    "Index Weighting": "指数加权方式",
    "Net Assets Date": "净资产日期",
    "Net Assets - Share Class Base Currency": "份额类别净资产（基础货币）",
    "Fund Size Date": "基金规模日期",
    "Fund Size Base Currency": "基金规模（基础货币）",
    "Shares Outstanding Date": "流通份额日期",
    "Shares Outstanding": "流通份额",
    "Holding of an Investment": "单项投资持仓",
    "Oldest Share Class": "最早份额类别",
    "Distribution Status": "分红状态",
    "12 Mo Yield Date": "12个月收益率日期",
    "12 Mo Yield": "12个月收益率",
    "Yield 12 Month At Market Price": "按市价计算的12个月收益率",
    "Yield 12 Month At Market Price End Date": "按市价计算的12个月收益率截止日期",
    "Latest Dividend Date": "最新分红日期",
    "Latest Dividend Base Currency": "最新分红（基础货币）",
    "Latest Dividend NAV Base Currency": "最新分红NAV（基础货币）",
    "Unsplit Dividend amount history - Korea": "未拆分分红金额历史-韩国",
    "Share Split": "份额拆分",
    "Latest Capital Gain Date": "最新资本利得日期",
    "Latest Capital Gain Base Currency": "最新资本利得（基础货币）",
    "Latest Capital Gain NAV Base Currency": "最新资本利得NAV（基础货币）",
    "Latest Capital Gain Date (Long Term)": "最新长期资本利得日期",
    "Latest Capital Gain (Long Term) Base Currency": "最新长期资本利得（基础货币）",
    "Latest Capital Gain Date (Middle Term)": "最新中期资本利得日期",
    "Latest Capital Gain (Middle Term) Base Currency": "最新中期资本利得（基础货币）",
    "Latest Capital Gain Date (Short Term)": "最新短期资本利得日期",
    "Latest Capital Gain (Short Term) Base Currency": "最新短期资本利得（基础货币）",
    "Latest ROC Date": "最新资本返还日期",
    "Latest ROC Base Currency": "最新资本返还（基础货币）",
    "Portfolio Date": "组合日期",
    "Return Date (Daily)": "收益日期（日频）",
    "NAV (Daily) Base Currency": "单位净值（日频，基础货币）",
    "Total Ret 1 Day (Daily) Base Currency": "1日总回报（日频，基础货币）",
    "Total Ret YTD (Daily) Base Currency": "年初至今总回报（日频，基础货币）",
    "Total Ret 1 Yr (Daily) Base Currency": "近1年总回报（日频，基础货币）",
    "Total Ret Annlzd 2 Yr (Daily) Base Currency": "近2年年化总回报（日频，基础货币）",
    "Total Ret Annlzd 3 Yr (Daily) Base Currency": "近3年年化总回报（日频，基础货币）",
    "Total Ret Annlzd 5 Yr (Daily) Base Currency": "近5年年化总回报（日频，基础货币）",
    "Management Fee": "管理费",
    "Annual Report Net Expense Ratio": "年报净费用率",
    "Annual Report Ongoing Charge": "年报持续费用",
    "Annual Report Ongoing Charge Date": "年报持续费用日期",
    "Annual Report Gross Expense Ratio": "年报总费用率",
    "Prospectus Net Expense Ratio": "招募说明书净费用率",
    "Prospectus Gross Expense Ratio": "招募说明书总费用率",
    "Prospectus Objective": "招募说明书投资目标",
    "Turnover Ratio % Date": "换手率日期",
    "Turnover Ratio %": "换手率（%）",
    "Semi-Annual Report Turnover Ratio %": "半年报换手率（%）",
    "Semi-Annual Report Turnover Ratio % Date": "半年报换手率日期",
    "Other Sources Turnover Ratio %": "其他来源换手率（%）",
    "Other Sources Turnover Ratio % Date": "其他来源换手率日期",
    "Share Class Type": "份额类别类型",
    "SecId": "证券ID",
    "PerformanceId": "业绩ID",
    "FundId": "基金ID",
    "CUSIP": "CUSIP代码",
    "Performance Data Ready": "业绩数据是否就绪",
    "Price Data Ready": "价格数据是否就绪",
    "Operations Data Ready": "运营数据是否就绪",
    "Portfolio Data Ready": "组合数据是否就绪",
    "Restructure Date": "重组日期",
    "Morningstar Category Start Date": "晨星分类起始日期",
    "Note Effective Date": "备注生效日期",
    "Primary Share in Market": "市场主要份额",
    "Morningstar Page": "晨星页面",
    "Target Date Report Date": "目标日期报告日期",
    "Defined Contribution": "固定缴款计划",
    "Public": "是否公开发行",
    "SMA": "是否SMA",
    "Currency Hedged": "是否汇率对冲",
    "Currency Hedged To": "对冲至货币",
    "Hedged Proportion": "对冲比例",
    "Duration Hedged": "是否久期对冲",
    "Scope Start Date": "覆盖范围起始日期",
    "Scope End Date": "覆盖范围结束日期",
    "Scope": "覆盖范围",
    "Strategy Name": "策略名称",
    "StrategyId": "策略ID",
    "Strategy Provider Company": "策略提供方公司",
    "External Management Company": "外部管理公司",
    "Track Record Extension": "业绩记录延展",
    "Track Record Inherited From": "业绩记录继承来源",
    "Track Record Extension Type": "业绩记录延展类型",
    "Tax Saving Fund": "节税基金",
    "Primary Share": "主要份额",
    "IRS Tax Classification": "美国国税局税务分类",
    "Prospectus Deferred Tax Expense": "招募说明书递延税费",
    "Construct Date": "构建日期",
    "Market Price (Daily) Base Currency": "市场价格（日频，基础货币）",
    "Prem/Disc 12 Mo Avg": "12个月平均溢价/折价",
    "Avg Daily Volume (3 Mo)": "日均成交量（3个月）",
    "Exchange": "交易所",
    "Exchange Country": "交易所所在国家",
    "ETF Qualitative Rating - Australasia": "ETF定性评级-澳大拉西亚",
    "ETF Qualitative Rating Date": "ETF定性评级日期",
    "Ranking Eligibility": "排名资格",
    "Latest Accounting Date": "最新会计日期",
    "Latest Gross Income": "最新总收入",
    "Capital Gain and Loss on securities": "证券资本利得及损失",
    "Futures Capital Gain and Loss": "期货资本利得及损失",
    "Profit and Loss": "损益",
    "Stock Futures Ratio": "股票期货比例",
    "Ratio in Underlying Funds": "底层基金占比",
    "Percent of Equity in Underlying Funds": "底层基金中权益占比",
    "Morningstar UCITs Page Date": "晨星UCITS页面日期",
    "Prospectus Adjusted Expense Ratio": "招募说明书调整后费用率",
    "Managed Investment Report Availability": "管理型投资报告可得性",
    "Managed Investment Report Analysis Date": "管理型投资报告分析日期",
    "Date: NAV 52 Wk High": "52周最高净值日期",
    "NAV 52 Wk High": "52周最高净值",
    "Date: NAV 52 Wk Low": "52周最低净值日期",
    "NAV 52 Wk Low": "52周最低净值",
    "Transaction Date": "交易日期",
    "Transaction Entry Price": "交易进入价格",
    "Transaction Exit Price": "交易退出价格",
    "Annual Report Adjusted Expense Ratio": "年报调整后费用率",
    "Most Recent Category Peer Group Name": "最新类别同业组名称",
    "Dividend Expense": "分红费用",
    "K1 Issued": "是否出具K-1表",
    "Target Maturity Bond": "目标到期债券",
    "Target Maturity Bond Year": "目标到期债券年份",
    "Currency Tracking Detail": "货币跟踪明细",
    "Listing Date": "上市日期",
    "Delisting Date": "退市日期",
    "Maximum AUM Redemption Percent": "最大AUM赎回比例",
    "Redemption Months": "赎回月份",
    "Redemption Conditions": "赎回条件",
    "Is India GIFT City Fund": "是否印度GIFT城基金",
}


def translate_header(header: Any) -> tuple[str, bool]:
    text = "" if header is None else str(header).strip()
    if text in HEADER_MAP:
        return HEADER_MAP[text], True

    m = re.fullmatch(r"Fund Size - aggr from share classes \(Monthly\) (\d{4}-\d{2}) HKD", text)
    if m:
        return f"基金规模-份额类别汇总（月度）{m.group(1)}（HKD）", True

    m = re.fullmatch(r"Net Assets - share class \(Monthly\) (\d{4}-\d{2}) Base Currency", text)
    if m:
        return f"份额类别净资产（月度）{m.group(1)}（基础货币）", True

    return text, False


def main() -> None:
    if not INPUT_XLSX.exists():
        raise FileNotFoundError(f"输入文件不存在: {INPUT_XLSX}")

    processed_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    shutil.copy2(INPUT_XLSX, OUTPUT_XLSX)

    wb = load_workbook(OUTPUT_XLSX)
    if SOURCE_SHEET not in wb.sheetnames:
        raise ValueError(f"输入文件缺少Sheet: {SOURCE_SHEET}")
    ws = wb[SOURCE_SHEET]

    original_headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    translated_headers: list[str] = []
    untranslated: list[tuple[int, str]] = []
    for col_idx, original in enumerate(original_headers, start=1):
        translated, ok = translate_header(original)
        translated_headers.append(translated)
        ws.cell(1, col_idx).value = translated
        ws.cell(1, col_idx).font = Font(bold=True)
        ws.cell(1, col_idx).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if not ok:
            untranslated.append((col_idx, "" if original is None else str(original)))

    ws.title = OUTPUT_SHEET
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, header in enumerate(translated_headers, start=1):
        width = min(max(len(str(header)) + 2, 10), 28)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(OUTPUT_XLSX)

    check_wb = load_workbook(OUTPUT_XLSX, read_only=True, data_only=True)
    check_ws = check_wb[OUTPUT_SHEET]
    output_headers = [check_ws.cell(1, col).value for col in range(1, check_ws.max_column + 1)]
    translated_count = len(original_headers) - len(untranslated)

    log_lines = [
        "Morningstar 美国ETF合并版中文表头处理日志",
        "=" * 50,
        f"输入文件路径: {INPUT_XLSX}",
        f"输出文件路径: {OUTPUT_XLSX}",
        f"处理完成时间: {processed_at}",
        f"原始主Sheet: {SOURCE_SHEET}",
        f"输出主Sheet: {OUTPUT_SHEET}",
        f"数据行数: {check_ws.max_row - 1}",
        f"列数: {check_ws.max_column}",
        f"预期列数是否为396: {'是' if check_ws.max_column == EXPECTED_COLS else '否'}",
        f"已翻译表头数量: {translated_count}",
        f"未匹配表头数量: {len(untranslated)}",
        "未匹配表头明细:",
        *[f"- 列{idx}: {name}" for idx, name in untranslated[:200]],
        f"输出Sheet列表: {', '.join(check_wb.sheetnames)}",
        f"表头是否已写入中文: {'是' if output_headers == translated_headers else '否'}",
    ]
    OUTPUT_LOG.write_text("\n".join(log_lines) + "\n", encoding="utf-8-sig")

    print("中文表头版生成完成")
    print(f"输入文件: {INPUT_XLSX}")
    print(f"输出文件: {OUTPUT_XLSX}")
    print(f"数据行数: {check_ws.max_row - 1}")
    print(f"列数: {check_ws.max_column}")
    print(f"已翻译表头数量: {translated_count}")
    print(f"未匹配表头数量: {len(untranslated)}")
    print(f"日志: {OUTPUT_LOG}")


if __name__ == "__main__":
    main()
