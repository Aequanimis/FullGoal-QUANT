from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


INPUT_DIR = Path(r"C:\Users\lvdon\Desktop\Fullgoal\课题研究\处理后数据\境外策略ETF数据\原始数据\全量launcher")
OUTPUT_DIR = Path(r"C:\Users\lvdon\Desktop\Fullgoal\课题研究\处理后数据\境外策略ETF数据\处理结果")

OUTPUT_XLSX = OUTPUT_DIR / "ETF_Launches_汇总.xlsx"
OUTPUT_CSV = OUTPUT_DIR / "ETF_Launches_汇总.csv"
OUTPUT_LOG = OUTPUT_DIR / "ETF_Launches_汇总日志.txt"

REQUIRED_COLUMNS = ["Inception Date", "Ticker", "Fund Name"]
OUTPUT_COLUMNS = ["Inception Date", "Inception Year", "Ticker", "Fund Name", "Source_File"]


@dataclass
class FileReadStat:
    file_name: str
    original_rows: int
    rows_after_required_column_selection: int
    encoding: str


def normalize_col_name(value: object) -> str:
    return str(value).replace("\ufeff", "").strip().lower()


def is_blank_series(series: pd.Series) -> pd.Series:
    return series.fillna("").astype(str).str.strip().eq("")


def read_csv_with_fallback(path: Path) -> tuple[pd.DataFrame, str]:
    encodings = ["utf-8-sig", "utf-8", "gb18030", "latin1"]
    last_error: Exception | None = None
    for encoding in encodings:
        try:
            df = pd.read_csv(
                path,
                dtype=str,
                keep_default_na=False,
                encoding=encoding,
                skip_blank_lines=False,
            )
            return df, encoding
        except Exception as exc:  # noqa: BLE001 - keep fallback robust for vendor CSVs.
            last_error = exc
    raise RuntimeError(f"无法读取 CSV：{path}；最后错误：{last_error}") from last_error


def drop_empty_columns(df: pd.DataFrame) -> pd.DataFrame:
    kept_columns: list[str] = []
    for column in df.columns:
        column_name = str(column).strip()
        if column_name.startswith("Unnamed:") and is_blank_series(df[column]).all():
            continue
        if is_blank_series(df[column]).all():
            continue
        kept_columns.append(column)
    return df.loc[:, kept_columns].copy()


def select_required_columns(df: pd.DataFrame, file_name: str) -> pd.DataFrame:
    column_lookup = {normalize_col_name(col): col for col in df.columns}
    selected = pd.DataFrame()
    missing: list[str] = []

    for required in REQUIRED_COLUMNS:
        source_col = column_lookup.get(normalize_col_name(required))
        if source_col is None:
            missing.append(required)
        else:
            selected[required] = df[source_col]

    if missing:
        raise ValueError(f"{file_name} 缺少必要字段：{', '.join(missing)}")

    selected["Source_File"] = file_name
    return selected


def parse_inception_dates(values: Iterable[object]) -> pd.Series:
    raw = pd.Series(values, dtype="string").fillna("").str.strip()
    parsed = pd.Series(pd.NaT, index=raw.index, dtype="datetime64[ns]")

    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%Y/%m/%d"):
        remaining = parsed.isna() & raw.ne("")
        if not remaining.any():
            break
        parsed.loc[remaining] = pd.to_datetime(raw.loc[remaining], format=fmt, errors="coerce")

    remaining = parsed.isna() & raw.ne("")
    if remaining.any():
        parsed.loc[remaining] = pd.to_datetime(raw.loc[remaining], errors="coerce")

    return parsed


def format_anomaly_lines(anomaly_df: pd.DataFrame) -> list[str]:
    if anomaly_df.empty:
        return ["无"]

    lines: list[str] = []
    grouped = anomaly_df.sort_values(["Ticker", "Inception Date", "Fund Name"]).groupby("Ticker", sort=True)
    for ticker, group in grouped:
        dates = ", ".join(sorted(group["Inception Date"].dropna().astype(str).unique()))
        source_files = ", ".join(sorted(group["Source_File"].dropna().astype(str).unique()))
        lines.append(f"- {ticker}: dates=[{dates}], rows={len(group)}, source_files=[{source_files}]")
    return lines


def autosize_main_sheet(xlsx_path: Path) -> None:
    workbook = load_workbook(xlsx_path)
    worksheet = workbook["ETF_Launches_All"]

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        worksheet.column_dimensions[column_letter].width = min(max(max_len + 2, 12), 55)

    workbook.save(xlsx_path)


def main() -> None:
    if not INPUT_DIR.exists():
        raise FileNotFoundError(f"输入目录不存在：{INPUT_DIR}")

    csv_files = sorted(path for path in INPUT_DIR.iterdir() if path.is_file() and path.suffix.lower() == ".csv")
    if not csv_files:
        raise FileNotFoundError(f"输入目录下未找到 CSV 文件：{INPUT_DIR}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    file_stats: list[FileReadStat] = []
    frames: list[pd.DataFrame] = []

    for path in csv_files:
        raw_df, encoding = read_csv_with_fallback(path)
        original_rows = len(raw_df)
        cleaned_df = drop_empty_columns(raw_df)
        selected_df = select_required_columns(cleaned_df, path.name)
        file_stats.append(
            FileReadStat(
                file_name=path.name,
                original_rows=original_rows,
                rows_after_required_column_selection=len(selected_df),
                encoding=encoding,
            )
        )
        frames.append(selected_df)

    combined = pd.concat(frames, ignore_index=True)
    merged_before_rows = len(combined)

    for column in REQUIRED_COLUMNS + ["Source_File"]:
        combined[column] = combined[column].fillna("").astype(str)

    required_blank_mask = combined[REQUIRED_COLUMNS].apply(lambda col: col.str.strip().eq("")).all(axis=1)
    empty_rows_deleted = int(required_blank_mask.sum())
    combined = combined.loc[~required_blank_mask].copy()

    repeated_header_mask = (
        combined["Inception Date"].str.strip().str.lower().eq("inception date")
        & combined["Ticker"].str.strip().str.lower().eq("ticker")
        & combined["Fund Name"].str.strip().str.lower().eq("fund name")
    )
    repeated_header_rows_deleted = int(repeated_header_mask.sum())
    combined = combined.loc[~repeated_header_mask].copy()

    combined["Ticker"] = combined["Ticker"].str.strip().str.upper()
    combined["Fund Name"] = combined["Fund Name"].str.strip()
    combined["_raw_inception_date"] = combined["Inception Date"].str.strip()

    parsed_dates = parse_inception_dates(combined["_raw_inception_date"])
    date_unparseable_rows = int(parsed_dates.isna().sum())
    combined["_parsed_inception_date"] = parsed_dates
    combined["Inception Date"] = parsed_dates.dt.strftime("%Y-%m-%d").fillna("")
    combined["Inception Year"] = parsed_dates.dt.year.astype("Int64")

    before_dedup_rows = len(combined)
    deduped = combined.drop_duplicates(subset=REQUIRED_COLUMNS, keep="first").copy()
    duplicate_rows_deleted = before_dedup_rows - len(deduped)

    anomaly_keys = (
        deduped.loc[deduped["Ticker"].ne("") & deduped["Inception Date"].ne("")]
        .groupby("Ticker")["Inception Date"]
        .nunique()
    )
    anomaly_tickers = anomaly_keys[anomaly_keys > 1].index.tolist()
    anomaly_df = deduped[deduped["Ticker"].isin(anomaly_tickers)].copy()

    final_df = deduped[OUTPUT_COLUMNS].sort_values(
        ["Inception Date", "Ticker", "Fund Name", "Source_File"],
        kind="stable",
    )

    final_df.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        final_df.to_excel(writer, sheet_name="ETF_Launches_All", index=False)
    autosize_main_sheet(OUTPUT_XLSX)

    valid_dates = deduped["_parsed_inception_date"].dropna()
    earliest_date = valid_dates.min().strftime("%Y-%m-%d") if not valid_dates.empty else "无可解析日期"
    latest_date = valid_dates.max().strftime("%Y-%m-%d") if not valid_dates.empty else "无可解析日期"

    log_lines: list[str] = [
        "ETF Launches 汇总日志",
        "=" * 40,
        f"输入目录: {INPUT_DIR}",
        f"输出目录: {OUTPUT_DIR}",
        f"读取 CSV 文件数: {len(csv_files)}",
        "",
        "每个 CSV 文件原始行数:",
    ]
    for stat in file_stats:
        log_lines.append(
            f"- {stat.file_name}: 原始行数={stat.original_rows}, "
            f"必要字段选择后行数={stat.rows_after_required_column_selection}, 编码={stat.encoding}"
        )

    log_lines.extend(
        [
            "",
            f"合并前总行数: {merged_before_rows}",
            f"删除空行数量: {empty_rows_deleted}",
            f"删除重复表头行数量: {repeated_header_rows_deleted}",
            f"删除重复行数量: {duplicate_rows_deleted}",
            f"最终保留行数: {len(final_df)}",
            f"日期无法解析的行数: {date_unparseable_rows}",
            f"是否存在同一 Ticker 对应多个 Inception Date 的异常情况: {'是' if anomaly_tickers else '否'}",
            f"异常 Ticker 数量: {len(anomaly_tickers)}",
            "异常 Ticker 明细:",
            *format_anomaly_lines(anomaly_df),
            f"最早 Inception Date: {earliest_date}",
            f"最晚 Inception Date: {latest_date}",
            "",
            f"输出 Excel: {OUTPUT_XLSX}",
            f"输出 CSV: {OUTPUT_CSV}",
            f"输出日志: {OUTPUT_LOG}",
        ]
    )
    OUTPUT_LOG.write_text("\n".join(log_lines) + "\n", encoding="utf-8-sig")

    print("ETF Launches 处理完成")
    print(f"成功合并文件数: {len(csv_files)}")
    print(f"最终 ETF Launches 记录数: {len(final_df)}")
    print(f"日期范围: {earliest_date} 至 {latest_date}")
    print(f"输出文件保存目录: {OUTPUT_DIR}")
    print(f"Excel: {OUTPUT_XLSX}")
    print(f"CSV: {OUTPUT_CSV}")
    print(f"日志: {OUTPUT_LOG}")


if __name__ == "__main__":
    main()
