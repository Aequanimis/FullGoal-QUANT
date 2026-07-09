from __future__ import annotations

import csv
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT_DIR = Path(r"C:\Users\lvdon\Desktop\Fullgoal\课题研究\处理后数据\境外策略ETF数据\处理结果")
INPUT_XLSX = OUTPUT_DIR / "境外ETF全量数据_轻量分析版_v1.xlsx"
OUTPUT_XLSX = OUTPUT_DIR / "境外ETF全量数据_轻量分析版_v1_中文表头.xlsx"
OUTPUT_CSV = OUTPUT_DIR / "ETF_Master_Lite_v1_中文表头.csv"
OUTPUT_LOG = OUTPUT_DIR / "境外ETF全量数据_轻量分析版_v1_中文表头_日志.txt"

SHEET_RENAME_MAP = {
    "ETF_Master_Lite": "ETF主表_轻量版",
    "Yearly_Lifecycle_Summary": "年度生命周期统计",
    "Data_Notes": "数据口径说明",
}

MASTER_HEADER_MAP = {
    "Ticker": "ETF代码",
    "Fund Name": "ETF名称",
    "Issuer": "发行人",
    "Brand": "品牌",
    "Status": "产品状态",
    "Data_Type": "数据类型",
    "Inception Date": "成立日期",
    "Inception Year": "成立年份",
    "Closure Date": "清盘日期",
    "Closure Year": "清盘年份",
    "Asset Class": "资产类别",
    "Category": "产品类别",
    "Strategy": "策略类型",
    "Niche": "细分方向",
    "Segments": "市场分段",
    "Selection Criteria": "选样标准",
    "Weighting Scheme": "加权方式",
    "AUM_USD_100M": "资产规模（亿美元）",
    "Expense_Ratio_Pct": "费率（%）",
    "Avg_Daily_Dollar_Vol_USD_100M": "日均成交额（亿美元）",
    "Avg_Daily_Share_Vol": "日均成交份额",
    "Spread_Data_Pct": "买卖价差（%）",
    "Total_Return_1Y_Pct": "近1年总回报（%）",
    "Total_Return_3Y_Pct": "近3年总回报（%）",
    "Total_Return_5Y_Pct": "近5年总回报（%）",
    "Fund_Flow_1Y_USD": "近1年资金流（美元）",
    "Fund_Flow_YTD_USD": "年初至今资金流（美元）",
    "Fund_Flow_1MO_USD": "近1个月资金流（美元）",
    "Dividend_Yield_Pct": "股息率（%）",
    "PE": "市盈率",
    "PB": "市净率",
    "Is_ETN": "是否ETN",
    "Is_Leveraged": "是否杠杆",
    "Leveraged_Ratio_Num": "杠杆倍数",
    "Is_Inverse": "是否反向",
    "Region": "区域",
    "Geography": "地理市场",
    "Economic Development": "市场发展属性",
    "Lifecycle_Match_Flag": "生命周期匹配状态",
}

YEARLY_HEADER_MAP = {
    "Year": "年份",
    "Launch Count": "新发产品数量",
    "Closure Count": "清盘产品数量",
    "Active_Current_Count": "当前存续产品数量",
    "Net Launch Count": "净新发数量",
    "Note": "备注",
}

NOTES_HEADER_MAP = {
    "Item": "项目",
    "Description": "说明",
}

HEADER_MAP_BY_SHEET = {
    "ETF_Master_Lite": MASTER_HEADER_MAP,
    "Yearly_Lifecycle_Summary": YEARLY_HEADER_MAP,
    "Data_Notes": NOTES_HEADER_MAP,
}

DATE_HEADERS_CN = {"成立日期", "清盘日期"}


def cell_to_csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return value


def translate_headers(ws, header_map: dict[str, str]) -> list[str]:
    unmatched: list[str] = []
    for cell in ws[1]:
        original = str(cell.value or "").strip()
        if original in header_map:
            cell.value = header_map[original]
        else:
            unmatched.append(original)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = PatternFill(fill_type=None)
    return unmatched


def format_sheet(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header or ""))
        for cell in list(ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=min(ws.max_row, 500)))[0]:
            if header in DATE_HEADERS_CN:
                cell.number_format = "yyyy-mm-dd"
            max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 55)


def export_master_csv(workbook_path: Path) -> None:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb["ETF主表_轻量版"]
    with OUTPUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        for row in ws.iter_rows(values_only=True):
            writer.writerow([cell_to_csv_value(value) for value in row])


def main() -> None:
    if not INPUT_XLSX.exists():
        raise FileNotFoundError(f"输入文件不存在，请检查路径: {INPUT_XLSX}")

    run_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    shutil.copy2(INPUT_XLSX, OUTPUT_XLSX)

    wb = load_workbook(OUTPUT_XLSX)
    input_sheet_stats: dict[str, tuple[int, int]] = {}
    unmatched_headers: dict[str, list[str]] = {}

    for old_sheet_name, new_sheet_name in SHEET_RENAME_MAP.items():
        if old_sheet_name not in wb.sheetnames:
            unmatched_headers[old_sheet_name] = [f"缺少Sheet: {old_sheet_name}"]
            continue
        ws = wb[old_sheet_name]
        input_sheet_stats[new_sheet_name] = (ws.max_row - 1, ws.max_column)
        unmatched = translate_headers(ws, HEADER_MAP_BY_SHEET[old_sheet_name])
        if unmatched:
            unmatched_headers[new_sheet_name] = unmatched
        ws.title = new_sheet_name
        format_sheet(ws)

    wb.save(OUTPUT_XLSX)
    export_master_csv(OUTPUT_XLSX)

    out_wb = load_workbook(OUTPUT_XLSX, read_only=True, data_only=True)
    expected_sheet_names = list(SHEET_RENAME_MAP.values())
    sheet_names_converted = out_wb.sheetnames == expected_sheet_names
    master_ws = out_wb["ETF主表_轻量版"]
    master_rows = master_ws.max_row - 1
    master_cols = master_ws.max_column

    all_unmatched = {
        sheet: headers for sheet, headers in unmatched_headers.items() if headers
    }
    unmatched_text = "无" if not all_unmatched else "; ".join(
        f"{sheet}: {', '.join(headers)}" for sheet, headers in all_unmatched.items()
    )

    log_lines = [
        "境外 ETF 轻量分析版中文表头处理日志",
        "=" * 45,
        f"输入文件路径: {INPUT_XLSX}",
        f"输出 Excel 路径: {OUTPUT_XLSX}",
        f"输出 CSV 路径: {OUTPUT_CSV}",
        f"输出日志路径: {OUTPUT_LOG}",
        f"运行时间: {run_time}",
        f"Sheet名称是否已转换为中文: {'是' if sheet_names_converted else '否'}",
        "三个 Sheet 的行数和列数:",
    ]
    for sheet_name in expected_sheet_names:
        ws = out_wb[sheet_name]
        log_lines.append(f"- {sheet_name}: 行数={ws.max_row - 1}, 列数={ws.max_column}")
    log_lines.extend(
        [
            f"主表是否仍为 7284 行: {'是' if master_rows == 7284 else '否'}",
            f"主表是否仍为 39 列: {'是' if master_cols == 39 else '否'}",
            f"是否存在未成功翻译的表头: {'否' if unmatched_text == '无' else '是'}",
            f"未匹配字段: {unmatched_text}",
        ]
    )
    OUTPUT_LOG.write_text("\n".join(log_lines) + "\n", encoding="utf-8-sig")

    print("中文表头版生成完成")
    print(f"输入文件路径: {INPUT_XLSX}")
    print(f"输出文件路径: {OUTPUT_XLSX}")
    print("三个 Sheet 的行数和列数:")
    for sheet_name in expected_sheet_names:
        ws = out_wb[sheet_name]
        print(f"- {sheet_name}: 行数={ws.max_row - 1}, 列数={ws.max_column}")
    print(f"主表是否仍为 7284 行: {'是' if master_rows == 7284 else '否'}")
    print(f"主表是否仍为 39 列: {'是' if master_cols == 39 else '否'}")
    print(f"Sheet 名称是否已转换为中文: {'是' if sheet_names_converted else '否'}")
    print(f"是否存在未成功翻译的表头: {'否' if unmatched_text == '无' else '是'}")
    print(f"输出 CSV 路径: {OUTPUT_CSV}")
    print(f"输出日志路径: {OUTPUT_LOG}")


if __name__ == "__main__":
    main()
