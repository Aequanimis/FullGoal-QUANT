from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE_DIR = Path.home() / "Desktop" / "Fullgoal" / "ICI课题" / "课题更正"
US_INPUT = BASE_DIR / "美国ETF_TOP10维度评分提取版_20260706.xlsx"
CN_INPUT_PRIMARY = BASE_DIR / "中国ETF_TOP10维度评分提取版_补充交易跟踪版.xlsx"
CN_INPUT_FALLBACK = BASE_DIR / "中国ETF_TOP10维度评分提取版.xlsx"
OUTPUT = BASE_DIR / "中美ETF_TOP10六维评分_战略侧重版.xlsx"

OUTPUT_HEADERS = [
    "排名",
    "管理人",
    "产品布局",
    "规模沉淀",
    "费率竞争",
    "交易活跃",
    "能力迁移",
    "入口心智",
    "生态总分_战略侧重版",
    "评分口径说明",
]

DIMENSION_MAP = {
    "产品布局": [
        "产品矩阵广度",
        "Product Breadth",
        "产品矩阵广度_score",
    ],
    "规模沉淀": [
        "规模沉淀能力",
        "Scale Accumulation",
        "规模沉淀能力_score",
    ],
    "费率竞争": [
        "成本竞争力",
        "Cost Competitiveness",
        "成本竞争力_score",
    ],
    "交易活跃": [
        "交易流动性",
        "Trading Liquidity",
        "交易流动性_score",
    ],
    "能力迁移": [
        "能力迁移程度",
        "Capability Migration",
        "能力迁移程度_score",
    ],
    "入口心智": [
        "入口产品心智",
        "Flagship Mindshare",
        "入口产品心智_score",
    ],
}

MANAGER_FIELDS = [
    "issuer",
    "issuer_display_name",
    "manager",
    "管理人",
    "发行人",
    "entity",
]

LONG_DIM_FIELDS = ["dimension", "维度"]
LONG_SCORE_FIELDS = ["score", "分数", "评分"]

WEIGHTS = {
    "产品布局": 0.15,
    "规模沉淀": 0.20,
    "费率竞争": 0.15,
    "交易活跃": 0.15,
    "能力迁移": 0.20,
    "入口心智": 0.15,
}

US_PRIORITY = [
    "02_Dimension_Score_Wide",
    "01_TOP10_Score_Ranking",
    "03_Dimension_Score_Long",
]

CN_PRIORITY = [
    "02_Dimension_Score_Wide_Updated",
    "01_TOP10_Score_Ranking_Updated",
    "02_Dimension_Score_Wide",
    "01_TOP10_Score_Ranking",
    "03_Dimension_Score_Long_Updated",
    "03_Dimension_Score_Long",
]

US_NOTE = (
    "美国TOP10内部标准化评分，采用战略侧重权重；用于美国市场内部排序及与中国管理人画像映射，"
    "不代表中美绝对实力直接比较。"
)
CN_NOTE = (
    "中国TOP10内部标准化评分，采用战略侧重权重；交易活跃使用补充交易数据修正，用于中国市场内部排序"
    "及与美国管理人画像映射，不代表中美绝对实力直接比较。"
)


def norm_header(value: Any) -> str:
    return str(value or "").strip()


def to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("%", "")
    try:
        return float(text)
    except ValueError:
        return None


def read_rows(wb, sheet_name: str) -> list[dict[str, Any]]:
    ws = wb[sheet_name]
    headers = [norm_header(c.value) for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {h: values[i] if i < len(values) else None for i, h in enumerate(headers) if h}
        if any(v not in (None, "") for v in row.values()):
            rows.append(row)
    return rows


def find_first(headers: set[str], candidates: list[str]) -> str | None:
    for candidate in candidates:
        if candidate in headers:
            return candidate
    lowered = {h.lower(): h for h in headers}
    for candidate in candidates:
        found = lowered.get(candidate.lower())
        if found:
            return found
    return None


def choose_sheet(wb, priority: list[str]) -> str:
    for sheet in priority:
        if sheet in wb.sheetnames:
            return sheet
    raise ValueError(f"无法在工作簿中识别候选sheet；现有sheet={wb.sheetnames}")


def is_long_sheet(sheet_name: str) -> bool:
    return "Long" in sheet_name or "长表" in sheet_name


def resolve_dimension_name(raw: Any) -> str | None:
    raw_text = norm_header(raw)
    for output_dim, candidates in DIMENSION_MAP.items():
        if raw_text == output_dim or raw_text in candidates:
            return output_dim
    return None


def extract_wide(path: Path, priority: list[str], market: str, note: str) -> tuple[list[dict[str, Any]], str, list[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = choose_sheet(wb, priority)
    rows = read_rows(wb, sheet)
    warnings: list[str] = []
    if not rows:
        raise ValueError(f"{market}读取sheet为空：{sheet}")

    if is_long_sheet(sheet):
        extracted = pivot_long(rows, market, note, sheet, warnings)
    else:
        extracted = extract_from_wide_rows(rows, market, note, sheet, warnings)
    return extracted, sheet, warnings


def extract_from_wide_rows(
    rows: list[dict[str, Any]], market: str, note: str, sheet: str, warnings: list[str]
) -> list[dict[str, Any]]:
    headers = set(rows[0].keys())
    manager_field = find_first(headers, MANAGER_FIELDS)
    if not manager_field:
        raise ValueError(f"{market} {sheet} 无法识别管理人字段")

    dim_fields: dict[str, str] = {}
    for output_dim, candidates in DIMENSION_MAP.items():
        field = find_first(headers, candidates + [output_dim])
        if field:
            dim_fields[output_dim] = field
        else:
            warnings.append(f"WARN: {market} {sheet} 缺失维度字段：{output_dim}")
    missing = [dim for dim in WEIGHTS if dim not in dim_fields]
    if missing:
        raise ValueError(f"{market} {sheet} 缺失必要六维字段：{missing}")

    out = []
    for row in rows:
        manager = row.get(manager_field)
        if manager in (None, ""):
            continue
        item = {"管理人": str(manager).strip(), "评分口径说明": note}
        for dim, field in dim_fields.items():
            value = to_float(row.get(field))
            if value is None:
                warnings.append(f"WARN: {market} {sheet} {manager} 的 {dim} 分数为空")
                continue
            item[dim] = value
        if all(dim in item for dim in WEIGHTS):
            out.append(item)
    return rank_rows(out)


def pivot_long(
    rows: list[dict[str, Any]], market: str, note: str, sheet: str, warnings: list[str]
) -> list[dict[str, Any]]:
    headers = set(rows[0].keys())
    manager_field = find_first(headers, MANAGER_FIELDS)
    dim_field = find_first(headers, LONG_DIM_FIELDS)
    score_field = find_first(headers, LONG_SCORE_FIELDS)
    if not manager_field or not dim_field or not score_field:
        raise ValueError(f"{market} {sheet} 长表缺少 manager/dimension/score 字段")

    by_manager: dict[str, dict[str, Any]] = {}
    for row in rows:
        manager = row.get(manager_field)
        dim = resolve_dimension_name(row.get(dim_field))
        score = to_float(row.get(score_field))
        if manager in (None, "") or not dim:
            continue
        if score is None:
            warnings.append(f"WARN: {market} {sheet} {manager} 的 {dim} 分数为空")
            continue
        item = by_manager.setdefault(str(manager).strip(), {"管理人": str(manager).strip(), "评分口径说明": note})
        item[dim] = score

    out = []
    for manager, item in by_manager.items():
        missing = [dim for dim in WEIGHTS if dim not in item]
        if missing:
            warnings.append(f"WARN: {market} {sheet} {manager} 长表透视后缺失维度：{missing}")
            continue
        out.append(item)
    return rank_rows(out)


def rank_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    for row in rows:
        total = sum(float(row[dim]) * weight for dim, weight in WEIGHTS.items())
        row["生态总分_战略侧重版"] = total
    rows.sort(key=lambda r: r["生态总分_战略侧重版"], reverse=True)
    ranked = []
    for idx, row in enumerate(rows, 1):
        out = {"排名": idx}
        out["管理人"] = row["管理人"]
        for dim in WEIGHTS:
            out[dim] = round(float(row[dim]), 1)
        out["生态总分_战略侧重版"] = round(float(row["生态总分_战略侧重版"]), 1)
        out["评分口径说明"] = row["评分口径说明"]
        ranked.append(out)
    return ranked


def write_sheet(wb: Workbook, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(sheet_name)
    ws.append(OUTPUT_HEADERS)
    for row in rows:
        ws.append([row.get(header) for header in OUTPUT_HEADERS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    total_font = Font(bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=(cell.column == 10))
            if 3 <= cell.column <= 9:
                cell.number_format = "0.0"
            if cell.column == 1:
                cell.number_format = "0"
            if cell.column == 9:
                cell.font = total_font

    for col_idx, header in enumerate(OUTPUT_HEADERS, 1):
        max_len = len(str(header))
        for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
            for value_cell in cell:
                if value_cell.value is not None:
                    max_len = max(max_len, len(str(value_cell.value)))
        width = min(max(max_len + 2, 10), 45)
        if header == "评分口径说明":
            width = 45
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 24
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 42


def main() -> None:
    warnings: list[str] = []
    if not US_INPUT.exists():
        raise FileNotFoundError(f"美国输入文件不存在：{US_INPUT}")
    cn_input = CN_INPUT_PRIMARY if CN_INPUT_PRIMARY.exists() else CN_INPUT_FALLBACK
    if not cn_input.exists():
        raise FileNotFoundError(f"中国输入文件不存在：{CN_INPUT_PRIMARY} 或 {CN_INPUT_FALLBACK}")

    us_rows, us_sheet, us_warnings = extract_wide(US_INPUT, US_PRIORITY, "美国", US_NOTE)
    cn_rows, cn_sheet, cn_warnings = extract_wide(cn_input, CN_PRIORITY, "中国", CN_NOTE)
    warnings.extend(us_warnings)
    warnings.extend(cn_warnings)

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "美国TOP10六维评分", us_rows)
    write_sheet(wb, "中国TOP10六维评分", cn_rows)
    wb.save(OUTPUT)

    verify_wb = load_workbook(OUTPUT, read_only=True, data_only=True)
    print(f"美国输入文件路径: {US_INPUT}")
    print(f"中国输入文件路径: {cn_input}")
    print(f"输出文件路径: {OUTPUT}")
    print(f"美国读取的sheet名称: {us_sheet}")
    print(f"中国读取的sheet名称: {cn_sheet}")
    print("美国TOP10管理人及生态总分_战略侧重版:")
    for row in us_rows:
        print(f"- {row['排名']}. {row['管理人']}: {row['生态总分_战略侧重版']:.1f}")
    print("中国TOP10管理人及生态总分_战略侧重版:")
    for row in cn_rows:
        print(f"- {row['排名']}. {row['管理人']}: {row['生态总分_战略侧重版']:.1f}")
    print(f"输出sheet列表: {verify_wb.sheetnames}")
    if warnings:
        for warning in warnings:
            print(warning)
    else:
        print("WARN: 无字段缺失或无法识别问题")


if __name__ == "__main__":
    main()
