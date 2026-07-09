from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT_DIR = Path(r"C:\Users\lvdon\Desktop\Fullgoal\课题研究\处理后数据\境外策略ETF数据\处理结果")
INPUT_XLSX = OUTPUT_DIR / "境外ETF全量数据_清洗整合_含LaunchesClosures_v1.xlsx"
OUTPUT_XLSX = OUTPUT_DIR / "境外ETF全量数据_轻量分析版_v1.xlsx"
OUTPUT_CSV = OUTPUT_DIR / "ETF_Master_Lite_v1.csv"
OUTPUT_LOG = OUTPUT_DIR / "境外ETF全量数据_轻量分析版_v1_日志.txt"

MASTER_COLUMNS = [
    "Ticker",
    "Fund Name",
    "Issuer",
    "Brand",
    "Status",
    "Data_Type",
    "Inception Date",
    "Inception Year",
    "Closure Date",
    "Closure Year",
    "Asset Class",
    "Category",
    "Strategy",
    "Niche",
    "Segments",
    "Selection Criteria",
    "Weighting Scheme",
    "AUM_USD_100M",
    "Expense_Ratio_Pct",
    "Avg_Daily_Dollar_Vol_USD_100M",
    "Avg_Daily_Share_Vol",
    "Spread_Data_Pct",
    "Total_Return_1Y_Pct",
    "Total_Return_3Y_Pct",
    "Total_Return_5Y_Pct",
    "Fund_Flow_1Y_USD",
    "Fund_Flow_YTD_USD",
    "Fund_Flow_1MO_USD",
    "Dividend_Yield_Pct",
    "PE",
    "PB",
    "Is_ETN",
    "Is_Leveraged",
    "Leveraged_Ratio_Num",
    "Is_Inverse",
    "Region",
    "Geography",
    "Economic Development",
    "Lifecycle_Match_Flag",
]

YEARLY_COLUMNS = [
    "Year",
    "Launch Count",
    "Closure Count",
    "Active_Current_Count",
    "Net Launch Count",
    "Note",
]

NUMERIC_COLUMNS = [
    "Inception Year",
    "Closure Year",
    "AUM_USD_100M",
    "Expense_Ratio_Pct",
    "Avg_Daily_Dollar_Vol_USD_100M",
    "Avg_Daily_Share_Vol",
    "Spread_Data_Pct",
    "Total_Return_1Y_Pct",
    "Total_Return_3Y_Pct",
    "Total_Return_5Y_Pct",
    "Fund_Flow_1Y_USD",
    "Fund_Flow_YTD_USD",
    "Fund_Flow_1MO_USD",
    "Dividend_Yield_Pct",
    "PE",
    "PB",
    "Leveraged_Ratio_Num",
]

DATE_COLUMNS = ["Inception Date", "Closure Date"]

DELETED_HEAVY_SHEETS = [
    "Active_Current_Cleaned",
    "Closed_History_Cleaned",
    "Field_Coverage_Summary",
    "Ticker_Reuse_Anomalies",
    "Launch_Match_Audit",
    "Closure_Match_Audit",
    "Processing_Log",
]


def normalize_date_text(series: pd.Series) -> pd.Series:
    parsed = pd.to_datetime(series, errors="coerce")
    return parsed.dt.strftime("%Y-%m-%d").fillna("")


def ensure_columns(df: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    out = df.copy()
    for col in columns:
        if col not in out.columns:
            out[col] = pd.NA
    return out[columns].copy()


def count_nonblank(series: pd.Series) -> int:
    return int(series.fillna("").astype(str).str.strip().ne("").sum())


def read_optional_sheet(sheet_name: str) -> pd.DataFrame:
    try:
        return pd.read_excel(INPUT_XLSX, sheet_name=sheet_name)
    except ValueError:
        return pd.DataFrame()


def build_data_notes(master: pd.DataFrame) -> pd.DataFrame:
    active = master["Status"].eq("Active")
    closed = master["Status"].eq("Closed")
    active_count = int(active.sum())
    closed_count = int(closed.sum())
    total_count = len(master)
    active_matched = count_nonblank(master.loc[active, "Inception Date"])
    active_unmatched = active_count - active_matched
    closed_matched = count_nonblank(master.loc[closed, "Inception Date"])
    closed_unmatched = closed_count - closed_matched

    anomalies = read_optional_sheet("Ticker_Reuse_Anomalies")
    closure_multi_date = 0
    closure_multi_name = 0
    if not anomalies.empty and "Anomaly_Type" in anomalies.columns:
        closure_multi_date = int(
            anomalies["Anomaly_Type"].eq("Closures 中同一 Ticker 对应多个 Closure Date").sum()
        )
        closure_multi_name = int(
            anomalies["Anomaly_Type"].eq("Closures 中同一 Ticker 对应多个 Fund Name").sum()
        )

    rows = [
        {
            "Item": "数据来源",
            "Description": "主表基于 ETF.com Screener 当前截面数据，并追加 ETF.com Launches 与 ETF.com Closures 数据。",
        },
        {
            "Item": "数据口径",
            "Description": "Active 产品来自当前 Screener 数据；Closed 产品来自 Closures 数据，作为历史清盘产品追加，不用 Ticker 直接覆盖当前产品状态。",
        },
        {
            "Item": "总记录数",
            "Description": f"ETF_Master_Lite 总记录数 {total_count}，Active 数量 {active_count}，Closed 数量 {closed_count}。",
        },
        {
            "Item": "Launches 匹配情况",
            "Description": f"Active 产品中 Inception Date 匹配成功 {active_matched}，未匹配 {active_unmatched}。",
        },
        {
            "Item": "Closures 匹配情况",
            "Description": f"Closed 产品中 Inception Date 匹配成功 {closed_matched}，未匹配 {closed_unmatched}。",
        },
        {
            "Item": "Ticker复用风险",
            "Description": (
                "Closures 中存在同一 Ticker 对应多个 Closure Date 或多个 Fund Name 的情况，"
                f"当前检查到多清盘日期 {closure_multi_date} 个、多 Fund Name {closure_multi_name} 个；"
                "因此 Closed 产品不能简单按 Ticker 去重，也不能用 Ticker 直接覆盖 Active 状态。"
            ),
        },
        {
            "Item": "清盘数据时间范围",
            "Description": "Closures 数据最早从 2016 年开始，因此 2010-2015 年清盘数据可能缺失。",
        },
        {
            "Item": "使用建议",
            "Description": "后续做当前规模、费率、收益、流动性分析时，以 Status=Active 为主；做历史发行/清盘数量分析时，可结合 Inception Year 和 Closure Year。",
        },
    ]
    return pd.DataFrame(rows, columns=["Item", "Description"])


def apply_excel_formatting(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")

    for worksheet in workbook.worksheets:
        if worksheet.max_row >= 1:
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for column_cells in worksheet.columns:
            col_idx = column_cells[0].column
            column_letter = get_column_letter(col_idx)
            header = str(column_cells[0].value or "")
            max_len = 0
            for cell in column_cells[:500]:
                max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
                if header in DATE_COLUMNS:
                    cell.number_format = "yyyy-mm-dd"
                elif header in NUMERIC_COLUMNS or header in YEARLY_COLUMNS:
                    cell.number_format = "0.00" if header not in {"Inception Year", "Closure Year", "Year"} else "0"
            worksheet.column_dimensions[column_letter].width = min(max(max_len + 2, 11), 55)

    workbook.save(path)


def clean_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in out.columns:
        if str(out[col].dtype) == "Int64":
            out[col] = out[col].astype(object).where(out[col].notna(), "")
    return out


def main() -> None:
    if not INPUT_XLSX.exists():
        raise FileNotFoundError(f"输入文件不存在: {INPUT_XLSX}")

    run_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    universe = pd.read_excel(INPUT_XLSX, sheet_name="ETF_Universe_All")
    yearly = pd.read_excel(INPUT_XLSX, sheet_name="Yearly_Lifecycle_Summary")

    master = ensure_columns(universe, MASTER_COLUMNS)
    for col in DATE_COLUMNS:
        master[col] = normalize_date_text(master[col])
    for col in NUMERIC_COLUMNS:
        master[col] = pd.to_numeric(master[col], errors="coerce")

    master["Status"] = master["Status"].fillna("").astype(str).str.strip()
    master["Data_Type"] = master["Data_Type"].fillna("").astype(str).str.strip()
    master = master.loc[master["Status"].isin(["Active", "Closed"])].copy()
    master = master.loc[master["Data_Type"].isin(["Active_Current", "Closed_History"])].copy()

    status_order = master["Status"].map({"Active": 0, "Closed": 1}).fillna(2)
    master = (
        master.assign(_status_order=status_order)
        .sort_values(
            ["_status_order", "AUM_USD_100M", "Ticker"],
            ascending=[True, False, True],
            na_position="last",
            kind="stable",
        )
        .drop(columns=["_status_order"])
        .reset_index(drop=True)
    )

    yearly_lite = ensure_columns(yearly, YEARLY_COLUMNS)
    yearly_lite["Year"] = pd.to_numeric(yearly_lite["Year"], errors="coerce")
    for col in ["Launch Count", "Closure Count", "Active_Current_Count", "Net Launch Count"]:
        yearly_lite[col] = pd.to_numeric(yearly_lite[col], errors="coerce")
    yearly_lite = yearly_lite.sort_values("Year", ascending=True, kind="stable").reset_index(drop=True)

    data_notes = build_data_notes(master)

    master.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        clean_for_excel(master).to_excel(writer, sheet_name="ETF_Master_Lite", index=False)
        clean_for_excel(yearly_lite).to_excel(writer, sheet_name="Yearly_Lifecycle_Summary", index=False)
        data_notes.to_excel(writer, sheet_name="Data_Notes", index=False)
    apply_excel_formatting(OUTPUT_XLSX)

    active_count = int(master["Status"].eq("Active").sum())
    closed_count = int(master["Status"].eq("Closed").sum())
    log_lines = [
        "境外 ETF 轻量分析版生成日志",
        "=" * 40,
        f"输入文件路径: {INPUT_XLSX}",
        f"输出 Excel 路径: {OUTPUT_XLSX}",
        f"输出 CSV 路径: {OUTPUT_CSV}",
        f"输出日志路径: {OUTPUT_LOG}",
        f"ETF_Master_Lite 行数: {len(master)}",
        f"Active 行数: {active_count}",
        f"Closed 行数: {closed_count}",
        f"Yearly_Lifecycle_Summary 行数: {len(yearly_lite)}",
        f"保留字段数量: {len(MASTER_COLUMNS)}",
        "删除的重型审计 Sheet: " + ", ".join(DELETED_HEAVY_SHEETS),
        f"运行时间: {run_time}",
    ]
    OUTPUT_LOG.write_text("\n".join(log_lines) + "\n", encoding="utf-8-sig")

    print("境外 ETF 轻量分析版生成完成")
    print(f"轻量版主表行数: {len(master)}")
    print(f"Active 行数: {active_count}")
    print(f"Closed 行数: {closed_count}")
    print(f"保留字段数量: {len(MASTER_COLUMNS)}")
    print(f"输出 Excel 路径: {OUTPUT_XLSX}")
    print(f"输出 CSV 路径: {OUTPUT_CSV}")
    print(f"输出日志路径: {OUTPUT_LOG}")


if __name__ == "__main__":
    main()
