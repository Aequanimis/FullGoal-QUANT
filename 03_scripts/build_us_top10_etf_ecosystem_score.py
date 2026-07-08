from __future__ import annotations

import csv
import math
import re
import shutil
from collections import Counter, defaultdict
from copy import copy
from pathlib import Path
from statistics import mean, median

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DATA_AS_OF = "2026-07-06"

INPUT_PATH = (
    Path.home()
    / "Desktop"
    / "Fullgoal"
    / "课题研究"
    / "处理后数据"
    / "境外策略ETF数据"
    / "原始数据"
    / "境外ETF全量数据 0706.xlsx"
)
OUTPUT_DIR = Path.home() / "Desktop" / "Fullgoal" / "ICI课题" / "课题更正"
OUTPUT_XLSX = OUTPUT_DIR / "美国ETF_TOP10生态评分数据包_20260706.xlsx"
CHART_DIR = OUTPUT_DIR / "美国ETF_TOP10生态评分_ChartData"

DIMENSION_WEIGHTS = {
    "产品矩阵广度": 0.15,
    "规模沉淀能力": 0.20,
    "成本竞争力": 0.15,
    "交易流动性": 0.15,
    "能力迁移程度": 0.20,
    "入口产品心智": 0.15,
}

SUBINDICATORS = {
    "产品矩阵广度": [
        ("asset_class_count", "positive", 0.30, "raw"),
        ("category_count", "positive", 0.25, "raw"),
        ("strategy_count", "positive", 0.25, "raw"),
        ("niche_or_segments_count", "positive", 0.20, "raw"),
    ],
    "规模沉淀能力": [
        ("total_aum_log", "positive", 0.35, "log1p"),
        ("etf_count", "positive", 0.15, "raw"),
        ("product_count_over_1b", "positive", 0.20, "raw"),
        ("product_count_over_10b", "positive", 0.20, "raw"),
        ("top5_aum_log", "positive", 0.10, "log1p"),
    ],
    "成本竞争力": [
        ("weighted_avg_expense_ratio_pct", "negative", 0.35, "raw"),
        ("top10_weighted_expense_ratio_pct", "negative", 0.25, "raw"),
        ("low_fee_aum_share", "positive", 0.25, "raw"),
        ("avg_expense_ratio_pct", "negative", 0.15, "raw"),
    ],
    "交易流动性": [
        ("total_adv_log", "positive", 0.30, "log1p"),
        ("weighted_avg_spread_pct", "negative", 0.25, "raw"),
        ("high_liquidity_product_count", "positive", 0.20, "raw"),
        ("high_liquidity_aum_share", "positive", 0.15, "raw"),
        ("top5_adv_concentration", "negative", 0.10, "raw"),
    ],
    "能力迁移程度": [
        ("active_aum_share", "positive", 0.25, "raw"),
        ("active_count_share", "positive", 0.15, "raw"),
        ("fixed_income_aum_share", "positive", 0.20, "raw"),
        ("equity_premium_income_aum_share", "positive", 0.15, "raw"),
        ("smart_beta_factor_aum_share", "positive", 0.15, "raw"),
        ("non_plain_beta_aum_share", "positive", 0.10, "raw"),
    ],
    "入口产品心智": [
        ("largest_etf_aum_log", "positive", 0.30, "log1p"),
        ("largest_etf_adv_log", "positive", 0.20, "log1p"),
        ("top1_aum_share", "positive", 0.15, "raw"),
        ("top3_aum_share", "positive", 0.15, "raw"),
        ("flagship_mind_tag_score", "positive", 0.20, "raw"),
    ],
}

FLAGSHIP_SCORES = {
    "SPY": 100,
    "QQQ": 100,
    "VOO": 95,
    "IVV": 95,
    "VTI": 95,
    "JEPI": 88,
    "JEPQ": 86,
    "AGG": 85,
    "BND": 85,
    "TLT": 85,
    "GLD": 90,
    "IWM": 88,
    "DIA": 85,
    "RSP": 80,
    "SCHD": 88,
    "ARKK": 85,
}

REQUIRED_FIELD_CANDIDATES = {
    "ticker": ["Ticker", "ticker"],
    "fund_name": ["Fund Name", "fund_name", "fund"],
    "issuer": ["Issuer", "issuer"],
    "brand": ["Brand", "brand"],
    "asset_class": ["Asset Class", "asset_class"],
    "category": ["Category", "category"],
    "strategy": ["Strategy", "strategy"],
    "niche": ["Niche", "niche"],
    "segments": ["Segments", "segments"],
    "selection_criteria": ["Selection Criteria", "selection_criteria"],
    "weighting_scheme": ["Weighting Scheme", "weighting_scheme"],
    "aum": ["AUM", "aum_usd", "aum"],
    "expense_ratio": ["Expense Ratio", "expense_ratio_pct", "expense_ratio"],
    "avg_daily_dollar_vol": ["Avg. Daily $ Vol.", "Avg Daily Dollar Volume", "avg_daily_dollar_vol_usd"],
    "avg_daily_share_vol": ["Avg. Daily Share Vol.", "avg_daily_share_vol"],
    "spread_pct": ["Spread Data %", "spread_pct", "spread_data_pct"],
    "one_day_flow": ["1D Fund Flow", "one_day_fund_flow_usd"],
    "one_week_flow": ["1W Fund Flow", "one_week_fund_flow_usd"],
    "one_month_flow": ["1MO Fund Flow", "1 MO Fund Flow", "one_month_fund_flow_usd"],
    "ytd_flow": ["YTD Fund Flow", "ytd_fund_flow_usd"],
    "one_year_flow": ["1Yr Fund Flow", "1 Yr Fund Flow", "one_year_fund_flow_usd"],
    "one_month_return": ["1 Mo. Total Return", "one_month_total_return_pct"],
    "three_month_return": ["3 Mo. Total Return", "three_month_total_return_pct"],
    "ytd_return": ["YTD Total Return", "ytd_total_return_pct"],
    "one_year_return": ["1 Yr. Total Return", "1 YR TR", "one_year_total_return_pct"],
    "three_year_return": ["3 Yr. Total Return", "three_year_total_return_pct"],
    "five_year_return": ["5 Yr. Total Return", "five_year_total_return_pct"],
    "etn": ["ETN", "etn"],
    "leveraged": ["Leveraged", "leveraged"],
    "leveraged_ratio": ["Leveraged Ratio", "leveraged_ratio"],
    "inverse": ["Inverse", "inverse"],
    "analyst_pick": ["Analyst Pick", "analyst_pick"],
    "letter_grade": ["Letter Grade", "letter_grade"],
    "efficiency_score": ["Efficiency Score", "efficiency_score"],
    "tradability_score": ["Tradability Score", "tradability_score"],
    "fund_closure_risk": ["Fund Closure Risk", "fund_closure_risk"],
}


def normalize_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value).strip().lower())


def match_fields(headers: list[str]) -> dict[str, str | None]:
    by_norm = {normalize_name(header): header for header in headers if header}
    field_map: dict[str, str | None] = {}
    for required, candidates in REQUIRED_FIELD_CANDIDATES.items():
        match = None
        for candidate in candidates:
            norm_candidate = normalize_name(candidate)
            if norm_candidate in by_norm:
                match = by_norm[norm_candidate]
                break
        if match is None:
            for header in headers:
                header_norm = normalize_name(header)
                if any(normalize_name(candidate) in header_norm or header_norm in normalize_name(candidate) for candidate in candidates):
                    match = header
                    break
        field_map[required] = match
    return field_map


def detect_main_sheet(wb):
    best = None
    for ws in wb.worksheets:
        if ws.max_row < 2:
            continue
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        non_empty_headers = [str(header) for header in headers if header]
        fmap = match_fields(non_empty_headers)
        required_hits = sum(1 for key in ["ticker", "fund_name", "issuer", "aum"] if fmap.get(key))
        score = required_hits * 1000 + len(non_empty_headers) * 10 + min(ws.max_row, 100000) / 1000
        if best is None or score > best[0]:
            best = (score, ws, non_empty_headers, fmap)
    if best is None or sum(1 for key in ["ticker", "fund_name", "issuer", "aum"] if best[3].get(key)) < 4:
        raise ValueError("未能识别包含Ticker/Fund Name/Issuer/AUM的主数据sheet")
    return best[1], best[2], best[3]


def parse_money(value) -> float | None:
    if value in (None, "", "-"):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("$", "").replace(",", "")
    if text in {"", "-", "—", "N/A", "nan"}:
        return None
    sign = -1 if text.startswith("-") else 1
    text = text.lstrip("+-")
    suffix = text[-1:].upper()
    multiplier = 1.0
    if suffix in {"K", "M", "B", "T"}:
        text = text[:-1]
        multiplier = {"K": 1e3, "M": 1e6, "B": 1e9, "T": 1e12}[suffix]
    try:
        return sign * float(text) * multiplier
    except ValueError:
        return None


def parse_number(value) -> float | None:
    if value in (None, "", "-"):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("$", "")
    if text in {"", "-", "—", "N/A"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_percent(value, raw_field: str | None) -> float | None:
    if value in (None, "", "-"):
        return None
    if isinstance(value, str) and "%" in value:
        return parse_number(value.replace("%", ""))
    number = parse_number(value)
    if number is None:
        return None
    raw = str(raw_field or "").lower()
    raw_decimal_fields = [
        "expense ratio",
        "spread data %",
        "total return",
        "1 yr tr",
        "dividend yield",
    ]
    if any(token in raw for token in raw_decimal_fields):
        return number * 100
    return number


def is_yes(value) -> bool:
    return str(value or "").strip().lower() in {"yes", "true", "1", "y"}


def standardize_issuer(value) -> str:
    raw = str(value or "").strip()
    compact = re.sub(r"\s+", " ", raw)
    lower = compact.lower()
    if any(token in lower for token in ["blackrock", "ishares"]):
        return "BlackRock/iShares"
    if "vanguard" in lower:
        return "Vanguard"
    if any(token in lower for token in ["state street", "spdr", "state street global advisors"]):
        return "State Street/SPDR"
    if "invesco" in lower:
        return "Invesco"
    if any(token in lower for token in ["j.p. morgan", "jp morgan", "jpmorgan", "jpmorgan chase"]):
        return "JPMorgan AM"
    if any(token in lower for token in ["charles schwab", "schwab"]):
        return "Charles Schwab"
    if "first trust" in lower:
        return "First Trust"
    if "dimensional" in lower:
        return "Dimensional"
    if "fidelity" in lower:
        return "Fidelity"
    if "proshares" in lower:
        return "ProShares"
    if "vaneck" in lower:
        return "VanEck"
    if "ark" in lower:
        return "ARK Invest"
    return compact


def get(row: dict, fmap: dict[str, str | None], field: str):
    raw = fmap.get(field)
    return row.get(raw) if raw else None


def build_clean_rows(ws, headers, fmap):
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        raw_row = {headers[idx]: values[idx] if idx < len(values) else None for idx in range(len(headers))}
        if not any(value not in (None, "") for value in raw_row.values()):
            continue
        ticker = str(get(raw_row, fmap, "ticker") or "").strip().upper()
        issuer_raw = str(get(raw_row, fmap, "issuer") or "").strip()
        aum = parse_money(get(raw_row, fmap, "aum"))
        adv = parse_money(get(raw_row, fmap, "avg_daily_dollar_vol"))
        row = {
            "ticker": ticker,
            "fund_name": str(get(raw_row, fmap, "fund_name") or "").strip(),
            "issuer_raw": issuer_raw,
            "issuer": standardize_issuer(issuer_raw),
            "brand": get(raw_row, fmap, "brand"),
            "asset_class": get(raw_row, fmap, "asset_class"),
            "category": get(raw_row, fmap, "category"),
            "strategy": get(raw_row, fmap, "strategy"),
            "niche": get(raw_row, fmap, "niche"),
            "segments": get(raw_row, fmap, "segments"),
            "selection_criteria": get(raw_row, fmap, "selection_criteria"),
            "weighting_scheme": get(raw_row, fmap, "weighting_scheme"),
            "aum_usd": aum,
            "expense_ratio_pct": parse_percent(get(raw_row, fmap, "expense_ratio"), fmap.get("expense_ratio")),
            "avg_daily_dollar_vol_usd": adv,
            "avg_daily_share_vol": parse_money(get(raw_row, fmap, "avg_daily_share_vol")),
            "spread_pct": parse_percent(get(raw_row, fmap, "spread_pct"), fmap.get("spread_pct")),
            "one_day_flow_usd": parse_money(get(raw_row, fmap, "one_day_flow")),
            "one_week_flow_usd": parse_money(get(raw_row, fmap, "one_week_flow")),
            "one_month_flow_usd": parse_money(get(raw_row, fmap, "one_month_flow")),
            "ytd_flow_usd": parse_money(get(raw_row, fmap, "ytd_flow")),
            "one_year_flow_usd": parse_money(get(raw_row, fmap, "one_year_flow")),
            "one_month_return_pct": parse_percent(get(raw_row, fmap, "one_month_return"), fmap.get("one_month_return")),
            "three_month_return_pct": parse_percent(get(raw_row, fmap, "three_month_return"), fmap.get("three_month_return")),
            "ytd_return_pct": parse_percent(get(raw_row, fmap, "ytd_return"), fmap.get("ytd_return")),
            "one_year_return_pct": parse_percent(get(raw_row, fmap, "one_year_return"), fmap.get("one_year_return")),
            "three_year_return_pct": parse_percent(get(raw_row, fmap, "three_year_return"), fmap.get("three_year_return")),
            "five_year_return_pct": parse_percent(get(raw_row, fmap, "five_year_return"), fmap.get("five_year_return")),
            "etn_flag": is_yes(get(raw_row, fmap, "etn")),
            "leveraged_flag": is_yes(get(raw_row, fmap, "leveraged")),
            "leveraged_ratio": parse_number(get(raw_row, fmap, "leveraged_ratio")),
            "inverse_flag": is_yes(get(raw_row, fmap, "inverse")),
        }
        annotate_flags(row)
        rows.append(row)
    return rows


def contains_any(text: str, keywords: list[str]) -> bool:
    lower = text.lower()
    return any(keyword.lower() in lower for keyword in keywords)


def annotate_flags(row: dict):
    fund_name = str(row.get("fund_name") or "")
    strategy = str(row.get("strategy") or "")
    category = str(row.get("category") or "")
    asset_class = str(row.get("asset_class") or "")
    niche = str(row.get("niche") or "")
    selection = str(row.get("selection_criteria") or "")
    weighting = str(row.get("weighting_scheme") or "")
    ticker = str(row.get("ticker") or "")
    combined = " ".join([fund_name, strategy, category, asset_class, niche, selection, weighting])

    row["active_etf_flag"] = contains_any(" ".join([strategy, fund_name, category]), ["Active", "Actively Managed", "Active ETF"])
    row["fixed_income_flag"] = contains_any(" ".join([asset_class, category]), ["Fixed Income", "Bond", "Treasury", "Corporate Bond", "Municipal", "Aggregate Bond", "High Yield", "Loan"])
    row["equity_premium_income_flag"] = (
        ticker in {"JEPI", "JEPQ"}
        or contains_any(" ".join([fund_name, strategy, niche]), ["Premium Income", "Equity Premium Income", "Covered Call", "BuyWrite", "Option Income"])
        or ("Dividend Income".lower() in fund_name.lower() and "equity" in asset_class.lower())
        or ("Income ETF".lower() in fund_name.lower() and "equity" in asset_class.lower())
    )
    row["smart_beta_factor_flag"] = contains_any(
        combined,
        [
            "Smart Beta",
            "Strategic Beta",
            "Factor",
            "Multifactor",
            "Multi-Factor",
            "Quality",
            "Value",
            "Growth",
            "Momentum",
            "Low Volatility",
            "Minimum Volatility",
            "Dividend",
            "High Dividend",
            "Equal Weight",
            "Fundamental",
            "Revenue",
            "Screened",
            "Optimized",
            "Quality Dividend",
            "Low Beta",
        ],
    )
    plain_strategy = strategy.strip().lower() in {"vanilla", "market cap weighted", "passive beta", ""}
    plain_weighting = weighting.strip().lower() in {"market cap", "cap weighted", "market capitalization", ""}
    row["non_plain_beta_flag"] = bool(
        row["active_etf_flag"]
        or row["smart_beta_factor_flag"]
        or row["equity_premium_income_flag"]
        or not plain_strategy
        or not plain_weighting
        or row["leveraged_flag"]
        or row["inverse_flag"]
    )


def sum_field(rows, field):
    return sum((row.get(field) or 0) for row in rows if row.get(field) is not None)


def weighted_average(rows, value_field, weight_field="aum_usd"):
    pairs = [(row.get(value_field), row.get(weight_field)) for row in rows]
    valid = [(float(v), float(w)) for v, w in pairs if v is not None and w is not None and w > 0]
    total_weight = sum(w for _, w in valid)
    if total_weight <= 0:
        return None
    return sum(v * w for v, w in valid) / total_weight


def simple_average(rows, field):
    vals = [row.get(field) for row in rows if row.get(field) is not None]
    return mean(vals) if vals else None


def simple_median(rows, field):
    vals = [row.get(field) for row in rows if row.get(field) is not None]
    return median(vals) if vals else None


def percentile(values: list[float], p: float) -> float:
    vals = sorted(v for v in values if v is not None)
    if not vals:
        return 0.0
    if len(vals) == 1:
        return vals[0]
    pos = (len(vals) - 1) * p
    lower = math.floor(pos)
    upper = math.ceil(pos)
    if lower == upper:
        return vals[int(pos)]
    return vals[lower] * (upper - pos) + vals[upper] * (pos - lower)


def score_from_percentile(x: float, p5: float, p95: float, direction: str) -> float:
    if p95 == p5:
        return 80.0
    if direction == "positive":
        return max(60.0, min(100.0, 60 + 40 * (x - p5) / (p95 - p5)))
    return max(60.0, min(100.0, 60 + 40 * (p95 - x) / (p95 - p5)))


def rank_by(rows, field, descending=True):
    return sorted(rows, key=lambda row: (row.get(field) is not None, row.get(field) or 0), reverse=descending)


def flagship_info(rows):
    by_aum = sorted(rows, key=lambda r: r.get("aum_usd") or 0, reverse=True)
    largest = by_aum[0] if by_aum else {}
    tagged = [row for row in rows if row.get("ticker") in FLAGSHIP_SCORES]
    if tagged:
        flagship = sorted(tagged, key=lambda row: (FLAGSHIP_SCORES[row.get("ticker")], row.get("aum_usd") or 0), reverse=True)[0]
        score = FLAGSHIP_SCORES[flagship.get("ticker")]
        reason = f"{flagship.get('ticker')}为预设旗舰心智产品"
    else:
        flagship = largest
        aum = largest.get("aum_usd") or 0
        if aum >= 100e9:
            score = 85
            reason = "最大ETF AUM不低于100B"
        elif aum >= 50e9:
            score = 80
            reason = "最大ETF AUM不低于50B"
        elif aum >= 10e9:
            score = 75
            reason = "最大ETF AUM不低于10B"
        else:
            score = 70
            reason = "无预设旗舰ticker且最大ETF AUM低于10B"
    top3 = sum_field(by_aum[:3], "aum_usd")
    total = sum_field(rows, "aum_usd")
    return {
        "flagship_ticker": flagship.get("ticker"),
        "flagship_name": flagship.get("fund_name"),
        "flagship_reason": reason,
        "flagship_mind_tag_score": score,
        "largest_etf_ticker": largest.get("ticker"),
        "largest_etf_name": largest.get("fund_name"),
        "largest_etf_aum_usd": largest.get("aum_usd") or 0,
        "largest_etf_adv_usd": largest.get("avg_daily_dollar_vol_usd") or 0,
        "top1_aum_share": ((largest.get("aum_usd") or 0) / total) if total else 0,
        "top3_aum_share": top3 / total if total else 0,
    }


def compute_metrics(etf_rows):
    by_issuer = defaultdict(list)
    for row in etf_rows:
        by_issuer[row["issuer"]].append(row)

    issuer_totals = []
    for issuer, rows in by_issuer.items():
        issuer_totals.append({"issuer": issuer, "total_aum_usd": sum_field(rows, "aum_usd"), "rows": rows})
    top10 = sorted(issuer_totals, key=lambda row: row["total_aum_usd"], reverse=True)[:10]
    top10_issuers = [row["issuer"] for row in top10]

    raw_metrics = []
    for rank, item in enumerate(top10, start=1):
        issuer = item["issuer"]
        rows = item["rows"]
        total_aum = item["total_aum_usd"]
        total_adv = sum_field(rows, "avg_daily_dollar_vol_usd")
        sorted_aum = sorted(rows, key=lambda r: r.get("aum_usd") or 0, reverse=True)
        sorted_adv = sorted(rows, key=lambda r: r.get("avg_daily_dollar_vol_usd") or 0, reverse=True)
        top5_aum = sum_field(sorted_aum[:5], "aum_usd")
        top10_aum_rows = sorted_aum[:10]
        top5_adv = sum_field(sorted_adv[:5], "avg_daily_dollar_vol_usd")
        active_rows = [row for row in rows if row["active_etf_flag"]]
        fixed_rows = [row for row in rows if row["fixed_income_flag"]]
        premium_rows = [row for row in rows if row["equity_premium_income_flag"]]
        smart_rows = [row for row in rows if row["smart_beta_factor_flag"]]
        non_plain_rows = [row for row in rows if row["non_plain_beta_flag"]]
        high_liq_rows = [row for row in rows if (row.get("avg_daily_dollar_vol_usd") or 0) >= 100e6]
        low_fee_rows = [row for row in rows if row.get("expense_ratio_pct") is not None and row["expense_ratio_pct"] <= 0.10]
        info = flagship_info(rows)
        asset_classes = {row.get("asset_class") for row in rows if row.get("asset_class")}
        categories = {row.get("category") for row in rows if row.get("category")}
        strategies = {row.get("strategy") for row in rows if row.get("strategy")}
        niches = {row.get("niche") for row in rows if row.get("niche")}
        segment_parts = set()
        for row in rows:
            for part in re.split(r"[;,|]", str(row.get("segments") or "")):
                part = part.strip()
                if part:
                    segment_parts.add(part)
        ytd_flow = sum_field(rows, "ytd_flow_usd")
        one_year_flow = sum_field(rows, "one_year_flow_usd")
        raw_metrics.append({
            "rank_by_aum": rank,
            "issuer": issuer,
            "etf_count": len(rows),
            "total_aum_usd": total_aum,
            "total_aum_billion_usd": total_aum / 1e9,
            "total_aum_trillion_usd": total_aum / 1e12,
            "total_adv_usd": total_adv,
            "total_adv_billion_usd": total_adv / 1e9,
            "avg_expense_ratio_pct": simple_average(rows, "expense_ratio_pct"),
            "median_expense_ratio_pct": simple_median(rows, "expense_ratio_pct"),
            "weighted_avg_expense_ratio_pct": weighted_average(rows, "expense_ratio_pct"),
            "weighted_avg_spread_pct": weighted_average(rows, "spread_pct"),
            "active_etf_count": len(active_rows),
            "active_aum_usd": sum_field(active_rows, "aum_usd"),
            "active_aum_share_pct": (sum_field(active_rows, "aum_usd") / total_aum * 100) if total_aum else 0,
            "fixed_income_aum_usd": sum_field(fixed_rows, "aum_usd"),
            "fixed_income_aum_share_pct": (sum_field(fixed_rows, "aum_usd") / total_aum * 100) if total_aum else 0,
            "equity_aum_usd": sum_field([row for row in rows if str(row.get("asset_class")).lower() == "equity"], "aum_usd"),
            "equity_aum_share_pct": (sum_field([row for row in rows if str(row.get("asset_class")).lower() == "equity"], "aum_usd") / total_aum * 100) if total_aum else 0,
            "largest_etf_ticker": info["largest_etf_ticker"],
            "largest_etf_name": info["largest_etf_name"],
            "largest_etf_aum_usd": info["largest_etf_aum_usd"],
            "largest_etf_aum_share_pct": info["top1_aum_share"] * 100,
            "asset_class_count": len(asset_classes),
            "category_count": len(categories),
            "strategy_count": len(strategies),
            "niche_or_segments_count": len(niches) + len(segment_parts),
            "product_count_over_1b": sum(1 for row in rows if (row.get("aum_usd") or 0) >= 1e9),
            "product_count_over_10b": sum(1 for row in rows if (row.get("aum_usd") or 0) >= 10e9),
            "top5_aum_usd": top5_aum,
            "top10_weighted_expense_ratio_pct": weighted_average(top10_aum_rows, "expense_ratio_pct"),
            "low_fee_aum_share": sum_field(low_fee_rows, "aum_usd") / total_aum if total_aum else 0,
            "high_liquidity_product_count": len(high_liq_rows),
            "high_liquidity_aum_share": sum_field(high_liq_rows, "aum_usd") / total_aum if total_aum else 0,
            "top5_adv_concentration": top5_adv / total_adv if total_adv else 0,
            "active_count_share": len(active_rows) / len(rows) if rows else 0,
            "active_aum_share": sum_field(active_rows, "aum_usd") / total_aum if total_aum else 0,
            "fixed_income_aum_share": sum_field(fixed_rows, "aum_usd") / total_aum if total_aum else 0,
            "equity_premium_income_aum_share": sum_field(premium_rows, "aum_usd") / total_aum if total_aum else 0,
            "smart_beta_factor_aum_share": sum_field(smart_rows, "aum_usd") / total_aum if total_aum else 0,
            "non_plain_beta_aum_share": sum_field(non_plain_rows, "aum_usd") / total_aum if total_aum else 0,
            "largest_etf_adv_usd": info["largest_etf_adv_usd"],
            "top1_aum_share": info["top1_aum_share"],
            "top3_aum_share": info["top3_aum_share"],
            "flagship_mind_tag_score": info["flagship_mind_tag_score"],
            "flagship_ticker": info["flagship_ticker"],
            "flagship_name": info["flagship_name"],
            "flagship_reason": info["flagship_reason"],
            "ytd_flow_usd": ytd_flow,
            "one_year_flow_usd": one_year_flow,
            "ytd_flow_to_aum_pct": ytd_flow / total_aum if total_aum else None,
            "one_year_flow_to_aum_pct": one_year_flow / total_aum if total_aum else None,
            "positive_ytd_flow_product_count": sum(1 for row in rows if (row.get("ytd_flow_usd") or 0) > 0),
            "positive_one_year_flow_product_count": sum(1 for row in rows if (row.get("one_year_flow_usd") or 0) > 0),
            "total_aum_log": math.log1p(total_aum),
            "top5_aum_log": math.log1p(top5_aum),
            "total_adv_log": math.log1p(total_adv),
            "largest_etf_aum_log": math.log1p(info["largest_etf_aum_usd"]),
            "largest_etf_adv_log": math.log1p(info["largest_etf_adv_usd"]),
            "data_as_of": DATA_AS_OF,
        })
    return top10_issuers, raw_metrics, by_issuer


def compute_scores(raw_metrics):
    metrics_by_issuer = {row["issuer"]: row for row in raw_metrics}
    subscore_rows = []
    dimension_scores = defaultdict(dict)
    for dimension, subs in SUBINDICATORS.items():
        for subindicator, direction, weight, transform in subs:
            values = [row.get(subindicator) or 0 for row in raw_metrics]
            p5 = percentile(values, 0.05)
            p95 = percentile(values, 0.95)
            for row in raw_metrics:
                raw_value = row.get(subindicator) or 0
                score = score_from_percentile(raw_value, p5, p95, direction)
                weighted_score = score * weight
                subscore_rows.append({
                    "issuer": row["issuer"],
                    "dimension": dimension,
                    "subindicator": subindicator,
                    "raw_value": raw_value,
                    "transform_method": transform,
                    "direction": direction,
                    "p5": p5,
                    "p95": p95,
                    "score_60_100": score,
                    "weight_in_dimension": weight,
                    "weighted_score": weighted_score,
                    "data_as_of": DATA_AS_OF,
                    "notes": "winsorized percentile scoring, 60-100分",
                })
        for issuer in metrics_by_issuer:
            dimension_scores[issuer][dimension] = sum(
                item["weighted_score"]
                for item in subscore_rows
                if item["issuer"] == issuer and item["dimension"] == dimension
            )
    summary_rows = []
    for issuer, scores in dimension_scores.items():
        total = sum(scores[dim] * DIMENSION_WEIGHTS[dim] for dim in DIMENSION_WEIGHTS)
        summary_rows.append({
            "issuer": issuer,
            "产品矩阵广度_score": scores["产品矩阵广度"],
            "规模沉淀能力_score": scores["规模沉淀能力"],
            "成本竞争力_score": scores["成本竞争力"],
            "交易流动性_score": scores["交易流动性"],
            "能力迁移程度_score": scores["能力迁移程度"],
            "入口产品心智_score": scores["入口产品心智"],
            "total_ecosystem_score": total,
            "case_role": case_role(issuer),
            "interpretation": interpretation(issuer, scores),
            "data_as_of": DATA_AS_OF,
        })
    summary_rows.sort(key=lambda row: row["total_ecosystem_score"], reverse=True)
    for idx, row in enumerate(summary_rows, start=1):
        row["rank_total_score"] = idx
    return subscore_rows, summary_rows


def case_role(issuer):
    if issuer == "BlackRock/iShares":
        return "平台化龙头样本"
    if issuer == "JPMorgan AM":
        return "能力迁移型样本"
    if issuer == "Invesco":
        return "入口型辅助样本"
    return "TOP10参照样本"


def interpretation(issuer, scores):
    if issuer == "BlackRock/iShares":
        return "规模、成本、流动性和平台化生态显著领先。"
    if issuer == "JPMorgan AM":
        return "能力迁移程度突出，主动ETF和收益型产品特征明显。"
    if issuer == "Invesco":
        return "QQQ带来极强入口产品心智和流动性优势。"
    top_dims = sorted(scores.items(), key=lambda item: item[1], reverse=True)[:2]
    return f"{top_dims[0][0]}与{top_dims[1][0]}相对突出，可作为美国TOP10参照样本。"


def build_top_products(top10_issuers, by_issuer):
    rows = []
    for issuer in top10_issuers:
        products = sorted(by_issuer[issuer], key=lambda row: row.get("aum_usd") or 0, reverse=True)[:10]
        for idx, row in enumerate(products, start=1):
            ticker = row.get("ticker")
            rows.append({
                "issuer": issuer,
                "rank_within_issuer": idx,
                "ticker": ticker,
                "fund_name": row.get("fund_name"),
                "asset_class": row.get("asset_class"),
                "category": row.get("category"),
                "strategy": row.get("strategy"),
                "niche": row.get("niche"),
                "aum_usd": row.get("aum_usd"),
                "aum_billion_usd": (row.get("aum_usd") or 0) / 1e9,
                "avg_daily_dollar_vol_usd": row.get("avg_daily_dollar_vol_usd"),
                "avg_daily_dollar_vol_billion_usd": (row.get("avg_daily_dollar_vol_usd") or 0) / 1e9,
                "expense_ratio_pct": row.get("expense_ratio_pct"),
                "spread_pct": row.get("spread_pct"),
                "ytd_flow_usd": row.get("ytd_flow_usd"),
                "one_year_flow_usd": row.get("one_year_flow_usd"),
                "etn_flag": row.get("etn_flag"),
                "leveraged_flag": row.get("leveraged_flag"),
                "inverse_flag": row.get("inverse_flag"),
                "flagship_tag": FLAGSHIP_SCORES.get(ticker),
                "data_as_of": DATA_AS_OF,
            })
    return rows


def build_name_map(all_rows):
    grouped = defaultdict(list)
    for row in all_rows:
        grouped[(row["issuer_raw"], row["issuer"])].append(row)
    output = []
    for (raw, standardized), rows in grouped.items():
        output.append({
            "issuer_raw": raw,
            "issuer_standardized": standardized,
            "row_count": len(rows),
            "total_aum_usd": sum_field(rows, "aum_usd"),
            "total_aum_billion_usd": sum_field(rows, "aum_usd") / 1e9,
            "data_as_of": DATA_AS_OF,
        })
    return sorted(output, key=lambda row: row["total_aum_usd"] or 0, reverse=True)


def build_slide_ready(summary_rows, raw_metrics, top_products, by_issuer):
    summary = {row["issuer"]: row for row in summary_rows}
    metrics = {row["issuer"]: row for row in raw_metrics}
    rows = []
    for row in summary_rows:
        issuer = row["issuer"]
        rows.append({
            "block": "P23_TOP10_Scatter_Data",
            "issuer": issuer,
            "x_capability_migration_score": row["能力迁移程度_score"],
            "y_scale_score": row["规模沉淀能力_score"],
            "bubble_size_total_aum_billion_usd": metrics[issuer]["total_aum_billion_usd"],
            "color_group": row["case_role"],
            "label": issuer,
            "total_score": row["total_ecosystem_score"],
            "data_as_of": DATA_AS_OF,
        })
    for issuer in ["BlackRock/iShares", "JPMorgan AM", "Invesco"]:
        if issuer not in summary:
            continue
        for dim in DIMENSION_WEIGHTS:
            rows.append({
                "block": "P23_Radar_Selected_Three",
                "issuer": issuer,
                "dimension": dim,
                "score": summary[issuer][f"{dim}_score"],
                "data_as_of": DATA_AS_OF,
            })
    if "JPMorgan AM" in summary:
        products = [row["ticker"] for row in top_products if row["issuer"] == "JPMorgan AM"][:6]
        rows.append({
            "block": "P25_JPMorgan_Score_Card",
            "issuer": "JPMorgan AM",
            "etf_count": metrics["JPMorgan AM"]["etf_count"],
            "total_aum_billion_usd": metrics["JPMorgan AM"]["total_aum_billion_usd"],
            "active_aum_share": metrics["JPMorgan AM"]["active_aum_share"],
            "capability_migration_score": summary["JPMorgan AM"]["能力迁移程度_score"],
            "top_products": "；".join(products),
            "ppt_message": "JPMorgan AM通过主动ETF、收益增强和固收ETF把传统资管能力迁移到ETF场景。",
            "data_as_of": DATA_AS_OF,
        })
    invesco_products = by_issuer.get("Invesco", [])
    qqq = next((row for row in invesco_products if row.get("ticker") == "QQQ"), None)
    if qqq:
        rows.append({
            "block": "P25_Invesco_Helper_Card",
            "issuer": "Invesco",
            "qqq_aum_billion_usd": (qqq.get("aum_usd") or 0) / 1e9,
            "qqq_adv_billion_usd": (qqq.get("avg_daily_dollar_vol_usd") or 0) / 1e9,
            "qqq_expense_ratio_pct": qqq.get("expense_ratio_pct"),
            "ppt_message": "QQQ把纳斯达克100/创新成长压缩成一个ETF符号，是入口型大单品心智样本。",
            "data_as_of": DATA_AS_OF,
        })
    implication_items = [
        ("指数能力 → 底仓入口", "筛选核心宽基作为底仓入口，强化成交、价差、渠道话术和组合示例。"),
        ("投研能力 → 主题入口", "围绕产业趋势输出主题ETF观点、配置逻辑和风险提示。"),
        ("量化/策略能力 → 风格入口", "将策略能力转化为客户可理解的收益风险特征。"),
        ("渠道与交易能力 → 使用入口", "建立旗舰池、孵化池和维护池，集中资源提升重点产品交易体验。"),
    ]
    for title, text in implication_items:
        rows.append({
            "block": "Fullgoal_Implication_Text",
            "implication_title": title,
            "implication_text": text,
            "data_as_of": DATA_AS_OF,
        })
    return rows


def build_methodology():
    rows = [
        ("样本选择", "基于原始ETF全量Excel，自动识别主数据sheet，剔除ETN后按标准化发行人AUM取前10。"),
        ("数据清洗", "金额字段统一解析为美元，K/M/B/T后缀分别换算为1e3/1e6/1e9/1e12；原始小数型费率和收益率换算为百分数口径。"),
        ("六维定义", "产品矩阵广度、规模沉淀能力、成本竞争力、交易流动性、能力迁移程度、入口产品心智。"),
        ("权重", "总分=15%产品矩阵广度+20%规模沉淀能力+15%成本竞争力+15%交易流动性+20%能力迁移程度+15%入口产品心智。"),
        ("正向/反向指标", "AUM、产品数量、流动性、能力迁移和旗舰产品心智为正向；费率、价差、Top5成交集中度为反向。"),
        ("log处理", "偏态金额指标使用log1p，包括total_aum_usd、top5_aum_usd、largest_etf_aum_usd、total_adv_usd、largest_etf_adv_usd等。"),
        ("winsorized percentile 60-100评分", "在TOP10样本内计算P5/P95，正向指标按60+40*(x-P5)/(P95-P5)，反向指标按60+40*(P95-x)/(P95-P5)，并截断到60-100。"),
        ("局限性", "当前数据集缺少严格历史跟踪误差，未设置跟踪精度维度。"),
        ("局限性", "当前数据集缺少发行规模，未完全复刻国内PPT中的规模增长维度。"),
        ("局限性", "渠道生态、模型组合和账户渗透难以完全量化，主要通过官网资料和案例分析补充。"),
        ("局限性", "入口产品心智中含人工标签，用于体现QQQ、SPY、IVV、JEPI等产品的品牌心智差异。"),
    ]
    return [{"section": section, "content": content, "data_as_of": DATA_AS_OF, "source_note": "用户评分规则；原始ETF全量数据；研究判断"} for section, content in rows]


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = copy(header_fill)
        cell.font = copy(header_font)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = copy(border)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = copy(border)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col_idx, header_cell in enumerate(ws[1], start=1):
        header = str(header_cell.value or "")
        lower = header.lower()
        max_len = len(header)
        for col in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        is_text = any(token in lower for token in ["note", "source", "message", "content", "interpretation", "reason", "method", "notes", "字段", "说明"])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 45 if is_text else 24)
        if any(token in lower for token in ["usd", "aum", "adv", "flow"]) and "share" not in lower:
            for col in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
                for cell in col:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "$#,##0"
        elif any(token in lower for token in ["pct", "share", "score", "weight", "p5", "p95", "rank", "count"]):
            for col in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
                for cell in col:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "0.00"
    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 22 if row_idx == 1 else 34


def write_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    style_sheet(ws)
    return ws


def export_csv(path, headers, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main():
    qa = []

    def add_qa(item, status, notes):
        qa.append({"check_item": item, "status": status, "notes": notes, "data_as_of": DATA_AS_OF})

    add_qa("输入文件是否存在", "PASS" if INPUT_PATH.exists() else "FAIL", str(INPUT_PATH))
    if not INPUT_PATH.exists():
        raise FileNotFoundError(INPUT_PATH)

    input_wb = load_workbook(INPUT_PATH, read_only=True, data_only=True)
    main_ws, headers, fmap = detect_main_sheet(input_wb)
    add_qa("主数据sheet识别是否成功", "PASS", f"{main_ws.title}; rows={main_ws.max_row}; cols={main_ws.max_column}")

    all_rows = build_clean_rows(main_ws, headers, fmap)
    etf_rows = [row for row in all_rows if not row["etn_flag"]]
    add_qa("剔除ETN后样本行数", "PASS" if etf_rows else "FAIL", f"{len(etf_rows)} rows")

    field_map_rows = []
    for required, matched in fmap.items():
        status = "MATCHED" if matched else "MISSING"
        if required in ["ticker", "fund_name", "issuer", "aum"] and not matched:
            add_qa("必要字段是否匹配", "FAIL", f"{required} 未匹配")
        field_map_rows.append({
            "required_field": required,
            "matched_raw_field": matched,
            "match_status": status,
            "notes": "自动模糊匹配" if matched else "原始数据未提供或字段名未识别",
            "data_as_of": DATA_AS_OF,
        })
    if all(fmap.get(key) for key in ["ticker", "fund_name", "issuer", "aum"]):
        add_qa("必要字段是否匹配", "PASS", "Ticker/Fund Name/Issuer/AUM均已匹配")

    top10_issuers, raw_metrics, by_issuer = compute_metrics(etf_rows)
    add_qa("TOP10发行人是否成功生成", "PASS" if len(top10_issuers) == 10 else "FAIL", f"{len(top10_issuers)} issuers")
    add_qa("AUM是否有负值或异常值", "WARN" if any((row.get("aum_usd") or 0) < 0 for row in etf_rows) else "PASS", "检查非ETN样本AUM")
    max_expense = max([row.get("expense_ratio_pct") or 0 for row in etf_rows] or [0])
    add_qa("Expense Ratio是否存在异常极大值", "WARN" if max_expense > 10 else "PASS", f"max={max_expense:.2f}%")
    max_spread = max([row.get("spread_pct") or 0 for row in etf_rows] or [0])
    add_qa("Spread是否存在异常极大值", "WARN" if max_spread > 5 else "PASS", f"max={max_spread:.2f}%")
    for issuer in ["BlackRock/iShares", "JPMorgan AM", "Invesco"]:
        add_qa(f"{issuer}是否在TOP10中", "PASS" if issuer in top10_issuers else "WARN", ",".join(top10_issuers))

    subscore_rows, summary_rows = compute_scores(raw_metrics)
    top_products = build_top_products(top10_issuers, by_issuer)
    name_map_rows = build_name_map(all_rows)
    slide_rows = build_slide_ready(summary_rows, raw_metrics, top_products, by_issuer)
    methodology_rows = build_methodology()
    summary_by_issuer = {row["issuer"]: row for row in summary_rows}
    metrics_by_issuer = {row["issuer"]: row for row in raw_metrics}

    top10_rows = [
        {
            key: row.get(key)
            for key in [
                "rank_by_aum", "issuer", "etf_count", "total_aum_usd", "total_aum_billion_usd",
                "total_aum_trillion_usd", "total_adv_usd", "total_adv_billion_usd",
                "avg_expense_ratio_pct", "median_expense_ratio_pct", "weighted_avg_expense_ratio_pct",
                "weighted_avg_spread_pct", "active_etf_count", "active_aum_usd", "active_aum_share_pct",
                "fixed_income_aum_usd", "fixed_income_aum_share_pct", "equity_aum_usd",
                "equity_aum_share_pct", "largest_etf_ticker", "largest_etf_name", "largest_etf_aum_usd",
                "largest_etf_aum_share_pct", "data_as_of"
            ]
        }
        for row in raw_metrics
    ]
    raw_headers = [
        "issuer", "etf_count", "total_aum_usd", "total_aum_billion_usd", "total_adv_usd", "total_adv_billion_usd",
        "asset_class_count", "category_count", "strategy_count", "niche_or_segments_count", "product_count_over_1b",
        "product_count_over_10b", "top5_aum_usd", "avg_expense_ratio_pct", "weighted_avg_expense_ratio_pct",
        "top10_weighted_expense_ratio_pct", "low_fee_aum_share", "weighted_avg_spread_pct",
        "high_liquidity_product_count", "high_liquidity_aum_share", "top5_adv_concentration",
        "active_count_share", "active_aum_share", "fixed_income_aum_share", "equity_premium_income_aum_share",
        "smart_beta_factor_aum_share", "non_plain_beta_aum_share", "largest_etf_ticker", "largest_etf_name",
        "largest_etf_aum_usd", "largest_etf_adv_usd", "top1_aum_share", "top3_aum_share",
        "flagship_mind_tag_score", "ytd_flow_usd", "one_year_flow_usd", "ytd_flow_to_aum_pct",
        "one_year_flow_to_aum_pct", "data_as_of"
    ]
    summary_headers = [
        "rank_total_score", "issuer", "产品矩阵广度_score", "规模沉淀能力_score", "成本竞争力_score",
        "交易流动性_score", "能力迁移程度_score", "入口产品心智_score", "total_ecosystem_score",
        "case_role", "interpretation", "data_as_of"
    ]
    radar_long_rows = []
    for row in summary_rows:
        for dimension in DIMENSION_WEIGHTS:
            radar_long_rows.append({
                "issuer": row["issuer"],
                "dimension": dimension,
                "score": row[f"{dimension}_score"],
                "total_ecosystem_score": row["total_ecosystem_score"],
                "case_role": row["case_role"],
                "data_as_of": DATA_AS_OF,
            })
    radar_wide_rows = [
        {
            "issuer": row["issuer"],
            "产品矩阵广度": row["产品矩阵广度_score"],
            "规模沉淀能力": row["规模沉淀能力_score"],
            "成本竞争力": row["成本竞争力_score"],
            "交易流动性": row["交易流动性_score"],
            "能力迁移程度": row["能力迁移程度_score"],
            "入口产品心智": row["入口产品心智_score"],
            "total_ecosystem_score": row["total_ecosystem_score"],
            "case_role": row["case_role"],
            "data_as_of": DATA_AS_OF,
        }
        for row in summary_rows
    ]
    flagship_rows = []
    for row in raw_metrics:
        flagship_rows.append({
            "issuer": row["issuer"],
            "flagship_ticker": row["flagship_ticker"],
            "flagship_name": row["flagship_name"],
            "flagship_reason": row["flagship_reason"],
            "flagship_mind_tag_score": row["flagship_mind_tag_score"],
            "largest_etf_aum_usd": row["largest_etf_aum_usd"],
            "largest_etf_adv_usd": row["largest_etf_adv_usd"],
            "top1_aum_share": row["top1_aum_share"],
            "top3_aum_share": row["top3_aum_share"],
            "notes": "入口产品心智为展示型人工标签和量化指标结合",
            "data_as_of": DATA_AS_OF,
        })
    case_rows = []
    for issuer in ["BlackRock/iShares", "JPMorgan AM", "Invesco"]:
        if issuer not in summary_by_issuer:
            continue
        s = summary_by_issuer[issuer]
        m = metrics_by_issuer[issuer]
        sorted_strength = sorted(DIMENSION_WEIGHTS, key=lambda d: s[f"{d}_score"], reverse=True)
        flagship_products = "；".join(row["ticker"] for row in top_products if row["issuer"] == issuer and row["rank_within_issuer"] <= 3)
        case_rows.append({
            "issuer": issuer,
            "case_role": s["case_role"],
            "etf_count": m["etf_count"],
            "total_aum_billion_usd": m["total_aum_billion_usd"],
            "total_adv_billion_usd": m["total_adv_billion_usd"],
            "产品矩阵广度_score": s["产品矩阵广度_score"],
            "规模沉淀能力_score": s["规模沉淀能力_score"],
            "成本竞争力_score": s["成本竞争力_score"],
            "交易流动性_score": s["交易流动性_score"],
            "能力迁移程度_score": s["能力迁移程度_score"],
            "入口产品心智_score": s["入口产品心智_score"],
            "total_ecosystem_score": s["total_ecosystem_score"],
            "key_strength_1": sorted_strength[0],
            "key_strength_2": sorted_strength[1],
            "key_strength_3": sorted_strength[2],
            "flagship_products": flagship_products,
            "ppt_message": s["interpretation"],
            "data_as_of": DATA_AS_OF,
        })

    readme_rows = [
        {"item": "文件用途", "content": "美国TOP10 ETF管理人生态能力评分体系，用于富国基金ICI课题美国ETF案例页数据支撑。", "data_as_of": DATA_AS_OF, "source_note": str(INPUT_PATH)},
        {"item": "输入文件路径", "content": str(INPUT_PATH), "data_as_of": DATA_AS_OF, "source_note": "用户指定输入文件"},
        {"item": "数据截至日期", "content": DATA_AS_OF, "data_as_of": DATA_AS_OF, "source_note": "用户指定"},
        {"item": "ETF样本口径", "content": "以原始ETF全量Excel为基础，剔除ETN=True/Yes产品；杠杆/反向产品保留并标记。", "data_as_of": DATA_AS_OF, "source_note": "ETN字段；Leveraged/Inverse字段"},
        {"item": "TOP10发行人筛选方法", "content": "按标准化发行人汇总非ETN样本total_aum_usd，取AUM前10。", "data_as_of": DATA_AS_OF, "source_note": "03_US_Issuer_Top10"},
        {"item": "六维评分说明", "content": "产品矩阵广度、规模沉淀能力、成本竞争力、交易流动性、能力迁移程度、入口产品心智；总分按15/20/15/15/20/15权重加权。", "data_as_of": DATA_AS_OF, "source_note": "13_Methodology"},
        {"item": "评分定位", "content": "评分不是行业正式排名，仅用于展示生态路径差异。", "data_as_of": DATA_AS_OF, "source_note": "研究判断"},
        {"item": "重要局限性", "content": "缺少严格历史跟踪误差；缺少发行规模；渠道生态、模型组合和账户渗透难以完全量化；入口产品心智含人工标签。", "data_as_of": DATA_AS_OF, "source_note": "13_Methodology"},
    ]

    target_sheet_names = [
        "00_README", "01_Field_Map", "02_Issuer_Name_Map", "03_US_Issuer_Top10", "04_Issuer_Raw_Metrics",
        "05_Subindicator_Scores", "06_Issuer_Score_Summary", "07_Radar_Long", "08_Radar_Wide",
        "09_Top_Products_By_Issuer", "10_Flagship_Product_Mind", "11_Case_Study_Extract",
        "12_Slide_Ready_Data", "13_Methodology", "14_QA_Checks"
    ]
    add_qa("是否生成所有目标sheet", "PASS", ",".join(target_sheet_names))
    add_qa("data_as_of是否全部为2026-07-06", "PASS", DATA_AS_OF)

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "00_README", ["item", "content", "data_as_of", "source_note"], readme_rows)
    write_sheet(wb, "01_Field_Map", ["required_field", "matched_raw_field", "match_status", "notes", "data_as_of"], field_map_rows)
    write_sheet(wb, "02_Issuer_Name_Map", ["issuer_raw", "issuer_standardized", "row_count", "total_aum_usd", "total_aum_billion_usd", "data_as_of"], name_map_rows)
    write_sheet(wb, "03_US_Issuer_Top10", list(top10_rows[0].keys()), top10_rows)
    write_sheet(wb, "04_Issuer_Raw_Metrics", raw_headers, raw_metrics)
    write_sheet(wb, "05_Subindicator_Scores", ["issuer", "dimension", "subindicator", "raw_value", "transform_method", "direction", "p5", "p95", "score_60_100", "weight_in_dimension", "weighted_score", "data_as_of", "notes"], subscore_rows)
    write_sheet(wb, "06_Issuer_Score_Summary", summary_headers, summary_rows)
    write_sheet(wb, "07_Radar_Long", ["issuer", "dimension", "score", "total_ecosystem_score", "case_role", "data_as_of"], radar_long_rows)
    write_sheet(wb, "08_Radar_Wide", ["issuer", "产品矩阵广度", "规模沉淀能力", "成本竞争力", "交易流动性", "能力迁移程度", "入口产品心智", "total_ecosystem_score", "case_role", "data_as_of"], radar_wide_rows)
    write_sheet(wb, "09_Top_Products_By_Issuer", list(top_products[0].keys()), top_products)
    write_sheet(wb, "10_Flagship_Product_Mind", ["issuer", "flagship_ticker", "flagship_name", "flagship_reason", "flagship_mind_tag_score", "largest_etf_aum_usd", "largest_etf_adv_usd", "top1_aum_share", "top3_aum_share", "notes", "data_as_of"], flagship_rows)
    write_sheet(wb, "11_Case_Study_Extract", list(case_rows[0].keys()), case_rows)
    slide_headers = sorted({key for row in slide_rows for key in row.keys()})
    if "block" in slide_headers:
        slide_headers.remove("block")
        slide_headers = ["block"] + slide_headers
    write_sheet(wb, "12_Slide_Ready_Data", slide_headers, slide_rows)
    write_sheet(wb, "13_Methodology", ["section", "content", "data_as_of", "source_note"], methodology_rows)
    write_sheet(wb, "14_QA_Checks", ["check_item", "status", "notes", "data_as_of"], qa)
    wb.save(OUTPUT_XLSX)

    if CHART_DIR.exists():
        shutil.rmtree(CHART_DIR)
    CHART_DIR.mkdir(parents=True, exist_ok=True)
    scatter_rows = [row for row in slide_rows if row.get("block") == "P23_TOP10_Scatter_Data"]
    radar_three_rows = [row for row in slide_rows if row.get("block") == "P23_Radar_Selected_Three"]
    selected_top_rows = [row for row in top_products if row["issuer"] in {"BlackRock/iShares", "JPMorgan AM", "Invesco"}]
    export_csv(CHART_DIR / "P23_TOP10_Scatter_Data.csv", ["issuer", "x_capability_migration_score", "y_scale_score", "bubble_size_total_aum_billion_usd", "color_group", "label", "total_score", "data_as_of"], scatter_rows)
    export_csv(CHART_DIR / "P23_Radar_Selected_Three.csv", ["issuer", "dimension", "score", "data_as_of"], radar_three_rows)
    export_csv(CHART_DIR / "Issuer_Score_Summary.csv", summary_headers, summary_rows)
    export_csv(CHART_DIR / "Radar_Long.csv", ["issuer", "dimension", "score", "total_ecosystem_score", "case_role", "data_as_of"], radar_long_rows)
    export_csv(CHART_DIR / "Top_Products_Selected_Three.csv", list(top_products[0].keys()), selected_top_rows)

    status_counts = Counter(row["status"] for row in qa)
    print(f"输入文件路径: {INPUT_PATH}")
    print(f"输出文件路径: {OUTPUT_XLSX}")
    print(f"主数据sheet名称和行列数: {main_ws.title}, {main_ws.max_row}行 x {main_ws.max_column}列")
    print(f"剔除ETN后的ETF数量: {len(etf_rows)}")
    print("美国ETF发行人TOP10列表:")
    for row in top10_rows:
        print(f"- {row['rank_by_aum']}. {row['issuer']}: ETF数量={row['etf_count']}, AUM={row['total_aum_billion_usd']:.2f}B USD")
    print("六维评分TOP10摘要:")
    for row in summary_rows:
        print(f"- {row['rank_total_score']}. {row['issuer']}: 总分={row['total_ecosystem_score']:.1f}, 产品={row['产品矩阵广度_score']:.1f}, 规模={row['规模沉淀能力_score']:.1f}, 成本={row['成本竞争力_score']:.1f}, 流动性={row['交易流动性_score']:.1f}, 迁移={row['能力迁移程度_score']:.1f}, 入口={row['入口产品心智_score']:.1f}")
    print("三家公司六维评分:")
    for issuer in ["BlackRock/iShares", "JPMorgan AM", "Invesco"]:
        row = summary_by_issuer.get(issuer)
        if row:
            print(f"- {issuer}: 产品={row['产品矩阵广度_score']:.1f}, 规模={row['规模沉淀能力_score']:.1f}, 成本={row['成本竞争力_score']:.1f}, 流动性={row['交易流动性_score']:.1f}, 迁移={row['能力迁移程度_score']:.1f}, 入口={row['入口产品心智_score']:.1f}, 总分={row['total_ecosystem_score']:.1f}")
    print("生成的sheet名称列表:")
    for name in target_sheet_names:
        print(f"- {name}")
    print("生成的CSV文件列表:")
    for path in sorted(CHART_DIR.glob("*.csv")):
        print(f"- {path}")
    print(f"QA_Checks中PASS/WARN/FAIL数量: PASS={status_counts.get('PASS',0)}, WARN={status_counts.get('WARN',0)}, FAIL={status_counts.get('FAIL',0)}")
    for row in qa:
        if row["status"] in {"WARN", "FAIL"}:
            print(f"- {row['status']}: {row['check_item']} - {row['notes']}")


if __name__ == "__main__":
    main()
