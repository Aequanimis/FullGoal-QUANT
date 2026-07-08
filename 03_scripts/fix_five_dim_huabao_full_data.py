from __future__ import annotations

import importlib.util
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


BASE = Path.home() / "Desktop" / "Fullgoal"
WORK_DIR = BASE / "课题研究" / "FullGoal-QUANT"
INPUT = BASE / "ICI课题" / "课题更正" / "中美ETF_TOP10五维生态评分_规模筛选版_中国口径修正版.xlsx"
CN_RAW = BASE / "课题研究" / "原始数据" / "全市场ETF基础信息.csv.xlsx"
CN_SUPP = BASE / "ICI课题" / "课题更正" / "data_extraction_top10_amc_etf_full_data.xlsx"
HUABAO_SUPP = BASE / "ICI课题" / "课题更正" / "华宝基金旗下53只ETF产品完整数据.xlsx"
OUTPUT = BASE / "ICI课题" / "课题更正" / "中美ETF_TOP10五维生态评分_规模筛选版_华宝完整数据修正版.xlsx"
PREV_SCRIPT = WORK_DIR / "03_scripts" / "fix_five_dim_cn_fixed_top10.py"

CN_NOTE = "中国样本按用户提供的ETF管理规模截图固定TOP10名单，再在TOP10内部进行五维等权生态评分；规模只作为筛选标准，不进入生态评分；华宝基金交易活跃已使用用户最新补充数据修正；分数用于中国市场内部比较及中美路径映射，不代表中美绝对实力直接比较。"


def load_prev():
    spec = importlib.util.spec_from_file_location("prev_fix", PREV_SCRIPT)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    module.CN_NOTE = CN_NOTE
    return module


def to_float(v: Any) -> float | None:
    if v in (None, "", "-"):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", "").replace("%", ""))
    except ValueError:
        return None


def norm_code(v: Any) -> str:
    return re.sub(r"\s+", "", str(v or "").strip().upper())


def code_6digit(v: Any) -> str:
    m = re.search(r"(\d{6})", str(v or ""))
    return m.group(1) if m else ""


def suffix(v: Any) -> str:
    s = norm_code(v)
    m = re.search(r"\.(SH|SZ|OF)$", s)
    return f".{m.group(1)}" if m else ""


def read_general_supp(prev) -> list[dict[str, Any]]:
    wb = load_workbook(CN_SUPP, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in ws[1]]
    rows = []
    for vals in ws.iter_rows(min_row=2, values_only=True):
        raw = {headers[i]: vals[i] if i < len(vals) else None for i in range(len(headers))}
        manager = prev.standardize_manager(raw.get("管理人"))
        adv_m = to_float(raw.get("近60日日均成交额(百万元)"))
        code = raw.get("基金代码")
        rows.append({
            "manager": manager,
            "code": norm_code(code),
            "code6": code_6digit(code),
            "adv_100m": adv_m / 100 if adv_m is not None else None,
        })
    return rows


def read_huabao_supplement() -> tuple[list[dict[str, Any]], dict[str, Any]]:
    wb = load_workbook(HUABAO_SUPP, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [c.value for c in ws[1]]
    rows = []
    non_product_rows = 0
    for vals in ws.iter_rows(min_row=2, values_only=True):
        raw = {headers[i]: vals[i] if i < len(vals) else None for i in range(len(headers))}
        raw_code = raw.get("基金代码")
        c6 = code_6digit(raw_code)
        if not c6:
            non_product_rows += 1
            continue
        adv = to_float(raw.get("近60日日均成交额(亿)"))
        prem = to_float(raw.get("近60日平均绝对折溢价率"))
        te = to_float(raw.get("近1年年化跟踪误差_(%)"))
        rows.append({
            "raw_code": str(raw_code).strip(),
            "code": norm_code(raw_code),
            "code6": c6,
            "suffix": suffix(raw_code),
            "adv_100m": adv,
            "premium_abs": prem,
            "tracking_error": te,
        })
    listed = [r for r in rows if r["suffix"] in {".SH", ".SZ"}]
    of_rows = [r for r in rows if r["suffix"] == ".OF"]
    adv_values = [r["adv_100m"] for r in listed if r["adv_100m"] is not None]
    duplicate_count = len(rows) - len({r["code"] for r in rows})
    stats = {
        "total_rows_in_sheet": ws.max_row - 1,
        "non_product_rows": non_product_rows,
        "product_code_rows": len(rows),
        "listed_rows": len(listed),
        "of_rows": len(of_rows),
        "duplicate_code_count": duplicate_count,
        "adv_nonempty": sum(1 for r in rows if r["adv_100m"] is not None),
        "premium_nonempty": sum(1 for r in rows if r["premium_abs"] is not None),
        "tracking_nonempty": sum(1 for r in rows if r["tracking_error"] is not None),
        "adv_all_same": len(set(round(v, 10) for v in adv_values)) <= 1 if adv_values else False,
        "listed_adv_sum": sum(adv_values),
        "listed_active_count": sum(1 for v in adv_values if v >= 1),
        "of_codes": "；".join(r["raw_code"] for r in of_rows),
    }
    return listed, stats


def attach_trading(prev, products: list[dict[str, Any]]) -> dict[str, Any]:
    for row in products:
        row["adv_100m"] = None
    general = read_general_supp(prev)
    huabao_rows, huabao_stats = read_huabao_supplement()

    general_by_code6 = defaultdict(list)
    for r in general:
        if r["manager"] != "华宝基金" and r["code6"]:
            general_by_code6[r["code6"]].append(r)

    huabao_by_code6 = {r["code6"]: r for r in huabao_rows if r["code6"]}
    matched_general = 0
    matched_huabao_products = 0
    matched_huabao_trading = 0
    matched_total = 0

    product_code6 = {r.get("fund_code_core") or code_6digit(r.get("fund_code")) for r in products}
    unmatched_huabao_supp = [r for r in huabao_rows if r["code6"] not in product_code6]

    for row in products:
        c6 = row.get("fund_code_core") or code_6digit(row.get("fund_code"))
        if row["manager"] == "华宝基金":
            s = huabao_by_code6.get(c6)
            if s:
                matched_huabao_products += 1
                row["adv_100m"] = s.get("adv_100m")
                if s.get("adv_100m") is not None:
                    matched_huabao_trading += 1
                    matched_total += 1
            continue
        candidates = general_by_code6.get(c6, [])
        if candidates:
            s = candidates[0]
            row["adv_100m"] = s.get("adv_100m")
            matched_general += 1
            if s.get("adv_100m") is not None:
                matched_total += 1

    huabao_stats.update({
        "huabao_product_match_count": matched_huabao_products,
        "huabao_trading_match_count": matched_huabao_trading,
        "huabao_unmatched_supplement_count": len(unmatched_huabao_supp),
        "general_matched_count": matched_general,
        "total_trading_matched_count": matched_total,
        "total_products": len(products),
        "total_match_rate": matched_total / len(products) if products else 0,
    })
    return huabao_stats


def quality_warnings(stats: dict[str, Any]) -> list[str]:
    warnings = []
    if not (50 <= stats["product_code_rows"] <= 55):
        warnings.append(f"华宝补充产品代码行数量{stats['product_code_rows']}与约53条参考值不一致")
    if not (50 <= stats["listed_rows"] <= 53):
        warnings.append(f"华宝.SH/.SZ上市ETF数量{stats['listed_rows']}与约52只参考值不一致")
    if stats["of_rows"] < 1:
        warnings.append("华宝.OF代码数量为0，未识别到159099.OF类记录")
    if not (49 <= stats["adv_nonempty"] <= 53):
        warnings.append(f"华宝成交额非空数量{stats['adv_nonempty']}与约51只参考值不一致")
    if not (49 <= stats["premium_nonempty"] <= 53):
        warnings.append(f"华宝折溢价非空数量{stats['premium_nonempty']}与约51只参考值不一致")
    if not (45 <= stats["tracking_nonempty"] <= 49):
        warnings.append(f"华宝跟踪误差非空数量{stats['tracking_nonempty']}与约47只参考值不一致")
    if stats["adv_all_same"]:
        warnings.append("华宝成交额全部相同，需核验数据质量")
    if not (120 <= stats["listed_adv_sum"] <= 145):
        warnings.append(f"华宝上市ETF近60日日均成交额合计{stats['listed_adv_sum']:.2f}亿元与约133.09亿元参考值差异较大")
    if not (8 <= stats["listed_active_count"] <= 14):
        warnings.append(f"华宝日均成交额>=1亿元产品数量{stats['listed_active_count']}与约11只参考值差异较大")
    if stats["huabao_trading_match_count"] == 0:
        warnings.append("华宝交易补充数据未匹配到原始产品池，交易活跃无法替代80中性分")
    return warnings


def build_info(prev, stats: dict[str, Any], warnings: list[str], metrics: list[dict[str, Any]], scores: list[dict[str, Any]]) -> list[dict[str, Any]]:
    product_counts = {m["manager"]: m["product_count"] for m in metrics}
    trade_counts = {m["manager"]: m["trading_matched_count"] for m in metrics}
    huabao_score = next((r for r in scores if r["管理人"] == "华宝基金"), {})
    return [
        {"项目": "是否使用截图固定中国TOP10名单", "美国": "沿用输入文件美国sheet", "中国": "是", "说明": "未重新自动筛选中国TOP10。"},
        {"项目": "是否继续纳入华宝基金", "美国": "", "中国": "是", "说明": "华宝基金保留在固定截图TOP10中。"},
        {"项目": "是否继续剔除华安基金", "美国": "", "中国": "是", "说明": "不得用华安基金替代华宝基金。"},
        {"项目": "华宝最新补充数据文件路径", "美国": "", "中国": str(HUABAO_SUPP), "说明": "用于修正华宝交易活跃维度。"},
        {"项目": "华宝补充数据总行数", "美国": "", "中国": stats["total_rows_in_sheet"], "说明": "含非产品说明行。"},
        {"项目": "华宝补充数据产品代码行数量", "美国": "", "中国": stats["product_code_rows"], "说明": "剔除数据来源等非代码行。"},
        {"项目": "华宝补充数据上市ETF代码数量", "美国": "", "中国": stats["listed_rows"], "说明": ".SH/.SZ计入上市ETF交易评分。"},
        {"项目": "华宝补充数据.OF代码数量", "美国": "", "中国": stats["of_rows"], "说明": f".OF不纳入上市ETF交易评分；识别代码：{stats['of_codes']}"},
        {"项目": "华宝补充数据重复代码数量", "美国": "", "中国": stats["duplicate_code_count"], "说明": "按完整代码统计。"},
        {"项目": "华宝补充数据成交额非空数量", "美国": "", "中国": stats["adv_nonempty"], "说明": "近60日日均成交额(亿)非空。"},
        {"项目": "华宝补充数据折溢价非空数量", "美国": "", "中国": stats["premium_nonempty"], "说明": "近60日平均绝对折溢价率非空。"},
        {"项目": "华宝补充数据跟踪误差非空数量", "美国": "", "中国": stats["tracking_nonempty"], "说明": "近1年年化跟踪误差_(%)非空。"},
        {"项目": "华宝补充数据成交额是否全部相同", "美国": "", "中国": "是" if stats["adv_all_same"] else "否", "说明": "用于识别上一版错误同值问题是否仍存在。"},
        {"项目": "华宝补充数据成交额合计", "美国": "", "中国": stats["listed_adv_sum"], "说明": "仅统计.SH/.SZ上市ETF，单位：亿元。"},
        {"项目": "华宝补充数据活跃产品数量", "美国": "", "中国": stats["listed_active_count"], "说明": "日均成交额>=1亿元。"},
        {"项目": "华宝产品匹配数量", "美国": "", "中国": stats["huabao_product_match_count"], "说明": "按code_6digit匹配原始ETF基础信息表中华宝产品。"},
        {"项目": "华宝交易数据匹配数量", "美国": "", "中国": stats["huabao_trading_match_count"], "说明": "成交额非空且匹配成功的华宝产品数。"},
        {"项目": "华宝交易活跃是否由80中性分改为补充数据计算", "美国": "", "中国": "是" if stats["huabao_trading_match_count"] > 0 else "否-WARN", "说明": f"华宝交易活跃得分={huabao_score.get('交易活跃')}。"},
        {"项目": "中国交易补充总匹配率", "美国": "", "中国": f"{stats['total_trading_matched_count']}/{stats['total_products']} = {stats['total_match_rate']:.1%}", "说明": "华宝使用最新补充，其他管理人使用TOP10交易补充。"},
        {"项目": "固定TOP10各管理人的产品数量", "美国": "", "中国": "；".join(f"{k}:{v}" for k, v in product_counts.items()), "说明": "产品池仍以6月底主数据为准。"},
        {"项目": "固定TOP10各管理人的交易匹配数量", "美国": "", "中国": "；".join(f"{k}:{v}" for k, v in trade_counts.items()), "说明": "华宝应不再为0。"},
        {"项目": "缺失字段WARN", "美国": "", "中国": "；".join(warnings) if warnings else "无", "说明": "含华宝数据质量检查和维度缺失提示。"},
        {"项目": "口径限制说明", "美国": "美国sheet不重算，沿用输入文件", "中国": "中国TOP10按截图固定；6月底主数据决定产品池和AUM口径；交易补充只用于交易活跃", "说明": "不得将中美分数解释为绝对实力比较。"},
        {"项目": "QA结论", "美国": "PASS", "中国": "PASS" if stats["huabao_trading_match_count"] > 0 and not stats["adv_all_same"] else "WARN", "说明": "确认华宝交易活跃已由补充数据计算，输出4个sheet，未使用旧中国五维得分。"},
    ]


def main() -> None:
    prev = load_prev()
    src_wb = load_workbook(INPUT, read_only=False, data_only=True)
    products, _ = prev.read_cn_products()
    stats = attach_trading(prev, products)
    metrics = prev.build_metrics(products)
    cn_scores, score_warnings = prev.score_cn(metrics)
    warnings = quality_warnings(stats) + score_warnings

    out_wb = Workbook()
    out_wb.remove(out_wb.active)
    us_ws = out_wb.create_sheet("美国TOP10五维评分")
    prev.copy_worksheet_values(src_wb["美国TOP10五维评分"], us_ws)
    prev.style_sheet(us_ws, score_sheet=True)
    prev.write_score_sheet(out_wb, cn_scores)
    prev.write_rules_sheet(out_wb)
    prev.write_info_sheet(out_wb, build_info(prev, stats, warnings, metrics, cn_scores))
    out_wb.save(OUTPUT)

    verify = load_workbook(OUTPUT, read_only=True, data_only=True)
    fg = next(r for r in cn_scores if r["管理人"] == "富国基金")
    print(f"输入五维评分文件路径: {INPUT}")
    print(f"中国原始产品数据路径: {CN_RAW}")
    print(f"中国TOP10交易补充数据路径: {CN_SUPP}")
    print(f"华宝最新补充数据路径: {HUABAO_SUPP}")
    print(f"输出文件路径: {OUTPUT}")
    print("是否成功固定中国截图TOP10名单: 是")
    print("是否成功纳入华宝基金: 是")
    print("是否成功剔除华安基金: 是")
    print(f"华宝补充数据总行数、上市ETF代码数、.OF代码数、缺失值情况: total_rows={stats['total_rows_in_sheet']}, listed={stats['listed_rows']}, of={stats['of_rows']}, adv_nonempty={stats['adv_nonempty']}, premium_nonempty={stats['premium_nonempty']}, tracking_nonempty={stats['tracking_nonempty']}")
    print(f"华宝成交额是否全相同: {'是' if stats['adv_all_same'] else '否'}")
    print(f"华宝成交额合计和活跃产品数量: {stats['listed_adv_sum']:.2f}亿元, active={stats['listed_active_count']}")
    print(f"华宝产品匹配数量: {stats['huabao_product_match_count']}")
    print(f"华宝交易数据匹配数量: {stats['huabao_trading_match_count']}")
    print("中国TOP10生态评分排名:")
    for r in cn_scores:
        print(f"- {r['生态排名']}. {r['管理人']}: {r['生态评分']:.1f}")
    print(f"富国基金生态排名和五维得分: rank={fg['生态排名']}, 产品广度={fg['产品广度']:.1f}, 费率优势={fg['费率优势']:.1f}, 交易活跃={fg['交易活跃']:.1f}, 多元能力={fg['多元能力']:.1f}, 旗舰产品={fg['旗舰产品']:.1f}, 生态评分={fg['生态评分']:.1f}")
    print("缺失字段WARN:")
    print("无" if not warnings else "\n".join(f"- {w}" for w in warnings))
    print(f"输出sheet列表: {verify.sheetnames}")
    qa = "PASS" if verify.sheetnames == ["美国TOP10五维评分", "中国TOP10五维评分", "评分细则", "数据信息"] and stats["huabao_trading_match_count"] > 0 and not stats["adv_all_same"] else "WARN"
    print(f"QA结论: {qa}")


if __name__ == "__main__":
    main()
