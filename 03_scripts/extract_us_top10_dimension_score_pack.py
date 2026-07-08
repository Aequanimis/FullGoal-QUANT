from __future__ import annotations

import csv
import shutil
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DATA_AS_OF = "2026-07-06"

INPUT_PATH = (
    Path.home()
    / "Desktop"
    / "Fullgoal"
    / "ICI课题"
    / "课题更正"
    / "美国ETF_TOP10生态评分数据包_20260706.xlsx"
)
OUTPUT_PATH = (
    Path.home()
    / "Desktop"
    / "Fullgoal"
    / "ICI课题"
    / "课题更正"
    / "美国ETF_TOP10维度评分提取版_20260706.xlsx"
)
CSV_DIR = (
    Path.home()
    / "Desktop"
    / "Fullgoal"
    / "ICI课题"
    / "课题更正"
    / "美国ETF_TOP10维度评分提取_ChartData"
)

REQUIRED_SOURCE_SHEETS = [
    "03_US_Issuer_Top10",
    "04_Issuer_Raw_Metrics",
    "05_Subindicator_Scores",
    "06_Issuer_Score_Summary",
    "07_Radar_Long",
    "08_Radar_Wide",
    "13_Methodology",
    "14_QA_Checks",
]

DIMENSION_WEIGHTS = {
    "产品矩阵广度": 0.15,
    "规模沉淀能力": 0.20,
    "成本竞争力": 0.15,
    "交易流动性": 0.15,
    "能力迁移程度": 0.20,
    "入口产品心智": 0.15,
}

SUBINDICATOR_NOTES = {
    "asset_class_count": "覆盖资产类别越多，说明产品货架越完整。",
    "category_count": "覆盖产品类别越多，说明可服务的配置场景越多。",
    "strategy_count": "策略类型越丰富，说明ETF能力封装方式越多。",
    "niche_or_segments_count": "细分赛道越丰富，说明长尾场景覆盖越强。",
    "total_aum_log": "总AUM越大，说明发行人规模沉淀越强。",
    "etf_count": "ETF数量越多，说明货架铺设越充分。",
    "product_count_over_1b": "超过10亿美元的产品越多，说明规模产品梯队越厚。",
    "product_count_over_10b": "超过100亿美元的产品越多，说明成熟大单品更多。",
    "top5_aum_log": "前五大产品AUM越大，说明头部产品规模承载更强。",
    "weighted_avg_expense_ratio_pct": "AUM加权费率越低，成本竞争力越强。",
    "top10_weighted_expense_ratio_pct": "头部产品加权费率越低，主力产品成本优势越强。",
    "low_fee_aum_share": "低费率产品AUM占比越高，说明低成本资产基础越厚。",
    "avg_expense_ratio_pct": "简单平均费率越低，整体产品费率越有竞争力。",
    "total_adv_log": "总成交额越高，说明真实交易使用更活跃。",
    "weighted_avg_spread_pct": "AUM加权买卖价差越低，交易体验越好。",
    "high_liquidity_product_count": "高流动性产品越多，说明流动性不是单点依赖。",
    "high_liquidity_aum_share": "高流动性产品AUM占比越高，说明规模资产交易承载更强。",
    "top5_adv_concentration": "成交过度集中于前五产品会削弱整体流动性生态分散度。",
    "active_aum_share": "主动ETF AUM占比越高，说明传统资管能力迁移到ETF的程度越高。",
    "active_count_share": "主动ETF数量占比越高，说明主动能力产品化更充分。",
    "fixed_income_aum_share": "固收ETF AUM占比越高，说明固收配置能力迁移更明显。",
    "equity_premium_income_aum_share": "收益增强ETF AUM占比越高，说明收益场景封装能力越强。",
    "smart_beta_factor_aum_share": "Smart Beta/因子ETF AUM占比越高，说明策略能力封装越强。",
    "non_plain_beta_aum_share": "非纯Beta产品AUM占比越高，说明产品能力更偏主动、策略或场景化。",
    "largest_etf_aum_log": "最大ETF AUM越高，说明旗舰产品规模心智越强。",
    "largest_etf_adv_log": "最大ETF成交额越高，说明旗舰产品交易心智越强。",
    "top1_aum_share": "最大ETF占比越高，说明入口产品对发行人品牌代表性越强。",
    "top3_aum_share": "前三产品占比越高，说明头部产品心智更集中。",
    "flagship_mind_tag_score": "旗舰产品识别度标签用于体现SPY、QQQ、IVV、JEPI等产品的品牌心智差异。",
}

METHOD_MAP_ROWS = [
    ("产品矩阵广度", 0.15, "asset_class_count", 0.30, "正向", "raw", "04_Issuer_Raw_Metrics.asset_class_count", "覆盖资产类别越多，产品矩阵越完整。"),
    ("产品矩阵广度", 0.15, "category_count", 0.25, "正向", "raw", "04_Issuer_Raw_Metrics.category_count", "覆盖产品类别越多，可服务的投资场景越多。"),
    ("产品矩阵广度", 0.15, "strategy_count", 0.25, "正向", "raw", "04_Issuer_Raw_Metrics.strategy_count", "策略类型越多，说明ETF能力封装方式越丰富。"),
    ("产品矩阵广度", 0.15, "niche_or_segments_count", 0.20, "正向", "raw", "04_Issuer_Raw_Metrics.niche_or_segments_count", "细分赛道越多，说明场景覆盖越广。"),
    ("规模沉淀能力", 0.20, "total_aum_log", 0.35, "正向", "log", "04_Issuer_Raw_Metrics.total_aum_usd", "总AUM越大，说明发行人规模沉淀越强。"),
    ("规模沉淀能力", 0.20, "etf_count", 0.15, "正向", "raw", "04_Issuer_Raw_Metrics.etf_count", "ETF数量越多，说明产品货架越充分。"),
    ("规模沉淀能力", 0.20, "product_count_over_1b", 0.20, "正向", "raw", "04_Issuer_Raw_Metrics.product_count_over_1b", "AUM超过10亿美元产品越多，说明规模产品梯队越厚。"),
    ("规模沉淀能力", 0.20, "product_count_over_10b", 0.20, "正向", "raw", "04_Issuer_Raw_Metrics.product_count_over_10b", "AUM超过100亿美元产品越多，说明成熟大单品更多。"),
    ("规模沉淀能力", 0.20, "top5_aum_log", 0.10, "正向", "log", "04_Issuer_Raw_Metrics.top5_aum_usd", "前五大产品AUM越大，头部规模承载越强。"),
    ("成本竞争力", 0.15, "weighted_avg_expense_ratio_pct", 0.35, "反向", "raw", "04_Issuer_Raw_Metrics.weighted_avg_expense_ratio_pct", "费率越低，成本竞争力越强。"),
    ("成本竞争力", 0.15, "top10_weighted_expense_ratio_pct", 0.25, "反向", "raw", "04_Issuer_Raw_Metrics.top10_weighted_expense_ratio_pct", "头部产品费率越低，核心货架成本优势越强。"),
    ("成本竞争力", 0.15, "low_fee_aum_share", 0.25, "正向", "raw", "04_Issuer_Raw_Metrics.low_fee_aum_share", "低费率AUM占比越高，低成本资产基础越厚。"),
    ("成本竞争力", 0.15, "avg_expense_ratio_pct", 0.15, "反向", "raw", "04_Issuer_Raw_Metrics.avg_expense_ratio_pct", "平均费率越低，整体成本越有竞争力。"),
    ("交易流动性", 0.15, "total_adv_log", 0.30, "正向", "log", "04_Issuer_Raw_Metrics.total_adv_usd", "总成交额越高，说明真实交易使用更活跃。"),
    ("交易流动性", 0.15, "weighted_avg_spread_pct", 0.25, "反向", "raw", "04_Issuer_Raw_Metrics.weighted_avg_spread_pct", "买卖价差越低，交易体验越好。"),
    ("交易流动性", 0.15, "high_liquidity_product_count", 0.20, "正向", "raw", "04_Issuer_Raw_Metrics.high_liquidity_product_count", "高流动性产品越多，说明流动性生态更分散。"),
    ("交易流动性", 0.15, "high_liquidity_aum_share", 0.15, "正向", "raw", "04_Issuer_Raw_Metrics.high_liquidity_aum_share", "高流动性AUM占比越高，规模资产交易承载越强。"),
    ("交易流动性", 0.15, "top5_adv_concentration", 0.10, "反向", "raw", "04_Issuer_Raw_Metrics.top5_adv_concentration", "成交额过度集中于少数产品，说明流动性生态依赖单点产品。"),
    ("能力迁移程度", 0.20, "active_aum_share", 0.25, "正向", "raw", "04_Issuer_Raw_Metrics.active_aum_share", "主动ETF AUM占比越高，传统资管能力迁移到ETF的程度越高。"),
    ("能力迁移程度", 0.20, "active_count_share", 0.15, "正向", "raw", "04_Issuer_Raw_Metrics.active_count_share", "主动ETF数量占比越高，主动能力产品化越充分。"),
    ("能力迁移程度", 0.20, "fixed_income_aum_share", 0.20, "正向", "raw", "04_Issuer_Raw_Metrics.fixed_income_aum_share", "固收ETF AUM占比越高，固收能力ETF化越明显。"),
    ("能力迁移程度", 0.20, "equity_premium_income_aum_share", 0.15, "正向", "raw", "04_Issuer_Raw_Metrics.equity_premium_income_aum_share", "收益增强ETF AUM占比越高，收益场景封装能力越强。"),
    ("能力迁移程度", 0.20, "smart_beta_factor_aum_share", 0.15, "正向", "raw", "04_Issuer_Raw_Metrics.smart_beta_factor_aum_share", "Smart Beta/因子ETF AUM占比越高，策略能力封装越强。"),
    ("能力迁移程度", 0.20, "non_plain_beta_aum_share", 0.10, "正向", "raw", "04_Issuer_Raw_Metrics.non_plain_beta_aum_share", "非纯Beta产品AUM占比越高，说明产品能力更场景化。"),
    ("入口产品心智", 0.15, "largest_etf_aum_log", 0.30, "正向", "log", "04_Issuer_Raw_Metrics.largest_etf_aum_usd", "最大ETF AUM越高，旗舰产品规模心智越强。"),
    ("入口产品心智", 0.15, "largest_etf_adv_log", 0.20, "正向", "log", "04_Issuer_Raw_Metrics.largest_etf_adv_usd", "最大ETF成交额越高，旗舰产品交易心智越强。"),
    ("入口产品心智", 0.15, "top1_aum_share", 0.15, "正向", "raw", "04_Issuer_Raw_Metrics.top1_aum_share", "最大ETF占比越高，入口产品对品牌代表性越强。"),
    ("入口产品心智", 0.15, "top3_aum_share", 0.15, "正向", "raw", "04_Issuer_Raw_Metrics.top3_aum_share", "前三产品占比越高，头部产品心智更集中。"),
    ("入口产品心智", 0.15, "flagship_mind_tag_score", 0.20, "正向", "raw", "04_Issuer_Raw_Metrics.flagship_mind_tag_score", "旗舰产品识别度越高，客户越容易记住发行人的ETF品牌入口。"),
]


def read_sheet_rows(wb, sheet_name: str) -> list[dict]:
    ws = wb[sheet_name]
    headers = [cell.value for cell in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {header: values[idx] for idx, header in enumerate(headers) if header}
        if any(value not in (None, "") for value in row.values()):
            rows.append(row)
    return rows


def project_rows(rows: list[dict], headers: list[str], round_scores: bool = False) -> list[dict]:
    projected = []
    for row in rows:
        new = {}
        for header in headers:
            value = row.get(header)
            if round_scores and isinstance(value, (int, float)) and header != "rank_total_score" and (
                "score" in header.lower()
                or header in {"产品矩阵广度", "规模沉淀能力", "成本竞争力", "交易流动性", "能力迁移程度", "入口产品心智"}
            ):
                value = round(float(value), 1)
            new[header] = value
        projected.append(new)
    return projected


def sort_by(rows: list[dict], field: str, reverse: bool = True) -> list[dict]:
    return sorted(rows, key=lambda row: row.get(field) or 0, reverse=reverse)


def build_readme_rows() -> list[dict]:
    return [
        {"item": "文件名称", "content": "美国ETF_TOP10维度评分提取版_20260706", "data_as_of": DATA_AS_OF, "notes": "维度评分展示轻量版"},
        {"item": "数据截至日期", "content": DATA_AS_OF, "data_as_of": DATA_AS_OF, "notes": "统一使用评分包日期"},
        {"item": "样本范围", "content": "剔除ETN后的美国ETF发行人AUM前10名", "data_as_of": DATA_AS_OF, "notes": "来源于03_US_Issuer_Top10"},
        {"item": "文件用途", "content": "用于PPT展示美国TOP10 ETF管理人生态能力评分、雷达图、散点图和评分方法说明", "data_as_of": DATA_AS_OF, "notes": "PPT作图/汇报解释/后续雷达图与散点图"},
        {"item": "评分性质", "content": "展示型、解释型评分，不作为行业正式排名", "data_as_of": DATA_AS_OF, "notes": "内部研究汇报用途"},
        {"item": "六个维度", "content": "产品矩阵广度；规模沉淀能力；成本竞争力；交易流动性；能力迁移程度；入口产品心智", "data_as_of": DATA_AS_OF, "notes": "06/07/08评分表字段"},
        {"item": "总分公式", "content": "ETF生态能力总分 = 15% × 产品矩阵广度 + 20% × 规模沉淀能力 + 15% × 成本竞争力 + 15% × 交易流动性 + 20% × 能力迁移程度 + 15% × 入口产品心智", "data_as_of": DATA_AS_OF, "notes": "维度权重"},
        {"item": "注意事项", "content": "所有分数为60-100标准化分数；金额类极端偏态指标已做log处理；标准化使用TOP10样本中的P5/P95截尾方法；入口产品心智包含旗舰产品识别度人工标签。", "data_as_of": DATA_AS_OF, "notes": "方法说明"},
        {"item": "局限性", "content": "当前数据集缺少严格跟踪误差和发行规模，因此未设置“跟踪精度”和“规模增长”维度。", "data_as_of": DATA_AS_OF, "notes": "方法说明"},
    ]


def build_methodology_summary() -> list[dict]:
    rows = [
        ("样本选择", "样本选择方法", "剔除ETN后的美国ETF发行人AUM前10名。", "按issuer汇总total_aum_usd后排序取TOP10", "来源03_US_Issuer_Top10"),
        ("样本选择", "数据截至日期", DATA_AS_OF, DATA_AS_OF, "所有表统一标注"),
        ("样本选择", "ETF样本口径", "剔除ETN；杠杆/反向产品保留并标记。", "ETN=True/Yes剔除", "来源原始ETF全量数据清洗"),
        ("六维定义", "产品矩阵广度", "衡量ETF货架覆盖资产类别、产品类别、策略和细分场景的广度。", "四个子指标加权", "总分权重15%"),
        ("六维定义", "规模沉淀能力", "衡量发行人是否真正沉淀出规模，而非只有产品数量。", "五个子指标加权", "总分权重20%"),
        ("六维定义", "成本竞争力", "衡量发行人低成本产品和低费率主力货架能力。", "四个子指标加权", "总分权重15%"),
        ("六维定义", "交易流动性", "衡量ETF真实交易使用和买卖便利程度。", "五个子指标加权", "总分权重15%"),
        ("六维定义", "能力迁移程度", "衡量主动、固收、收益增强、Smart Beta等能力迁移进ETF的程度。", "六个子指标加权", "总分权重20%"),
        ("六维定义", "入口产品心智", "衡量是否拥有客户容易记住、代表发行人ETF品牌的旗舰入口产品。", "五个子指标加权", "总分权重15%"),
        ("维度权重", "总分权重", "15%/20%/15%/15%/20%/15%。", "total_score = sum(dimension_score * dimension_weight)", "来源评分体系"),
        ("子指标权重", "每个维度的子指标及权重", "详见07_Dimension_Method_Map。", "维度内加权求和", "用于解释05_Subindicator_Scores"),
        ("标准化", "正向指标标准化公式", "x越大越好。", "score = 60 + 40 × (x - P5) / (P95 - P5)，低于P5记60，高于P95记100", "若P95=P5则记80"),
        ("标准化", "反向指标标准化公式", "x越小越好。", "score = 60 + 40 × (P95 - x) / (P95 - P5)，低于P5记100，高于P95记60", "若P95=P5则记80"),
        ("标准化", "log处理规则", "偏态金额指标先取log1p。", "log1p(total_aum_usd/top5_aum_usd/largest_etf_aum_usd/total_adv_usd/largest_etf_adv_usd)", "降低极端规模值影响"),
        ("标准化", "P5/P95截尾规则", "在TOP10样本内计算P5和P95后做60-100标准化。", "winsorized percentile scoring", "避免极端值过度拉开分数"),
        ("总分", "总分计算公式", "ETF生态能力总分 = 15% × 产品矩阵广度 + 20% × 规模沉淀能力 + 15% × 成本竞争力 + 15% × 交易流动性 + 20% × 能力迁移程度 + 15% × 入口产品心智", "加权平均", "用于06_Issuer_Score_Summary"),
        ("局限性", "跟踪精度缺失", "当前数据集缺少严格跟踪误差数据，未设置跟踪精度维度。", "", "重要局限性"),
        ("局限性", "规模增长缺失", "当前数据集缺少发行规模，未完全复刻国内PPT中的规模增长维度。", "", "重要局限性"),
        ("局限性", "渠道生态难量化", "渠道生态、模型组合和账户渗透难以完全量化，主要通过官网资料和案例分析补充。", "", "重要局限性"),
        ("局限性", "人工标签", "入口产品心智包含旗舰产品识别度人工标签。", "SPY/QQQ/IVV/JEPI等预设标签", "展示型解释变量"),
    ]
    return [
        {"section": section, "item": item, "description": desc, "formula_or_rule": formula, "notes": notes}
        for section, item, desc, formula, notes in rows
    ]


def build_dimension_method_map() -> list[dict]:
    return [
        {
            "dimension": dimension,
            "dimension_weight_in_total": dim_weight,
            "subindicator": subindicator,
            "subindicator_weight_in_dimension": sub_weight,
            "direction": direction,
            "transform_method": transform,
            "raw_metric_source": source,
            "plain_language_explanation": explanation,
        }
        for dimension, dim_weight, subindicator, sub_weight, direction, transform, source, explanation in METHOD_MAP_ROWS
    ]


def build_qa_rows(source_rows: list[dict]) -> list[dict]:
    output = []
    for row in source_rows:
        status = row.get("status")
        notes = row.get("notes")
        if status == "PASS":
            explanation = "校验通过。"
        elif status == "WARN":
            explanation = "存在提示项，需在引用时保留说明，但不影响本次评分包生成。"
        elif status == "FAIL":
            explanation = "校验失败，需要优先复核。"
        else:
            explanation = "状态未识别，请人工复核。"
        output.append({
            "check_item": row.get("check_item"),
            "status": status,
            "value": notes,
            "notes": notes,
            "qa_explanation_cn": explanation,
        })
    return output


def build_chart_index() -> list[dict]:
    return [
        {"chart_name": "TOP10总分排名柱状图", "source_sheet": "01_TOP10_Score_Ranking", "chart_type": "bar", "x_field": "issuer", "y_field": "total_ecosystem_score", "size_field": "", "category_field": "", "recommended_usage": "展示美国TOP10 ETF管理人生态能力总分排序。"},
        {"chart_name": "TOP10六维雷达图", "source_sheet": "03_Dimension_Score_Long", "chart_type": "radar", "x_field": "", "y_field": "score", "size_field": "", "category_field": "dimension", "recommended_usage": "用长表生成TOP10发行人六维雷达图。"},
        {"chart_name": "三家案例雷达图", "source_sheet": "03_Dimension_Score_Long", "chart_type": "radar", "x_field": "", "y_field": "score", "size_field": "", "category_field": "dimension", "recommended_usage": "筛选 issuer in [BlackRock/iShares, JPMorgan AM, Invesco]，对比三家案例路径差异。"},
        {"chart_name": "能力迁移 vs 规模沉淀散点图", "source_sheet": "02_Dimension_Score_Wide", "chart_type": "bubble scatter", "x_field": "能力迁移程度", "y_field": "规模沉淀能力", "size_field": "total_ecosystem_score", "category_field": "case_role", "recommended_usage": "展示TOP10发行人在能力迁移与规模沉淀两个维度上的位置。"},
        {"chart_name": "三家公司维度条形对比图", "source_sheet": "02_Dimension_Score_Wide", "chart_type": "grouped bar", "x_field": "issuer", "y_field": "六个维度字段", "size_field": "", "category_field": "dimension", "recommended_usage": "对比BlackRock/iShares、JPMorgan AM、Invesco六维差异。"},
    ]


def add_dimension_notes(rows: list[dict]) -> list[dict]:
    output = []
    for row in rows:
        copied = dict(row)
        copied["dimension_weight_note"] = SUBINDICATOR_NOTES.get(row.get("subindicator"), "该子指标用于解释维度得分来源。")
        for key in ["score_60_100", "weighted_score"]:
            if isinstance(copied.get(key), (int, float)):
                copied[key] = round(copied[key], 1)
        output.append(copied)
    return output


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.fill = copy(header_fill)
        cell.font = copy(header_font)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = copy(border)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = copy(border)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col_idx, header_cell in enumerate(ws[1], start=1):
        header = str(header_cell.value or "")
        lower = header.lower()
        max_len = len(header)
        for col in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        is_text = any(token in lower for token in ["note", "source", "interpretation", "description", "formula", "rule", "explanation", "usage", "content", "item"])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 45 if is_text else 24)

        if lower.startswith("rank"):
            fmt = "0"
        elif any(token in lower for token in ["score"]) or header in {"产品矩阵广度", "规模沉淀能力", "成本竞争力", "交易流动性", "能力迁移程度", "入口产品心智"}:
            fmt = "0.0"
        elif any(token in lower for token in ["usd", "aum", "adv"]) and "share" not in lower:
            fmt = "$#,##0"
        elif any(token in lower for token in ["pct", "share", "weight", "p5", "p95", "raw_value"]):
            fmt = "0.00"
        else:
            fmt = None
        if fmt:
            for col in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
                for cell in col:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = fmt

    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 22 if row_idx == 1 else 34


def write_sheet(wb, title: str, headers: list[str], rows: list[dict]):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    style_sheet(ws)
    return ws


def export_csv(path: Path, headers: list[str], rows: list[dict]):
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main():
    if not INPUT_PATH.exists():
        raise FileNotFoundError(f"输入文件不存在：{INPUT_PATH}")

    src = load_workbook(INPUT_PATH, data_only=True)
    missing = [sheet for sheet in REQUIRED_SOURCE_SHEETS if sheet not in src.sheetnames]
    if missing:
        raise ValueError(f"输入文件缺少必要sheet：{missing}")

    summary_rows = sort_by(read_sheet_rows(src, "06_Issuer_Score_Summary"), "total_ecosystem_score", True)
    radar_wide_rows = sort_by(read_sheet_rows(src, "08_Radar_Wide"), "total_ecosystem_score", True)
    radar_long_rows = read_sheet_rows(src, "07_Radar_Long")
    sub_rows = read_sheet_rows(src, "05_Subindicator_Scores")
    raw_rows = read_sheet_rows(src, "04_Issuer_Raw_Metrics")
    qa_rows = read_sheet_rows(src, "14_QA_Checks")

    ranking_headers = [
        "rank_total_score", "issuer", "产品矩阵广度_score", "规模沉淀能力_score", "成本竞争力_score",
        "交易流动性_score", "能力迁移程度_score", "入口产品心智_score", "total_ecosystem_score",
        "case_role", "interpretation", "data_as_of",
    ]
    wide_headers = [
        "issuer", "产品矩阵广度", "规模沉淀能力", "成本竞争力", "交易流动性", "能力迁移程度",
        "入口产品心智", "total_ecosystem_score", "case_role", "data_as_of",
    ]
    long_headers = ["issuer", "dimension", "score", "total_ecosystem_score", "case_role", "data_as_of"]
    sub_headers = [
        "issuer", "dimension", "subindicator", "raw_value", "transform_method", "direction", "p5", "p95",
        "score_60_100", "weight_in_dimension", "weighted_score", "data_as_of", "notes", "dimension_weight_note",
    ]
    raw_headers = [
        "issuer", "etf_count", "total_aum_usd", "total_aum_billion_usd", "total_adv_usd", "total_adv_billion_usd",
        "asset_class_count", "category_count", "strategy_count", "niche_or_segments_count",
        "product_count_over_1b", "product_count_over_10b", "top5_aum_usd", "avg_expense_ratio_pct",
        "weighted_avg_expense_ratio_pct", "top10_weighted_expense_ratio_pct", "low_fee_aum_share",
        "weighted_avg_spread_pct", "high_liquidity_product_count", "high_liquidity_aum_share",
        "top5_adv_concentration", "active_count_share", "active_aum_share", "fixed_income_aum_share",
        "equity_premium_income_aum_share", "smart_beta_factor_aum_share", "non_plain_beta_aum_share",
        "largest_etf_ticker", "largest_etf_name", "largest_etf_aum_usd", "largest_etf_adv_usd",
        "top1_aum_share", "top3_aum_share", "flagship_mind_tag_score", "data_as_of",
    ]
    methodology_headers = ["section", "item", "description", "formula_or_rule", "notes"]
    method_map_headers = [
        "dimension", "dimension_weight_in_total", "subindicator", "subindicator_weight_in_dimension",
        "direction", "transform_method", "raw_metric_source", "plain_language_explanation",
    ]
    qa_headers = ["check_item", "status", "value", "notes", "qa_explanation_cn"]
    chart_index_headers = [
        "chart_name", "source_sheet", "chart_type", "x_field", "y_field", "size_field",
        "category_field", "recommended_usage",
    ]

    ranking = project_rows(summary_rows, ranking_headers, round_scores=True)
    wide = project_rows(radar_wide_rows, wide_headers, round_scores=True)
    long = project_rows(radar_long_rows, long_headers, round_scores=True)
    sub_detail = add_dimension_notes(project_rows(sub_rows, sub_headers[:-1], round_scores=True))
    raw_metrics = project_rows(raw_rows, raw_headers, round_scores=False)
    methodology = build_methodology_summary()
    method_map = build_dimension_method_map()
    qa = build_qa_rows(qa_rows)
    chart_index = build_chart_index()
    readme = build_readme_rows()

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "00_README", ["item", "content", "data_as_of", "notes"], readme)
    write_sheet(wb, "01_TOP10_Score_Ranking", ranking_headers, ranking)
    write_sheet(wb, "02_Dimension_Score_Wide", wide_headers, wide)
    write_sheet(wb, "03_Dimension_Score_Long", long_headers, long)
    write_sheet(wb, "04_Subindicator_Score_Detail", sub_headers, sub_detail)
    write_sheet(wb, "05_Raw_Metrics_TOP10", raw_headers, raw_metrics)
    write_sheet(wb, "06_Methodology_Summary", methodology_headers, methodology)
    write_sheet(wb, "07_Dimension_Method_Map", method_map_headers, method_map)
    write_sheet(wb, "08_QA_Checks", qa_headers, qa)
    write_sheet(wb, "09_Chart_Data_Index", chart_index_headers, chart_index)
    wb.save(OUTPUT_PATH)

    if CSV_DIR.exists():
        shutil.rmtree(CSV_DIR)
    CSV_DIR.mkdir(parents=True, exist_ok=True)
    export_csv(CSV_DIR / "TOP10_Score_Ranking.csv", ranking_headers, ranking)
    export_csv(CSV_DIR / "Dimension_Score_Wide.csv", wide_headers, wide)
    export_csv(CSV_DIR / "Dimension_Score_Long.csv", long_headers, long)
    export_csv(CSV_DIR / "Subindicator_Score_Detail.csv", sub_headers, sub_detail)
    export_csv(CSV_DIR / "Dimension_Method_Map.csv", method_map_headers, method_map)
    selected = {"BlackRock/iShares", "JPMorgan AM", "Invesco"}
    three_selected = [row for row in long if row.get("issuer") in selected]
    export_csv(CSV_DIR / "Three_Selected_Radar.csv", long_headers, three_selected)

    verify = load_workbook(OUTPUT_PATH, data_only=True)
    csv_files = sorted(CSV_DIR.glob("*.csv"))
    qa_counts = {}
    for row in qa:
        qa_counts[row["status"]] = qa_counts.get(row["status"], 0) + 1

    print(f"输入文件路径: {INPUT_PATH}")
    print(f"输出文件路径: {OUTPUT_PATH}")
    print("每个输出sheet的行数和列数（数据行数，不含表头）:")
    for ws in verify.worksheets:
        print(f"- {ws.title}: {max(ws.max_row - 1, 0)} 行 x {ws.max_column} 列")
    print("TOP10发行人及总分排名:")
    for row in ranking:
        print(f"- {row['rank_total_score']}. {row['issuer']}: {row['total_ecosystem_score']:.1f}")
    expected_csv = {
        "TOP10_Score_Ranking.csv",
        "Dimension_Score_Wide.csv",
        "Dimension_Score_Long.csv",
        "Subindicator_Score_Detail.csv",
        "Dimension_Method_Map.csv",
        "Three_Selected_Radar.csv",
    }
    actual_csv = {path.name for path in csv_files}
    print(f"是否成功生成所有CSV: {'是' if expected_csv == actual_csv else '否'}")
    for path in csv_files:
        print(f"- {path}")
    print(f"QA中PASS/WARN/FAIL数量: PASS={qa_counts.get('PASS', 0)}, WARN={qa_counts.get('WARN', 0)}, FAIL={qa_counts.get('FAIL', 0)}")
    for row in qa:
        if row["status"] in {"WARN", "FAIL"}:
            print(f"- {row['status']}: {row['check_item']} - {row['notes']}")


if __name__ == "__main__":
    main()
