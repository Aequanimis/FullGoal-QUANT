from pathlib import Path
from openpyxl import load_workbook

path = Path.home() / "Desktop" / "Fullgoal" / "ICI课题" / "课题更正" / "华宝基金旗下53只ETF产品完整数据.xlsx"
print(path)
print("exists", path.exists())
if path.exists():
    wb = load_workbook(path, read_only=True, data_only=True)
    print("sheets", wb.sheetnames)
    for ws in wb.worksheets:
        print("sheet", ws.title, ws.max_row, ws.max_column)
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            print(i, row)
            if i >= 8:
                break
