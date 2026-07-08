from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SNAPSHOT_DATE = "2026-07-06"

INPUT_PATH = (
    Path.home()
    / "Desktop"
    / "Fullgoal"
    / "ICI课题"
    / "课题更正"
    / "美国ETF案例数据收集_富国启示_展示整理版_v2.xlsx"
)
OUTPUT_PATH = (
    Path.home()
    / "Desktop"
    / "Fullgoal"
    / "ICI课题"
    / "课题更正"
    / "美国ETF案例数据收集_富国启示_作图sheet版.xlsx"
)

PLOT_SHEETS = [
    "01_Issuer_Overview_Clean",
    "03_Representative_ETFs_Clean",
    "04_Official_Sources_Clean",
    "05_P23_Case_Map_Data",
    "06_P24_iShares_OS_Data",
    "07_P25_JPMorgan_Migration_Data",
    "08_P25_Invesco_Helper_Data",
    "09_P26_Fullgoal_Implication_Data",
    "10_Slide_Text_Prompts",
    "11_QA_Checks",
]

README_ROWS = [
    ("文件名称", "美国ETF案例数据收集_富国启示_作图sheet版"),
    ("数据截至日期", SNAPSHOT_DATE),
    ("文件用途", "P23-P26作图和页面生成专用数据包"),
    ("包含页面", "P23 美国案例：ETF生态打造的三类路径"),
    ("包含页面", "P24 BlackRock/iShares：把ETF做成投顾账户的配置操作系统"),
    ("包含页面", "P25 JPMorgan AM：传统资管机构如何把能力迁移到ETF"),
    ("包含页面", "P26 对富国的启示：从产品铺设转向能力封装"),
    ("注意事项", "05_P23_Case_Map_Data中的坐标为研究判断，不是硬数据；"),
    ("注意事项", "JPMorgan主动ETF AUM占比基于本数据集strategy=Active分类口径；"),
    ("注意事项", "金额字段单位为美元，aum_billion_usd为十亿美元，aum_100m_usd为亿美元。"),
]


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False

    for cell in ws[1]:
        cell.fill = copy(header_fill)
        cell.font = copy(header_font)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = copy(border)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = copy(border)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col_idx, header_cell in enumerate(ws[1], start=1):
        header = str(header_cell.value or "")
        lower = header.lower()
        max_len = len(header)
        for col in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
            for cell in col:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))

        is_long_text = any(
            token in lower
            for token in [
                "url",
                "source",
                "source_note",
                "description",
                "annotation",
                "message",
                "text",
                "claim",
                "instruction",
                "basis",
                "suggestion",
                "followup",
            ]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 42 if is_long_text else 26)

        if "usd" in lower:
            number_format = "$#,##0"
            if lower.endswith("billion_usd") or lower.endswith("100m_usd") or lower.endswith("trillion_usd"):
                number_format = "0.00"
            for col in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
                for cell in col:
                    cell.number_format = number_format
                    cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=is_long_text)
        elif lower.endswith("_pct") or lower == "aum_share":
            number_format = "0.00%" if lower == "aum_share" else '0.00"%"'
            for col in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
                for cell in col:
                    cell.number_format = number_format
                    cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=is_long_text)
        elif is_long_text:
            for col in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
                for cell in col:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row_idx in range(1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 22 if row_idx == 1 else 34


def copy_values_and_formats(src_ws, dst_ws):
    for row in src_ws.iter_rows():
        for src_cell in row:
            dst_cell = dst_ws.cell(src_cell.row, src_cell.column, src_cell.value)
            if src_cell.has_style:
                dst_cell.font = copy(src_cell.font)
                dst_cell.fill = copy(src_cell.fill)
                dst_cell.border = copy(src_cell.border)
                dst_cell.alignment = copy(src_cell.alignment)
                dst_cell.number_format = src_cell.number_format
                dst_cell.protection = copy(src_cell.protection)

    for merged_range in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merged_range))


def header_map(ws):
    return {cell.value: cell.column for cell in ws[1] if cell.value}


def count_empty_sources(wb, sheet_names):
    empties = []
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        headers = header_map(ws)
        source_cols = [col for key, col in headers.items() if str(key).lower() in {"source", "source_note"}]
        for col in source_cols:
            header = ws.cell(1, col).value
            for row_idx in range(2, ws.max_row + 1):
                if ws.cell(row_idx, col).value in (None, ""):
                    empties.append((sheet_name, header, row_idx))
    return empties


def find_bad_data_as_of(wb, sheet_names):
    bad = []
    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        headers = header_map(ws)
        for key, col in headers.items():
            if str(key) in {"data_as_of", "data_as_of_or_access_date"}:
                for row_idx in range(2, ws.max_row + 1):
                    value = ws.cell(row_idx, col).value
                    if value != SNAPSHOT_DATE:
                        bad.append((sheet_name, key, row_idx, value))
    return bad


def count_qa_pass(wb):
    ws = wb["11_QA_Checks"]
    headers = header_map(ws)
    status_col = headers.get("status")
    if not status_col:
        return 0
    return sum(1 for row_idx in range(2, ws.max_row + 1) if ws.cell(row_idx, status_col).value == "PASS")


def main():
    if not INPUT_PATH.exists():
        raise FileNotFoundError(f"输入文件不存在：{INPUT_PATH}")

    src_wb = load_workbook(INPUT_PATH, data_only=False)
    missing = [sheet for sheet in PLOT_SHEETS if sheet not in src_wb.sheetnames]
    if missing:
        raise ValueError(f"输入文件缺少必要sheet：{missing}")

    out_wb = Workbook()
    default_sheet = out_wb.active
    out_wb.remove(default_sheet)

    readme = out_wb.create_sheet("README")
    readme.append(["项目", "内容"])
    for row in README_ROWS:
        readme.append(row)
    style_sheet(readme)

    for sheet_name in PLOT_SHEETS:
        dst_ws = out_wb.create_sheet(sheet_name)
        copy_values_and_formats(src_wb[sheet_name], dst_ws)
        style_sheet(dst_ws)

    out_wb.save(OUTPUT_PATH)

    # Reopen to verify the exact saved artifact.
    verify_wb = load_workbook(OUTPUT_PATH, data_only=True)
    all_sheets = ["README"] + PLOT_SHEETS
    empty_sources = count_empty_sources(verify_wb, PLOT_SHEETS)
    bad_data_as_of = find_bad_data_as_of(verify_wb, PLOT_SHEETS)
    qa_pass_count = count_qa_pass(verify_wb)

    print(f"输出文件路径: {OUTPUT_PATH}")
    print("\n每个sheet的行数和列数（数据行数，不含表头）:")
    for sheet_name in all_sheets:
        ws = verify_wb[sheet_name]
        print(f"- {sheet_name}: {max(ws.max_row - 1, 0)} 行 x {ws.max_column} 列")
    print(f"\nQA_Checks中PASS数量: {qa_pass_count}")
    print(f"是否存在空source/source_note字段: {'是' if empty_sources else '否'}")
    if empty_sources:
        for sheet_name, header, row_idx in empty_sources[:20]:
            print(f"- {sheet_name}!{header} 第{row_idx}行为空")
        if len(empty_sources) > 20:
            print(f"- 其余 {len(empty_sources) - 20} 处略")
    print(f"是否存在data_as_of不等于{SNAPSHOT_DATE}的情况: {'是' if bad_data_as_of else '否'}")
    if bad_data_as_of:
        for sheet_name, header, row_idx, value in bad_data_as_of[:20]:
            print(f"- {sheet_name}!{header} 第{row_idx}行 = {value!r}")
        if len(bad_data_as_of) > 20:
            print(f"- 其余 {len(bad_data_as_of) - 20} 处略")


if __name__ == "__main__":
    main()
