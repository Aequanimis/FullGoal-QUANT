from __future__ import annotations

import math
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


WORK_DIR = Path(r"C:\Users\lvdon\Desktop\Fullgoal\日常工作\7.8")
MAIN_XLSX = WORK_DIR / "境外ETF全量数据_轻量分析版_v1_中文表头.xlsx"
SUPPLEMENT_XLSX = WORK_DIR / "数据收集.xlsx"
OUTPUT_XLSX = WORK_DIR / "境外ETF全量数据_轻量分析版_v2_补充YTD总回报.xlsx"
OUTPUT_LOG = WORK_DIR / "境外ETF全量数据_补充YTD总回报_整合日志.txt"

MAIN_SHEET = "ETF主表_轻量版"
YEARLY_SHEET = "年度生命周期统计"
NOTES_SHEET = "数据口径说明"
SUPPLEMENT_SHEET = "Sheet1"
REQUIRED_SHEETS = [MAIN_SHEET, YEARLY_SHEET, NOTES_SHEET]

KEY_MAIN = "ETF代码"
KEY_SUPPLEMENT = "Ticker"
SOURCE_FIELD = "YTD Total Return"
TARGET_FIELD = "年初至今总回报（%）"
INSERT_AFTER_FIELD = "近5年总回报（%）"
SAMPLE_TICKERS = ["VOO", "IVV", "SPY", "VTI", "QQQ"]


def normalize_ticker(value: Any) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip().upper()


def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    text = str(value).strip()
    return text == "" or text in {"-", "--", "N/A", "n/a", "nan", "NaN"}


def parse_ytd_total_return(value: Any) -> float | None:
    """Return percentage-point numeric value for YTD Total Return.

    This helper is intentionally reusable for later return fields. Values with
    abs(value) <= 2 are treated as decimal returns and multiplied by 100.
    Values already larger than 2 are treated as percentage-point numbers.
    """
    if is_blank(value):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
    else:
        text = str(value).strip().replace(",", "").replace("−", "-").replace("—", "-").replace("–", "-")
        if text.endswith("%"):
            text = text[:-1].strip()
        text = re.sub(r"^\((.*)\)$", r"-\1", text)
        try:
            number = float(text)
        except ValueError:
            return None
    if abs(number) <= 2:
        number *= 100
    return round(number, 2)


def first_nonblank_ytd(values: pd.Series) -> float | None:
    for value in values:
        parsed = parse_ytd_total_return(value)
        if parsed is not None:
            return parsed
    return None


def write_workbook(main_df: pd.DataFrame, yearly_df: pd.DataFrame, notes_df: pd.DataFrame) -> None:
    with pd.ExcelWriter(OUTPUT_XLSX, engine="xlsxwriter") as writer:
        main_df.to_excel(writer, sheet_name=MAIN_SHEET, index=False)
        yearly_df.to_excel(writer, sheet_name=YEARLY_SHEET, index=False)
        notes_df.to_excel(writer, sheet_name=NOTES_SHEET, index=False)

        workbook = writer.book
        header_fmt = workbook.add_format({"bold": True, "align": "center", "valign": "vcenter"})
        numeric_fmt = workbook.add_format({"num_format": "0.00"})
        date_fmt = workbook.add_format({"num_format": "yyyy-mm-dd"})

        for sheet_name, df in [(MAIN_SHEET, main_df), (YEARLY_SHEET, yearly_df), (NOTES_SHEET, notes_df)]:
            worksheet = writer.sheets[sheet_name]
            worksheet.freeze_panes(1, 0)
            worksheet.autofilter(0, 0, len(df), max(len(df.columns) - 1, 0))
            for col_idx, col_name in enumerate(df.columns):
                width = min(max(len(str(col_name)) + 2, 10), 42)
                if len(df) > 0:
                    sample = df[col_name].head(300)
                    sample_width = int(sample.map(lambda value: 0 if pd.isna(value) else len(str(value))).max())
                    width = min(max(width, sample_width + 2), 42)
                fmt = None
                if col_name == TARGET_FIELD:
                    fmt = numeric_fmt
                elif col_name in {"成立日期", "清盘日期"}:
                    fmt = date_fmt
                worksheet.set_column(col_idx, col_idx, width, fmt)
                worksheet.write(0, col_idx, col_name, header_fmt)


def main() -> None:
    if not MAIN_XLSX.exists():
        raise FileNotFoundError(f"主文件不存在: {MAIN_XLSX}")
    if not SUPPLEMENT_XLSX.exists():
        raise FileNotFoundError(f"补充文件不存在: {SUPPLEMENT_XLSX}")

    completed_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    source_sheets = pd.ExcelFile(MAIN_XLSX).sheet_names
    missing_sheets = [sheet for sheet in REQUIRED_SHEETS if sheet not in source_sheets]
    if missing_sheets:
        raise ValueError(f"主文件缺少Sheet: {', '.join(missing_sheets)}")

    main_df = pd.read_excel(MAIN_XLSX, sheet_name=MAIN_SHEET)
    yearly_df = pd.read_excel(MAIN_XLSX, sheet_name=YEARLY_SHEET)
    notes_df = pd.read_excel(MAIN_XLSX, sheet_name=NOTES_SHEET)
    supplement_df = pd.read_excel(SUPPLEMENT_XLSX, sheet_name=SUPPLEMENT_SHEET)
    supplement_wb = load_workbook(SUPPLEMENT_XLSX, read_only=True, data_only=True)
    supplement_sheet_rows = supplement_wb[SUPPLEMENT_SHEET].max_row - 1

    for required_col in [KEY_MAIN]:
        if required_col not in main_df.columns:
            raise ValueError(f"主表缺少必要字段: {required_col}")
    for required_col in [KEY_SUPPLEMENT, SOURCE_FIELD]:
        if required_col not in supplement_df.columns:
            raise ValueError(f"补充表缺少必要字段: {required_col}")

    main_total_rows = len(main_df)
    supplement_total_rows = supplement_sheet_rows
    original_main_columns = list(main_df.columns)

    supplement_df["_ticker_norm"] = supplement_df[KEY_SUPPLEMENT].map(normalize_ticker)
    supplement_nonblank_ticker = supplement_df[supplement_df["_ticker_norm"].ne("")].copy()
    ytd_blank_count = int(supplement_nonblank_ticker[SOURCE_FIELD].map(is_blank).sum())

    duplicate_counts = supplement_nonblank_ticker["_ticker_norm"].value_counts()
    duplicate_tickers = sorted(duplicate_counts[duplicate_counts > 1].index.tolist())

    # One value per supplement ticker: keep the first non-empty converted YTD.
    ytd_mapping = (
        supplement_nonblank_ticker.groupby("_ticker_norm", sort=False)[SOURCE_FIELD]
        .apply(first_nonblank_ytd)
        .dropna()
        .to_dict()
    )

    main_df["_ticker_norm"] = main_df[KEY_MAIN].map(normalize_ticker)
    main_tickers = set(main_df.loc[main_df["_ticker_norm"].ne(""), "_ticker_norm"])
    supplement_only_tickers = sorted(set(supplement_nonblank_ticker["_ticker_norm"]) - main_tickers)

    ytd_series = main_df["_ticker_norm"].map(ytd_mapping)
    matched_and_filled = int(ytd_series.notna().sum())
    main_unmatched_or_blank = int(main_total_rows - matched_and_filled)

    if TARGET_FIELD in main_df.columns:
        main_df[TARGET_FIELD] = ytd_series
    else:
        if INSERT_AFTER_FIELD not in main_df.columns:
            raise ValueError(f"主表缺少插入定位字段: {INSERT_AFTER_FIELD}")
        insert_at = list(main_df.columns).index(INSERT_AFTER_FIELD) + 1
        main_df.insert(insert_at, TARGET_FIELD, ytd_series)

    main_df = main_df.drop(columns=["_ticker_norm"])
    # Keep all original fields and only add/update the target field.
    if TARGET_FIELD not in original_main_columns:
        ordered_cols = list(original_main_columns)
        ordered_cols.insert(ordered_cols.index(INSERT_AFTER_FIELD) + 1, TARGET_FIELD)
        main_df = main_df[ordered_cols]

    write_workbook(main_df, yearly_df, notes_df)

    # Validation after writing.
    written_sheets = pd.ExcelFile(OUTPUT_XLSX).sheet_names
    out_main = pd.read_excel(OUTPUT_XLSX, sheet_name=MAIN_SHEET)
    output_exists = OUTPUT_XLSX.exists()
    sheets_retained = all(sheet in written_sheets for sheet in REQUIRED_SHEETS)
    row_count_ok = len(out_main) == main_total_rows
    target_exists = TARGET_FIELD in out_main.columns
    sample_values = {}
    for ticker in SAMPLE_TICKERS:
        rows = out_main[out_main[KEY_MAIN].map(normalize_ticker).eq(ticker)]
        sample_values[ticker] = "未找到" if rows.empty else rows.iloc[0][TARGET_FIELD]
    non_numeric_count = int(
        out_main[TARGET_FIELD].dropna().map(lambda value: not isinstance(value, (int, float))).sum()
    )

    log_lines = [
        "境外ETF全量数据补充 YTD Total Return 整合日志",
        "=" * 55,
        f"主文件路径: {MAIN_XLSX}",
        f"补充文件路径: {SUPPLEMENT_XLSX}",
        f"输出文件路径: {OUTPUT_XLSX}",
        f"主表总行数: {main_total_rows}",
        f"补充表总行数: {supplement_total_rows}",
        f"主表成功匹配并补充 YTD Total Return 的数量: {matched_and_filled}",
        f"主表未匹配到补充数据的数量: {main_unmatched_or_blank}",
        f"补充表中存在但主表不存在的Ticker数量: {len(supplement_only_tickers)}",
        "补充表中存在但主表不存在的Ticker前50个示例: "
        + (", ".join(supplement_only_tickers[:50]) if supplement_only_tickers else "无"),
        f"补充表重复Ticker数量: {len(duplicate_tickers)}",
        "补充表重复Ticker前50个示例: " + (", ".join(duplicate_tickers[:50]) if duplicate_tickers else "无"),
        f"YTD Total Return 空值数量: {ytd_blank_count}",
        "YTD Total Return 数值转换规则说明: 绝对值 <= 2 的值按小数收益率处理并乘以100；绝对值 > 2 的值视为已是百分比数值，不再乘以100；最终写入数值并保留两位小数格式。",
        f"输出文件路径: {OUTPUT_XLSX}",
        f"处理完成时间: {completed_at}",
        "",
        "自动校验结果:",
        f"1. 输出文件是否存在: {'是' if output_exists else '否'}",
        f"2. 主表行数是否与原主表一致: {'是' if row_count_ok else '否'}，输出主表行数={len(out_main)}",
        f"3. 原有Sheet是否都保留: {'是' if sheets_retained else '否'}，Sheet列表={', '.join(written_sheets)}",
        f"4. 是否存在字段“{TARGET_FIELD}”: {'是' if target_exists else '否'}",
        "5. 抽查Ticker的YTD Total Return: " + ", ".join(f"{ticker}={sample_values[ticker]}" for ticker in SAMPLE_TICKERS),
        f"6. 字段是否为数值型: {'是' if non_numeric_count == 0 else '否'}，非数值数量={non_numeric_count}",
    ]
    OUTPUT_LOG.write_text("\n".join(log_lines) + "\n", encoding="utf-8-sig")

    print("YTD Total Return 补充完成")
    print(f"主表总行数: {main_total_rows}")
    print(f"补充表总行数: {supplement_total_rows}")
    print(f"成功补充数量: {matched_and_filled}")
    print(f"未匹配/无可用补充值数量: {main_unmatched_or_blank}")
    print(f"补充表存在但主表不存在Ticker数量: {len(supplement_only_tickers)}")
    print(f"补充表重复Ticker数量: {len(duplicate_tickers)}")
    print(f"YTD Total Return 空值数量: {ytd_blank_count}")
    print(f"输出文件: {OUTPUT_XLSX}")
    print(f"日志文件: {OUTPUT_LOG}")
    print("抽查Ticker:")
    for ticker in SAMPLE_TICKERS:
        print(f"- {ticker}: {sample_values[ticker]}")
    print(f"输出文件存在: {'是' if output_exists else '否'}")
    print(f"主表行数一致: {'是' if row_count_ok else '否'}")
    print(f"原有Sheet保留: {'是' if sheets_retained else '否'}")
    print(f"字段为数值型: {'是' if non_numeric_count == 0 else '否'}")


if __name__ == "__main__":
    main()
